"""创建完全独立、明确标记为模拟数据的演示输入，不触碰真实企业档案。

用法：python scripts/prepare_demo.py --output data/demo
输出仅包含虚构资料；通过系统正常导入、确认、生成、审查、导出。
"""
from __future__ import annotations

import argparse
import json
import zipfile
from pathlib import Path

import docx
from jb_kb.models import (
    Boilerplate,
    Certificate,
    CompanyProfile,
    FinancialYear,
    Performance,
    Person,
)
from jb_parser.trm import (
    TRM,
    FormatBlock,
    PackageTRM,
    ScoringItem,
    ScoringRef,
    ScoringTemplate,
    SubmissionItem,
)
from PIL import Image, ImageDraw, ImageFont

DEMO_NOTICE = "演示样例 · 模拟数据 · 不可用于正式投标"


def prepare(output: Path) -> None:
    output.mkdir(parents=True, exist_ok=True)
    company = "演示企业（模拟数据）"
    profile = CompanyProfile(
        name=company, is_demo=True, credit_code="DEMO-ONLY-NOT-A-REAL-CODE",
        legal_person="演示负责人", legal_person_gender="男", legal_person_id_no="演示证件号（非真实号码）",
        legal_person_since="2020-01-01", sgcc_relationship="演示声明：无关联（模拟）",
        authorized_rep="演示经办人", authorized_rep_title="项目经办人（模拟）",
        founded="2020-01-01", registered_capital_wan=100, company_type="模拟企业",
        address="演示地址（非真实经营地址）", bank="演示开户银行/演示账号",
        contact="演示联系人", phone="演示电话", staff_total=10, staff_technical=8,
        senior_engineers=1, engineers=3, business_scope="演示运维服务流程",
        sources=[DEMO_NOTICE],
        certificates=[Certificate(name="营业执照（演示）", includes_tax_registration=True,
                                  validity_kind="permanent", source=DEMO_NOTICE),
                      Certificate(name="基本存款账户信息（演示）", validity_kind="permanent", source=DEMO_NOTICE),
                      Certificate(name="演示资质说明", cert_type="资质等级", validity_kind="permanent", source=DEMO_NOTICE),
                      Certificate(name="质量管理体系认证证书（演示）", cert_type="体系认证", number="DEMO-Q",
                                  issuer="演示机构", valid_from="2025-01-01", valid_until="2030-12-31", source=DEMO_NOTICE),
                      *[Certificate(name=name+"（演示）", cert_type="信用查询材料", validity_kind="point_in_time",
                                    valid_from="2026-10-22", source=DEMO_NOTICE)
                        for name in ("信用中国查询", "企业信用信息公示报告", "中国裁判文书网查询")]],
        personnel=[Person(name="演示工程师", title="工程师", major="计算机", education="本科", available=True,
                          social_insurance_unit=company, social_insurance_months=["2026-08", "2026-09", "2026-10"],
                          credentials=["身份证（演示）", "学历证（演示）", "职称证（演示）", "社保（演示）", "劳动合同（演示）"], source=DEMO_NOTICE)],
        performances=[Performance(project="演示运维服务历史项目（模拟）", buyer="演示最终用户（模拟）",
                                  buyer_is_end_user=True, amount_wan=10, signed_date="2025-01-01",
                                  commissioned_date="2025-12-01", evidence=["合同", "发票", "验收报告"], source=DEMO_NOTICE)],
        financials=[FinancialYear(year=str(year), revenue_wan=100, net_profit_wan=10,
                                 asset_wan=100, liability_ratio="30%", source=DEMO_NOTICE)
                    for year in (2023, 2024, 2025)],
        boilerplates=[Boilerplate(topic=topic, title=topic+"演示方案", approved=True, source=DEMO_NOTICE,
                                 text=text) for topic, text in (
            ("售后服务", "演示服务流程：登记服务请求、确认影响范围、安排处理、复核处理结果并保存记录。"),
            ("质量保证", "演示质量流程：执行前核对任务，执行中保留记录，完成后按约定项目检查并登记整改结果。"),
            ("培训", "演示培训流程：说明操作步骤、安排实际演练、记录问题并确认培训材料已交付。"))])
    pkg = PackageTRM(sub_no="演示分标", sub_name="演示运维服务", pkg_no="包1",
                     project_name="演示运维服务项目（模拟）",
                     scoring_ref=ScoringRef(biz_template="演示商务", tech_template="演示技术", weight_biz=30, weight_tech=70, weight_price=0))
    trm = TRM(batch_name="客户演示项目（模拟数据）", batch_no="DEMO20261022", packages=[pkg],
              format_blocks=[FormatBlock(title="投标函", paragraphs=["致：演示采购单位（模拟）", "招标编号：，项目名称：，分标名称：，包名称：招标文件"]),
                             FormatBlock(title="授权委托书", paragraphs=["演示授权文本，仅用于流程演示，不构成真实授权。演示负责人授权演示经办人展示系统操作。"]),
                             FormatBlock(title="信用承诺书", paragraphs=[DEMO_NOTICE])],
              submission_table=[SubmissionItem(section="商务文件", item="商务文件", channels=["演示通道"], port="演示商务端口"),
                                SubmissionItem(section="技术文件", item="技术文件", channels=["演示通道"], port="演示技术端口")],
              scoring_templates=[ScoringTemplate(name="演示商务", kind="biz", items=[ScoringItem(element="材料完整性（0-10分）", content="演示人工复核项", score_min=0, score_max=10)]),
                                 ScoringTemplate(name="演示技术", kind="tech", items=[ScoringItem(element="服务方案（0-10分）", content="演示人工复核项", score_min=0, score_max=10)])])
    trm.key_terms.bid_deadline = "2026-10-31 09:00"
    trm.key_terms.bid_open_time = "2026-10-31 09:00"
    profile.save(str(output / "profile.json"))
    # 仅生成流程演示用说明页，不仿制证照、印章或真实查询页面。
    image = Image.new("RGB", (1100, 700), "white")
    draw = ImageDraw.Draw(image)
    fonts = [Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
             Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")]
    font_path = next((p for p in fonts if p.exists()), None)
    font = ImageFont.truetype(str(font_path), 30) if font_path else ImageFont.load_default()
    lines = ["DEMO ONLY - SYNTHETIC DATA", "NOT A CERTIFICATE OR VALID BID EVIDENCE",
             "This page demonstrates attachment insertion only.",
             "All names, amounts and declarations are fictional.",
             "No seal, signature or real authority is represented.",
             "演示样例 / 模拟数据 / 不可用于正式投标"]
    for i, line in enumerate(lines):
        draw.text((45, 65 + i * 90), line, font=font, fill="#b42318" if i in (0, 5) else "#344054")
    image.save(output / "模拟附件说明.png")
    (output / "trm.json").write_text(trm.model_dump_json(indent=2), encoding="utf-8")
    d = docx.Document()
    d.add_heading("客户演示招标输入", 0)
    d.add_paragraph(DEMO_NOTICE)
    d.add_paragraph("本文件为演示操作编制的模拟输入，没有真实采购或投标效力。")
    d.save(output / "演示说明.docx")
    with zipfile.ZipFile(output / "演示输入.zip", "w") as archive:
        archive.write(output / "演示说明.docx", "演示说明.docx")
    (output / "notice.json").write_text(json.dumps({"notice": DEMO_NOTICE, "company": company}, ensure_ascii=False), encoding="utf-8")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("data/demo"))
    prepare(parser.parse_args().output)
