"""jb-api：Jinbang 后端 HTTP 层。

薄 API 原则：业务逻辑全部在 packages/，本层只做路由、参数校验、任务调度。
存储：jb_store（SQLite 本地 / PostgreSQL compose）；长任务：jb_api.tasks（Celery 或进程内）。

启动：uvicorn jb_api.main:app --reload --port 8000
"""
from __future__ import annotations

import datetime as dt
import os
import shutil
from contextlib import asynccontextmanager
from typing import Optional

from fastapi import BackgroundTasks, FastAPI, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from jb_agents import qualify
from jb_kb import repo as kb_repo
from jb_parser.trm import TRM
from jb_store import Project, ProjectEvent, Task, init_db, session
from jb_store import auth as au
from jb_store import projects as pm
from sqlalchemy import select

from . import auth as auth_api
from . import config as config_api
from . import kb, tasks

UPLOAD_DIR = os.environ.get("UPLOAD_DIR", os.path.join(os.getcwd(), "uploads"))


@asynccontextmanager
async def lifespan(_: FastAPI):
    init_db()  # 开发便利；生产以 alembic upgrade head 为准
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    config_api.install_usage_hook()     # LLM 调用记账
    config_api.load_llm_settings()      # 配置中心保存的端点/模型覆盖
    if auth_api.bootstrap():
        print("[jb-api] 已创建默认管理员 admin（密码 admin，首次登录请修改）")
    yield


app = FastAPI(title="Jinbang API", version="0.2.0", lifespan=lifespan)
app.add_middleware(
    CORSMiddleware,
    allow_origins=os.environ.get("CORS_ORIGINS", "http://localhost:5173").split(","),
    allow_methods=["*"],
    allow_headers=["*"],
)
app.middleware("http")(auth_api.auth_middleware)   # 除 health/login 外全部要求登录
app.include_router(auth_api.router)     # 登录 / 用户管理：/api/auth/*、/api/users/*
app.include_router(kb.router)   # 企业知识库：/api/profiles/*、/api/attachments/*
app.include_router(config_api.router)   # 配置中心：/api/config/*


def _project_out(p: Project) -> dict:
    trm = p.trm_confirmed or p.trm or {}
    pkgs = trm.get("packages") or []
    return {"id": p.id, "filename": p.filename, "status": p.status, "batch_name": p.batch_name,
            "batch_no": p.batch_no, "error": p.error, "created_at": p.created_at.isoformat(),
            "updated_at": p.updated_at.isoformat(), "confirmed": p.trm_confirmed is not None,
            "deadline": p.deadline, "open_time": p.open_time, "deadline_manual": p.deadline_manual,
            "days_left": pm.days_left(p.deadline), "stage": p.stage, "stage_cn": pm.STAGE_CN.get(p.stage, p.stage),
            "outcome": p.outcome, "outcome_cn": pm.OUTCOME_CN.get(p.outcome, p.outcome), "active": pm.is_active(p),
            "notes": p.notes, "pkg_nos": [k.get("pkg_no", "") for k in pkgs], "packages": len(pkgs),
            "results": p.results or {}}


def _get_project(s, project_id: str) -> Project:
    p = s.get(Project, project_id)
    if not p:
        raise HTTPException(404, "项目不存在")
    return p


@app.get("/api/health")
def health():
    return {"status": "ok", "service": "jb-api", "task_mode": "celery" if tasks.celery_app else "inprocess"}


# ---------- 项目 / 解析 ----------

@app.post("/api/projects", status_code=202)
async def create_project(file: UploadFile, background: BackgroundTasks):
    """上传招标文件包 → 建项目 → 异步解析。返回 project_id + task_id 供轮询。"""
    if not (file.filename or "").lower().endswith(".zip"):
        raise HTTPException(400, "请上传招标文件包 zip")
    with session() as s:
        p = Project(filename=file.filename, zip_path="")
        s.add(p)
        s.flush()
        pdir = os.path.join(UPLOAD_DIR, p.id)
        os.makedirs(pdir, exist_ok=True)
        p.zip_path = os.path.join(pdir, file.filename)
        with open(p.zip_path, "wb") as f:
            shutil.copyfileobj(file.file, f)
        t = Task(project_id=p.id, kind="parse", actor=au.current_actor())
        s.add(t)
        s.flush()
        pid, tid = p.id, t.id
    mode = tasks.submit_parse(background, tid, pid)
    return {"id": pid, "task_id": tid, "mode": mode}


@app.get("/api/projects")
def list_projects(active: Optional[bool] = None):
    """项目列表：在投按截止日升序在前，已结束在后；active 过滤在投/已结束。"""
    with session() as s:
        rows = s.execute(select(Project)).scalars().all()
        if active is not None:
            rows = [p for p in rows if pm.is_active(p) == active]
        return [_project_out(p) for p in sorted(rows, key=pm.sort_key)]


@app.patch("/api/projects/{project_id}")
def patch_project(project_id: str, fields: dict):
    """项目管理字段：deadline / open_time（"YYYY-MM-DD HH:MM"）、outcome、notes。"""
    with session() as s:
        p = _get_project(s, project_id)
        for key in ("deadline", "open_time"):
            if key in fields:
                val = (fields[key] or "").strip()
                if val and pm.parse_deadline(val) is None:
                    raise HTTPException(400, f"{key} 须为 YYYY-MM-DD HH:MM")
                if getattr(p, key) != val:
                    pm.log_event(s, project_id, "deadline", f"{'截止' if key == 'deadline' else '开标'}时间改为 {val or '（清空）'}")
                setattr(p, key, val)
                p.deadline_manual = True
        if "outcome" in fields:
            val = fields["outcome"] or ""
            if val not in pm.OUTCOMES:
                raise HTTPException(400, "outcome 取值：submitted/won/lost/abandoned 或空")
            if val != p.outcome:
                p.outcome = val
                if val == "submitted":
                    pm.advance(p, "submitted")
                pm.log_event(s, project_id, "outcome", f"结果标记为「{pm.OUTCOME_CN[val]}」")
        if "notes" in fields:
            p.notes = fields["notes"] or ""
        return _project_out(p)


@app.delete("/api/projects/{project_id}")
def delete_project(project_id: str):
    """删除项目及其任务、事件、上传/产出文件。"""
    with session() as s:
        p = _get_project(s, project_id)
        for t in s.execute(select(Task).where(Task.project_id == project_id)).scalars():
            s.delete(t)
        for e in s.execute(select(ProjectEvent).where(ProjectEvent.project_id == project_id)).scalars():
            s.delete(e)
        s.delete(p)
    shutil.rmtree(os.path.join(UPLOAD_DIR, project_id), ignore_errors=True)
    return {"ok": True}


@app.get("/api/projects/{project_id}/detail")
def project_detail(project_id: str):
    """项目详情：概览 + 各包摘要 + 各环节结果 + 时间线 + 任务 + 产出文件。"""
    with session() as s:
        p = _get_project(s, project_id)
        out = _project_out(p)
        trm = p.trm_confirmed or p.trm or {}
        out["key_terms"] = trm.get("key_terms") or {}
        out["package_list"] = [{"pkg_no": k.get("pkg_no"), "sub_no": k.get("sub_no"), "sub_name": k.get("sub_name"),
                                "project_name": k.get("project_name"), "budget_yuan": k.get("budget_yuan"),
                                "max_price": k.get("max_price"), "materials": len(k.get("materials") or []),
                                "spec_docs": len(k.get("spec_docs") or [])} for k in trm.get("packages") or []]
        out["events"] = [{"id": e.id, "kind": e.kind, "message": e.message, "data": e.data, "actor": e.actor,
                          "created_at": e.created_at.isoformat()} for e in pm.events(s, project_id)]
        tasks_rows = s.execute(select(Task).where(Task.project_id == project_id).order_by(Task.created_at.desc())).scalars()
        out["tasks"] = [{"id": t.id, "kind": t.kind, "status": t.status, "progress": t.progress,
                         "message": t.message[:200], "created_at": t.created_at.isoformat()} for t in tasks_rows]
    out_dir = os.path.join(UPLOAD_DIR, project_id, "out")
    files = []
    if os.path.isdir(out_dir):
        for fn in sorted(os.listdir(out_dir)):
            st = os.stat(os.path.join(out_dir, fn))
            files.append({"name": fn, "size": st.st_size, "modified": dt.datetime.utcfromtimestamp(st.st_mtime).isoformat()})
    out["files"] = files
    return out


@app.get("/api/projects/{project_id}")
def get_project(project_id: str):
    with session() as s:
        p = s.get(Project, project_id)
        if not p:
            raise HTTPException(404, "项目不存在")
        return _project_out(p)


@app.get("/api/projects/{project_id}/trm")
def get_trm(project_id: str, confirmed: bool = False):
    with session() as s:
        p = s.get(Project, project_id)
        if not p:
            raise HTTPException(404, "项目不存在")
        data = p.trm_confirmed if confirmed and p.trm_confirmed else p.trm
        if data is None:
            raise HTTPException(409, f"TRM 尚未就绪（状态 {p.status}）")
        return data


@app.put("/api/projects/{project_id}/trm")
def confirm_trm(project_id: str, trm: dict):
    """确认页提交：保存人工确认版 TRM（整体覆盖；Schema 校验）。"""
    validated = TRM.model_validate(trm)
    with session() as s:
        p = s.get(Project, project_id)
        if not p:
            raise HTTPException(404, "项目不存在")
        p.trm_confirmed = validated.model_dump()
        p.status = "confirmed"
        pm.advance(p, "confirmed")
        pm.log_event(s, project_id, "confirmed", "TRM 人工确认")
    return {"ok": True}


@app.get("/api/tasks/{task_id}")
def get_task(task_id: str):
    with session() as s:
        t = s.get(Task, task_id)
        if not t:
            raise HTTPException(404, "任务不存在")
        return {"id": t.id, "project_id": t.project_id, "kind": t.kind, "status": t.status,
                "progress": t.progress, "message": t.message, "result": t.result}


# ---------- 资格自检（企业档案路由见 kb.py） ----------

@app.post("/api/projects/{project_id}/qualify")
def qualify_project(project_id: str, profile: str, llm: bool = False):
    """资格自检：优先用人工确认版 TRM。"""
    trm, cp = _load_trm_profile(project_id, profile)
    if llm:
        from jb_parser.llm_fallback import enrich
        enrich(trm)
    rep = qualify(trm, cp, use_llm=llm)
    with session() as s:
        p = _get_project(s, project_id)
        verdicts = {pk.pkg_no: pk.verdict for pk in rep.packages}
        pm.set_result(p, "qualify", {"profile": profile, "verdicts": verdicts, "llm": llm})
        pm.advance(p, "qualified")
        pm.log_event(s, project_id, "qualified", "资格自检：" + "，".join(f"{k} {v}" for k, v in verdicts.items()),
                     {"profile": profile, "verdicts": verdicts})
    return {"report": rep.model_dump(), "markdown": rep.markdown()}


# ---------- 文件生成 ----------

def _load_trm_profile(project_id: str, profile: str):
    with session() as s:
        p = s.get(Project, project_id)
        if not p:
            raise HTTPException(404, "项目不存在")
        data = p.trm_confirmed or p.trm
        if data is None:
            raise HTTPException(409, "TRM 尚未就绪")
        cp = kb_repo.load_profile(s, profile)
        if cp is None:
            raise HTTPException(404, "企业档案不存在")
        return TRM.model_validate(data), cp


@app.post("/api/projects/{project_id}/generate", status_code=202)
def generate_docs(project_id: str, profile: str, background: BackgroundTasks, pkg_index: int = 0,
                  product_model: Optional[str] = None, with_draft: bool = False):
    """生成商务/技术文件（异步任务；with_draft 时先由写作 Agent 起草）。轮询 /api/tasks/{id}。"""
    _load_trm_profile(project_id, profile)  # 参数校验（404/409）
    with session() as s:
        t = Task(project_id=project_id, kind="generate", actor=au.current_actor())
        s.add(t)
        s.flush()
        tid = t.id
    mode = tasks.submit_generate(background, tid, project_id, profile, pkg_index, with_draft, product_model or "")
    return {"task_id": tid, "mode": mode}


@app.post("/api/projects/{project_id}/review")
def review_docs(project_id: str, profile: str, pkg_index: int = 0):
    """合规审查：规则引擎扫描最近一次生成结果 + 档案 + TRM。"""
    import docx as _docx
    from jb_docgen.placeholders import scan_docx
    from jb_docgen.techparams import fill_spec
    from jb_rules import Context, review
    trm, cp = _load_trm_profile(project_id, profile)
    pkg = trm.packages[pkg_index]
    out_dir = os.path.join(UPLOAD_DIR, project_id, "out")
    docx_todos, doc_texts = {}, {}
    if os.path.isdir(out_dir):
        for fn in os.listdir(out_dir):
            if fn.endswith(".docx") and pkg.pkg_no in fn:
                path = os.path.join(out_dir, fn)
                docx_todos[fn] = scan_docx(path)
                d = _docx.Document(path)
                doc_texts[fn] = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    from jb_docgen import pick_product
    product = pick_product(cp, pkg)
    tech_params = [fill_spec(sd, product) for sd in pkg.spec_docs]
    ctx = Context(trm=trm, pkg=pkg, profile=cp, tech_params=tech_params, docx_todos=docx_todos, doc_texts=doc_texts)
    rep = review(ctx, settings=config_api.rule_settings_models())
    with session() as s:
        p = _get_project(s, project_id)
        pm.set_result(p, "review", {"blocked": rep.blocked, "counts": rep.counts(), "profile": profile}, pkg.pkg_no)
        pm.advance(p, "reviewed")
        pm.log_event(s, project_id, "reviewed",
                     f"{pkg.pkg_no} 合规审查：{'存在否决项' if rep.blocked else '无否决项'}，" + "，".join(f"{k} {v}" for k, v in rep.counts().items()),
                     {"pkg_no": pkg.pkg_no, "blocked": rep.blocked, "counts": rep.counts()})
    return {"blocked": rep.blocked, "counts": rep.counts(), "findings": [f.model_dump() for f in rep.findings],
            "markdown": rep.markdown()}


@app.post("/api/projects/{project_id}/price/validate")
def price_validate(request: Request, project_id: str, sheet: dict, peer_avg: Optional[float] = None, over_limit_pct: Optional[float] = None):
    auth_api.require_price(request)
    from jb_agents.price import PriceSheet, validate
    issues = validate(PriceSheet.model_validate(sheet), peer_avg=peer_avg, over_limit_pct=over_limit_pct)
    return {"issues": [i.model_dump() for i in issues], "blocked": any(i.level == "否决" for i in issues)}


@app.post("/api/price/simulate")
def price_simulate(request: Request, my_price: float, peer_prices: list[float], c_candidates: Optional[list[float]] = None,
                   weight: float = 30.0):
    auth_api.require_price(request)
    from jb_agents.price import simulate_interval_avg
    sim = simulate_interval_avg(my_price, peer_prices, c_candidates or [0.0, 0.01, 0.02, 0.03, 0.05], weight=weight)
    return {"benchmark_range": sim.benchmark_range, "score_range": sim.score_range, "detail": sim.detail}


@app.get("/api/projects/{project_id}/submission-matrix")
def submission_matrix(project_id: str, profile: str = "", pkg_index: int = 0):
    """递交矩阵：提交方式表 × 本包，标注系统已产出的文件。"""
    with session() as s:
        p = s.get(Project, project_id)
        if not p or not (p.trm_confirmed or p.trm):
            raise HTTPException(404, "项目/TRM 不存在")
        trm = TRM.model_validate(p.trm_confirmed or p.trm)
    out_dir = os.path.join(UPLOAD_DIR, project_id, "out")
    produced = os.listdir(out_dir) if os.path.isdir(out_dir) else []
    rows = []
    for it in trm.submission_table:
        if not it.item:
            continue
        sec = it.section
        gen = None
        if "商务" in sec and any("商务文件" in f for f in produced):
            gen = next(f for f in produced if "商务文件" in f)
        elif "技术" in sec and any("技术文件" in f for f in produced):
            gen = next(f for f in produced if "技术文件" in f)
        rows.append({"section": sec, "seq": it.seq, "item": it.item, "channels": it.channels, "port": it.port,
                     "generated_file": gen, "status": "系统产出(含于文件)" if gen else ("工具内填报" if "价格" in sec else "人工挂载")})
    return {"rows": rows}


@app.get("/api/projects/{project_id}/files/{filename}")
def download_file(project_id: str, filename: str):
    path = os.path.join(UPLOAD_DIR, project_id, "out", os.path.basename(filename))
    if not os.path.exists(path):
        raise HTTPException(404, "文件不存在")
    return FileResponse(path, filename=filename,
                        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


# ---------- 模拟评分 / 导出 ----------

@app.post("/api/projects/{project_id}/score")
def score_project(project_id: str, profile: str, pkg_index: int = 0, llm: bool = False, task_id: Optional[str] = None):
    """模拟评分：硬指标按档案计数，软指标对最近一次起草（task_id 的 draft）做 LLM 评审。"""
    from jb_agents.scorer import score_package
    from jb_agents.writer import DraftSection
    trm, cp = _load_trm_profile(project_id, profile)
    pkg = trm.packages[pkg_index]
    drafts = None
    if task_id:
        with session() as s:
            t = s.get(Task, task_id)
            md = (t.result or {}).get("draft_markdown", "") if t else ""
        if md:
            import re as _re
            drafts = [DraftSection(title=m.group(1).strip(), text=m.group(2))
                      for m in _re.finditer(r"## (.+?)\n(.*?)(?=\n## |\Z)", md, _re.S)]
    from jb_store import config as cfg
    with session() as s:
        library = cfg.library(s)
    rep = score_package(trm, pkg, cp, draft_sections=drafts, use_llm=llm, library=library)
    with session() as s:
        p = _get_project(s, project_id)
        summary = {"weighted": rep.weighted, "tech_total": rep.tech_total, "tech_max": rep.tech_max,
                   "biz_total": rep.biz_total, "biz_max": rep.biz_max, "llm": llm}
        pm.set_result(p, "score", summary, pkg.pkg_no)
        pm.log_event(s, project_id, "scored", f"{pkg.pkg_no} 模拟评分：技术 {rep.tech_total}/{rep.tech_max}，商务 {rep.biz_total}/{rep.biz_max}"
                     + (f"，加权 {rep.weighted}" if rep.weighted is not None else ""), {"pkg_no": pkg.pkg_no, **summary})
    return {"report": rep.model_dump(), "heatmap": rep.heatmap(), "markdown": rep.markdown()}


@app.post("/api/projects/{project_id}/export")
def export_project(project_id: str, filename: str, pdf: bool = True, force: bool = False):
    """导出加固：待补充阻断 → 元数据清理 → PDF（有转换服务时）。"""
    from jb_docgen.export import export
    path = os.path.join(UPLOAD_DIR, project_id, "out", os.path.basename(filename))
    if not os.path.exists(path):
        raise HTTPException(404, "文件不存在")
    res = export(path, want_pdf=pdf, force=force)
    with session() as s:
        pm.log_event(s, project_id, "exported", f"导出 {os.path.basename(filename)}：{'成功' if res.ok else f'被阻断（待补充 {len(res.blocked_by)} 处）'}",
                     {"filename": os.path.basename(filename), "ok": res.ok, "blocked_count": len(res.blocked_by), "force": force})
    return {"ok": res.ok, "blocked_by": res.blocked_by[:20], "blocked_count": len(res.blocked_by),
            "pdf": os.path.basename(res.pdf_path) if res.pdf_path else None, "notes": res.notes}
