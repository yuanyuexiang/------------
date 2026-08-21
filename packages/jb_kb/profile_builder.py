"""从投标文件成品 docx 建档：结构化表格直读 + 标题树遍历。

数据源为供应商自己的历史投标文件（商务文件含基本情况表/财务表，
技术文件的标题树含业绩[合同+发票]与人员[证件七件套]）。
"""
from __future__ import annotations

import re
from typing import Optional

import docx

from .models import Boilerplate, Certificate, CompanyProfile, FinancialYear, Performance, Person

_CREDENTIAL_WORDS = ("身份证", "学历证", "职称证书", "资格证书", "社保证明", "劳动合同", "服务项目", "毕业证")
_CERT_HEADING_PAT = re.compile(r"(证书|许可证)$")
_NUM_PAT = re.compile(r"[-+]?\d+(?:\.\d+)?")


def _num(text: str) -> Optional[float]:
    m = _NUM_PAT.search(text.replace(",", ""))
    return float(m.group(0)) if m else None


def _dedup_row(row) -> list[str]:
    out: list[str] = []
    for c in row.cells:
        t = c.text.strip().replace("\n", " ")
        if not out or out[-1] != t:
            out.append(t)
    while out and not out[-1]:
        out.pop()
    return out


def _kv_pairs(cells: list[str]) -> dict:
    """['统一社会信用代码','9132...','企业资质等级','AAA级'] → dict。
    奇数列行首列多为合并组名（如"职工概况"），跳过首列再配对。"""
    if len(cells) % 2 == 1 and len(cells) > 1:
        cells = cells[1:]
    return {cells[i]: cells[i + 1] for i in range(0, len(cells) - 1, 2)}


def read_basic_info(doc, profile: CompanyProfile, source: str) -> None:
    for t in doc.tables:
        kv: dict = {}
        for r in t.rows:
            kv.update(_kv_pairs(_dedup_row(r)))
        # 以"实际含企业主档键"判定基本情况表（避免人员信息表等误命中）
        if not kv.get("单位名称") and not kv.get("统一社会信用代码"):
            continue
        profile.name = profile.name or kv.get("单位名称", "")
        profile.credit_code = profile.credit_code or kv.get("统一社会信用代码", "")
        profile.founded = profile.founded or kv.get("成立时间", "")
        profile.address = profile.address or kv.get("单位地址及邮编", "")
        profile.company_type = profile.company_type or kv.get("单位性质", "")
        profile.bank = profile.bank or kv.get("开户银行及账号", "")
        profile.contact = profile.contact or kv.get("联系人", "")
        profile.phone = profile.phone or kv.get("电话", "")
        profile.email = profile.email or kv.get("E-mail", "")
        profile.website = profile.website or kv.get("网址", "")
        profile.business_scope = profile.business_scope or kv.get("经营范围", "")
        cap = kv.get("注册资金（万元）") or kv.get("注册资本金")
        if cap and profile.registered_capital_wan is None:
            profile.registered_capital_wan = _num(cap)
        for key, attr in (("员工总数", "staff_total"), ("其中：技术人员数", "staff_technical"),
                          ("高级工程师", "senior_engineers"), ("工程师", "engineers")):
            if kv.get(key) and getattr(profile, attr) is None:
                v = _num(kv[key])
                setattr(profile, attr, int(v) if v is not None else None)
        if kv.get("企业资质等级"):
            profile.certificates.append(Certificate(
                name="企业资质等级 " + kv["企业资质等级"], cert_type="信用/资质等级", source=source))
        if source not in profile.sources:
            profile.sources.append(source)
        return


def read_financials(doc, profile: CompanyProfile) -> None:
    for t in doc.tables:
        head = _dedup_row(t.rows[0])
        years = [c for c in head if re.fullmatch(r"(19|20)\d{2}年", c)]
        if not years or "财务指标" not in " ".join(head):
            continue
        data = {y: FinancialYear(year=y) for y in years}
        for r in t.rows[1:]:
            cells = _dedup_row(r)
            if len(cells) < len(years) + 1:
                continue
            label = cells[1] if cells[0].isdigit() else cells[0]
            vals = cells[-len(years) - 1: -1] if len(cells) > len(years) + 1 else cells[1:]
            for y, v in zip(years, vals):
                fy = data[y]
                if label.startswith("营业收入"):
                    fy.revenue_wan = _num(v)
                elif label.startswith("净利润"):
                    fy.net_profit_wan = _num(v)
                elif label.startswith("资产总额"):
                    fy.asset_wan = _num(v)
                elif label.startswith("资产负债率"):
                    fy.liability_ratio = v
        existing = {f.year for f in profile.financials}
        profile.financials.extend(v for y, v in data.items() if y not in existing)
        return


def read_headings_tree(doc, profile: CompanyProfile, source: str) -> None:
    """标题树遍历：业绩（子标题为 合同/发票）、人员（子标题为证件七件套）、证书。"""
    heads = [p.text.strip() for p in doc.paragraphs
             if p.style is not None and "Heading" in (p.style.name or "") and p.text.strip()]
    known_perf = {p.project for p in profile.performances}
    known_person = {p.name for p in profile.personnel}
    i = 0
    while i < len(heads):
        h = heads[i]
        nxt = heads[i + 1: i + 9]
        # 业绩：本标题后紧跟 合同/发票 子标题
        if nxt and nxt[0] in ("合同", "发票") and len(h) > 8:
            ev = []
            j = i + 1
            while j < len(heads) and heads[j] in ("合同", "发票", "中标通知书"):
                ev.append(heads[j])
                j += 1
            if h not in known_perf:
                profile.performances.append(Performance(project=h, evidence=sorted(set(ev)), source=source))
                known_perf.add(h)
            i = j
            continue
        # 人员：2-4 字姓名标题，后跟证件类子标题
        if 2 <= len(h) <= 4 and nxt and nxt[0] in _CREDENTIAL_WORDS:
            creds = []
            j = i + 1
            while j < len(heads) and heads[j] in _CREDENTIAL_WORDS:
                creds.append(heads[j])
                j += 1
            if h not in known_person:
                profile.personnel.append(Person(name=h, credentials=creds, source=source))
                known_person.add(h)
            i = j
            continue
        # 证书标题（排除模板性标题）
        if _CERT_HEADING_PAT.search(h) and "格式" not in h and "要求" not in h and len(h) < 25:
            if not any(c.name == h for c in profile.certificates):
                cert_type = "体系认证" if "体系" in h else ("许可证" if "许可证" in h else "证书")
                profile.certificates.append(Certificate(name=h, cert_type=cert_type, source=source))
        i += 1
    if source not in profile.sources:
        profile.sources.append(source)


_BP_TOPICS = {
    "售后服务": ("售后服务", "服务承诺", "服务体系", "技术服务", "响应时间"),
    "质量保证": ("质量保证", "质量保障", "质量管理", "质量控制"),
    "培训": ("培训",),
    "项目管理": ("项目管理", "进度", "组织机构", "实施方案"),
    "保密": ("保密", "信息安全"),
}


def read_boilerplates(doc, profile: CompanyProfile, source: str, min_len: int = 40) -> None:
    """历史标书中售后/质量/培训等章节的正文段落 → 话术候选（approved=False，需人工审核后启用）。
    只收叙述性段落，不收含具体数字事实（金额/日期）的句子以免误用。"""
    topic = ""
    known = {b.text for b in profile.boilerplates}
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        is_heading = p.style is not None and "Heading" in (p.style.name or "")
        if is_heading:
            topic = next((k for k, kws in _BP_TOPICS.items() if any(w in t for w in kws)), "")
            continue
        if not topic or len(t) < min_len or t in known:
            continue
        if re.search(r"\d{4}年|\d+万元|\d+元", t):
            continue
        profile.boilerplates.append(Boilerplate(topic=topic, title="", text=t, source=source, approved=False))
        known.add(t)


def build_profile(doc_paths: list[str]) -> CompanyProfile:
    profile = CompanyProfile()
    for path in doc_paths:
        doc = docx.Document(path)
        read_basic_info(doc, profile, path)
        read_financials(doc, profile)
        read_headings_tree(doc, profile, path)
        read_boilerplates(doc, profile, path)
    return profile
