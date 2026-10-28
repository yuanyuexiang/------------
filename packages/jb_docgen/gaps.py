"""同一团队事实集中办理；保留每条原始问题与位置，不靠删除缺项放行。"""
from __future__ import annotations

from docx import Document

from .placeholders import MARK_PAT, todo


def is_team_gap(value):
    if any(word in value for word in ('人员关系', '授权', 'AI角色', '引用', '数值')):
        return False
    return any(word in value for word in ('拟任', '拟投入', '拟派', '拟入场', '本包团队',
               '项目负责人', '项目经理', '技术服务人员', '学历', '毕业日期', '工作年限',
               '职称证书', '到岗', '可投入', '人员清单'))


def consolidate_team_gaps(path):
    doc = Document(path)
    locations = []
    # 按 XML 文本节点替换，保留原段落样式、图片、超链接与表格结构。
    for index, paragraph in enumerate(doc.element.xpath('.//w:p')):
        nodes = paragraph.xpath('.//w:t')
        text = ''.join(n.text or '' for n in nodes)
        matches = [m for m in MARK_PAT.finditer(text) if is_team_gap(m.group(1) or '')]
        for match in reversed(matches):
            locations.append({'message': match.group(1), 'paragraph': index + 1})
            offset = 0
            placed = False
            for node in nodes:
                value = node.text or ''
                end = offset + len(value)
                lo, hi = max(0, match.start() - offset), min(len(value), match.end() - offset)
                if lo < hi:
                    node.text = value[:lo] + ('（见“本次项目团队与资格核对”）' if not placed else '') + value[hi:]
                    placed = True
                offset = end
    if locations:
        doc.add_heading('本次项目团队与资格核对', level=2)
        doc.add_paragraph(todo('本包团队名单、岗位与资格证据（集中核对）'))
        doc.add_paragraph('以下事项须集中核实后重新生成；各章节引用本表，不代表已满足人员资格。')
        for value in dict.fromkeys(x['message'] for x in reversed(locations)):
            doc.add_paragraph('需核实：' + value)
        doc.save(path)
    return list(reversed(locations))
