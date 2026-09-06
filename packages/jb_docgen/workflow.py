"""生成版本、审查证据与正式导出门禁。HTTP 层只负责调用和错误映射。

版本清单存于 Project.results.artifacts；旧产出没有清单时仅可作为草稿下载。
内容指纹包含完整 TRM / 聚合档案 / 附件元数据，避免跨包或沿用旧资料。
"""
from __future__ import annotations

import hashlib
import json
import os
import shutil
import tempfile
import uuid
from pathlib import Path
from typing import Optional

import docx
from jb_kb import repo
from jb_parser.trm import TRM
from jb_rules import Context, RuleSetting, review
from jb_store import Project, Task
from jb_store import config as cfg
from jb_store import projects as pm
from jb_store.models import KbAttachment
from sqlalchemy import select

from . import pick_product
from .export import export
from .placeholders import scan_docx
from .techparams import fill_spec


class WorkflowError(ValueError):
    pass


def fingerprint(value) -> str:
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                    separators=(",", ":"), default=str).encode()).hexdigest()


def file_hash(path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def inputs(s, p: Project, profile: str):
    cp = repo.load_profile(s, profile)
    if cp is None or not (p.trm_confirmed or p.trm):
        raise WorkflowError("企业档案或 TRM 不存在，请补齐资料后重新生成")
    trm = TRM.model_validate(p.trm_confirmed or p.trm)
    attachments = [{"id": a.id, "sha256": a.sha256} for a in s.scalars(
        select(KbAttachment).where(KbAttachment.company == profile).order_by(KbAttachment.id))]
    digest = fingerprint({"trm": trm.model_dump(), "profile": cp.model_dump(),
                          "attachments": attachments, "open_time": p.open_time,
                          "deadline": p.deadline})
    return trm, cp, digest


def artifacts(p: Project) -> dict:
    return (p.results or {}).get("artifacts", {})


def latest(p: Project, pkg_index: int) -> Optional[dict]:
    matches = [a for a in artifacts(p).values() if a["pkg_index"] == pkg_index]
    return max(matches, key=lambda a: a["created_at"], default=None)


def save_artifact(p: Project, a: dict) -> None:
    pm.set_result(p, "artifacts", {**artifacts(p), a["id"]: a})


def settings(s) -> dict:
    return {k: RuleSetting(rule_id=k, enabled=v["enabled"], level_override=v["level_override"],
                           params=v["params"], note=v["note"])
            for k, v in cfg.rule_settings(s).items()}


def rules_hash(s) -> str:
    return fingerprint({k: v.model_dump() for k, v in settings(s).items()})


def source_blockers(s, p: Project, a: dict, out_dir: str, profile: str = "") -> list[str]:
    reasons = []
    if not p.trm_confirmed:
        reasons.append("TRM 尚未人工确认，请先确认招标要求")
    active = s.scalar(select(Task.id).where(Task.project_id == p.id, Task.kind == "generate",
                                           Task.status.in_(["queued", "running"])).limit(1))
    if active:
        reasons.append("文件正在生成，请等待任务完成")
    if latest(p, a["pkg_index"])["id"] != a["id"]:
        reasons.append("此文件已有更新版本，请使用最新生成结果")
    if profile and profile != a["profile"]:
        reasons.append("当前企业与生成文件使用的企业不一致，请重新生成")
    try:
        if inputs(s, p, a["profile"])[2] != a["input_hash"]:
            reasons.append("TRM、企业档案或附件已变化，请重新生成后复查")
    except WorkflowError as exc:
        reasons.append(str(exc))
    for name, digest in a["files"].items():
        path = Path(out_dir) / name
        if not path.is_file() or file_hash(path) != digest:
            reasons.append(f"文件缺失或已被修改，请重新生成：{name}")
    return reasons


def blockers(s, p: Project, a: dict, out_dir: str, profile: str = "") -> list[str]:
    reasons = source_blockers(s, p, a, out_dir, profile)
    rep = a.get("review")
    if not rep:
        reasons.append("本生成版本尚未审查，请执行合规审查")
    else:
        if rep["rules_hash"] != rules_hash(s):
            reasons.append("否决规则配置已变化，请重新审查")
        if rep["blocked"]:
            reasons.append("合规审查存在否决项，请整改后复查")
    # 待补充门禁独立于可配置规则，不可通过停用 DOC-01 绕过。
    for name in a["files"]:
        path = Path(out_dir) / name
        if path.is_file():
            try:
                count = len(scan_docx(str(path)))
                if count:
                    reasons.append(f"{name} 仍有 {count} 处【待补充】，请补齐资料后重新生成")
            except Exception:
                reasons.append(f"文件无法读取，请重新生成：{name}")
    return reasons


def review_generation(s, p: Project, pkg_index: int, profile: str, out_dir: str) -> dict:
    a = latest(p, pkg_index)
    if not a:
        raise WorkflowError("本包没有可追溯的生成版本，请先生成文件")
    reasons = source_blockers(s, p, a, out_dir, profile)
    if reasons:
        raise WorkflowError("；".join(reasons))
    trm, cp, _ = inputs(s, p, a["profile"])
    pkg = trm.packages[pkg_index]
    texts, todos = {}, {}
    for name in a["files"]:
        path = str(Path(out_dir) / name)
        todos[name] = scan_docx(path)
        d = docx.Document(path)
        texts[name] = "\n".join([x.text for x in d.paragraphs] +
                                [c.text for t in d.tables for r in t.rows for c in r.cells])
    product = pick_product(cp, pkg, a.get("product_model") or None)
    rule_settings = settings(s)
    rep = review(Context(trm=trm, pkg=pkg, profile=cp,
                         tech_params=[fill_spec(sd, product) for sd in pkg.spec_docs],
                         docx_todos=todos, doc_texts=texts,
                         open_date=(p.open_time or trm.key_terms.bid_open_time or "")[:10] or None),
                 settings=rule_settings)
    result = {"generation_id": a["id"], "blocked": rep.blocked, "counts": rep.counts(),
              "findings": [f.model_dump() for f in rep.findings], "markdown": rep.markdown(),
              "rules_hash": fingerprint({k: v.model_dump() for k, v in rule_settings.items()})}
    save_artifact(p, {**a, "review": result})
    pm.set_result(p, "review", {**result, "profile": profile}, pkg.pkg_no)
    pm.advance(p, "reviewed")
    pm.log_event(s, p.id, "reviewed", f"{pkg.pkg_no} 合规审查：{'存在否决项' if rep.blocked else '无否决项'}",
                 {"pkg_no": pkg.pkg_no, "generation_id": a["id"], "blocked": rep.blocked})
    return result


def workflow_state(s, p: Project, pkg_index: int, profile: str, out_dir: str) -> dict:
    active = s.scalar(select(Task.id).where(Task.project_id == p.id, Task.kind == "generate",
                                           Task.status.in_(["queued", "running"])).limit(1))
    a = latest(p, pkg_index)
    if not a:
        return {"generation_id": None, "generation": None, "review": None,
                "blockers": ["请先生成文件，再审查和正式导出"], "exports": [], "active_task_id": active,
                "profile": None}
    task = s.get(Task, a["id"])
    reasons = blockers(s, p, a, out_dir, profile)
    return {"generation_id": a["id"], "generation": task.result if task else None,
            "review": a.get("review"), "blockers": reasons,
            "exports": [] if reasons else a.get("exports", []), "active_task_id": active,
            "profile": a["profile"]}


def formal_export(s, p: Project, filename: str, out_dir: str, want_pdf: bool, force: bool) -> dict:
    a = next((a for a in artifacts(p).values() if filename in a["files"]), None)
    reasons = blockers(s, p, a, out_dir) if a else ["旧文件无生成版本记录，请重新生成并审查"]
    if force:
        reasons.insert(0, "正式导出不允许强制绕过检查")
    result = {"ok": False, "blocked_by": reasons, "blocked_count": len(reasons),
              "docx": None, "pdf": None, "notes": []}
    if not reasons:
        # 在副本上清理元数据，不改变已审查草稿的指纹；正式文件采用独立名称。
        with tempfile.TemporaryDirectory(dir=Path(out_dir).parent) as tmp:
            name = f"正式_{uuid.uuid4().hex[:12]}_{filename}"
            copy = Path(tmp) / name
            shutil.copyfile(Path(out_dir) / filename, copy)
            res = export(str(copy), want_pdf=want_pdf)
            # PDF 转换可能耗时数分钟，发布前再次读取资料与规则，防止期间变更。
            s.expire_all()
            s.refresh(p)
            current = artifacts(p).get(a["id"])
            changed = blockers(s, p, current, out_dir) if current else ["生成版本记录已变化"]
            if changed:
                result.update(blocked_by=changed, blocked_count=len(changed))
                pm.log_event(s, p.id, "exported", "正式导出被阻断：转换期间资料或审查状态变化",
                             {"generation_id": a["id"], "ok": False, "blocked_by": changed})
                return result
            a = current
            if res.ok:
                exported = {"docx": name, "pdf": None, "hashes": {name: file_hash(copy)}}
                os.replace(copy, Path(out_dir) / name)
                if res.pdf_path:
                    pdf_name = Path(res.pdf_path).name
                    exported["pdf"] = pdf_name
                    exported["hashes"][pdf_name] = file_hash(res.pdf_path)
                    os.replace(res.pdf_path, Path(out_dir) / pdf_name)
                save_artifact(p, {**a, "exports": [*a.get("exports", []), exported]})
                result.update(ok=True, docx=name, pdf=exported["pdf"], notes=res.notes)
            else:
                result.update(blocked_by=res.blocked_by, blocked_count=len(res.blocked_by), notes=res.notes)
    pm.log_event(s, p.id, "exported", f"正式导出 {filename}：{'成功' if result['ok'] else '被阻断'}",
                 {"filename": filename, "generation_id": a["id"] if a else None,
                  "ok": result["ok"], "blocked_by": result["blocked_by"]})
    return result


def validate_download(s, p: Project, filename: str, out_dir: str) -> None:
    """草稿允许下载以人工核对；正式件下载必须仍满足当前门禁。"""
    for a in artifacts(p).values():
        for exported in a.get("exports", []):
            if filename in exported["hashes"]:
                reasons = blockers(s, p, a, out_dir)
                if file_hash(Path(out_dir) / filename) != exported["hashes"][filename]:
                    reasons.append("正式文件已被修改，请重新导出")
                if reasons:
                    raise WorkflowError("；".join(reasons))
                return
    if filename.startswith("正式_"):
        raise WorkflowError("正式文件缺少有效导出记录，请重新导出")
