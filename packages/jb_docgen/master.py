"""母版切片与原位填空：以招标文件自身的"投标文件格式"章为母版，保留国网原版式。

实测（物资：江西/陕西；服务：江苏/福建）：格式章（物资第六章"投标文件格式"，服务竞谈第五章
"应答文件格式"）按 价格文件 / 商务文件 / 技术文件 三段排布，分隔行有 "价格文件"、"1.价格文件"、
"一、价格文件"、"三、技术文件" 几种写法；段内每个格式块 = 编号标题 + 原文/表格，表格表头四省一致。
投标文件成品就是把对应段落原样复制后填空（见 docs/国网真实招标文件结构分析.md §6）。

本模块只做 docx 机械操作（切片、合并 run 后正则填空、表格模板行填充、按锚点插入），
不含任何业务字段映射——映射在 master_builder。
"""
from __future__ import annotations

import copy
import re
from collections.abc import Iterable
from dataclasses import dataclass
from typing import Callable, Optional, Union

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.table import Table, _Cell, _Row
from docx.text.paragraph import Paragraph
from jb_parser import docx_utils

FORMAT_CHAPTER_KW = "文件格式"
PART_PAT = re.compile(r"^(?:[一二三四五六]、|\d\s*[.、．])?\s*(价格|商务|技术)(?:文件|部分)$")
PART_KEY = {"价格": "price", "商务": "commercial", "技术": "technical"}
# ECP 投标工具操作指引行（"（一）商务偏差表（上传投标工具路径：…）"），成品投标文件里不保留
GUIDE_PAT = re.compile(r"^（[一二三四五六七八九十]+）.*(投标工具|应答工具|上传路径|填写路径|在线填写)")
_WS = re.compile(r"\s+")
_TEMPLATE_LEAD = re.compile(r"^(\d{1,3}|[…\.]+)$")
_DROPPABLE_RELS = {getattr(RT, n) for n in ("IMAGE", "HYPERLINK", "CHART", "OLE_OBJECT", "PACKAGE", "DIAGRAM_DATA",
                                            "DIAGRAM_LAYOUT", "DIAGRAM_QUICK_STYLE", "DIAGRAM_COLORS") if hasattr(RT, n)}
Sub = tuple["re.Pattern[str]", Union[str, Callable]]


def norm(text: str) -> str:
    return _WS.sub("", text or "")


@dataclass
class PartRange:
    key: str
    start: int   # 段落下标（含分隔行）
    end: int     # 段落下标（不含）


# ---------- 定位与切片 ----------

def format_chapter(doc):
    """格式章：标题含"文件格式"的章（物资第六章 / 服务第五章）；找不到返回 None。"""
    for c in docx_utils.split_chapters(doc):
        if FORMAT_CHAPTER_KW in c.title:
            return c
    return None


def split_parts(doc) -> dict[str, PartRange]:
    """格式章内按分隔行切成 price / commercial / technical 三段（缺段则无该键）。"""
    paras = doc.paragraphs
    ch = format_chapter(doc)
    lo, hi = (ch.start, ch.end) if ch else (0, len(paras))
    seps: list[tuple[int, str]] = []
    seen: set[str] = set()
    for i in range(lo, hi):
        m = PART_PAT.match(norm(paras[i].text))
        if m and PART_KEY[m.group(1)] not in seen:
            seen.add(PART_KEY[m.group(1)])
            seps.append((i, PART_KEY[m.group(1)]))
    parts: dict[str, PartRange] = {}
    for k, (start, key) in enumerate(seps):
        end = seps[k + 1][0] if k + 1 < len(seps) else hi
        parts[key] = PartRange(key=key, start=start, end=end)
    return parts


def carve(path: str, key: str):
    """重新打开母版，只保留指定段的正文元素（样式/编号/页面设置全部沿用），返回 Document；无该段返回 None。"""
    doc = docx.Document(path)
    pr = split_parts(doc).get(key)
    if pr is None:
        return None
    paras = doc.paragraphs
    start_el = paras[pr.start]._p
    end_el = paras[pr.end]._p if pr.end < len(paras) else None
    body = doc.element.body
    keep = False
    for el in list(body.iterchildren()):
        if el is start_el:
            keep = True
        elif el is end_el:
            keep = False
        if el.tag == qn("w:sectPr"):
            continue
        if not keep:
            body.remove(el)
    strip_orphan_rels(doc)
    return doc


def strip_orphan_rels(doc) -> int:
    """删掉正文已不再引用的图片/超链接/图表关系，避免母版其他章节的媒体留在投标文件里。"""
    refs = set(doc.element.xpath("//@r:id | //@r:embed | //@r:link | //@r:dm | //@r:lo | //@r:qs | //@r:cs"))
    dropped = 0
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype in _DROPPABLE_RELS and rid not in refs:
            del doc.part.rels[rid]
            dropped += 1
    return dropped


def remove_guidance(doc) -> int:
    n = 0
    for p in list(doc.paragraphs):
        if GUIDE_PAT.match(norm(p.text)):
            p._p.getparent().remove(p._p)
            n += 1
    return n


# ---------- 段落填空（跨 run） ----------

def iter_paragraphs(doc_or_cell) -> Iterable[Paragraph]:
    """正文段落 + 一层表格单元格段落。"""
    yield from doc_or_cell.paragraphs
    for t in doc_or_cell.tables:
        for row in t.rows:
            for c in distinct_cells(row):
                yield from c.paragraphs


def merge_runs(p: Paragraph):
    """把段落全部 run 文本并入第一个 run（沿用其字体），返回该 run；含图片/域/超链接的段落不动（返回 None）。"""
    runs = p.runs
    if not runs:
        return None
    if p._p.xpath(".//w:drawing | .//w:pict | .//w:fldChar | .//w:fldSimple | .//w:hyperlink | .//w:sdt"):
        return None
    text = "".join(r.text for r in runs)
    first = runs[0]
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    first.text = text
    return first


def fill_paragraph(p: Paragraph, subs: list[Sub]) -> bool:
    """按整段文本做正则替换；有变化才合并 run 并回写。"""
    text = p.text
    new = text
    for pat, rep in subs:
        new = pat.sub(rep, new)
    if new == text:
        return False
    run = merge_runs(p)
    if run is None:
        return False
    run.text = new
    return True


def fill_all(doc, subs: list[Sub]) -> int:
    return sum(1 for p in iter_paragraphs(doc) if fill_paragraph(p, subs))


def label_subs(labels: dict[str, str]) -> list[Sub]:
    """"标签：" 后为空（紧跟另一个"标签："或行尾）时补值；标签内允许排版空格（"包    号："）。

    所有标签合成一个正则单遍匹配——逐个替换会让已填的值挡住下一个标签的前置判断
    （"项目名称：X项目编号："）。"投标人：（盖章）" 这类已有内容的不动；"联系人及电话："不会被"电话"误中
    （标签前不能是汉字/字母/数字）。
    """
    items = {norm(k): v for k, v in labels.items() if v}
    if not items:
        return []
    alts = "|".join(r"\s*".join(re.escape(ch) for ch in k) for k in sorted(items, key=len, reverse=True))
    pat = re.compile(rf"(?<![一-龥A-Za-z0-9])(?P<label>{alts})\s*[：:]"
                     rf"(?=[ \t\u3000]*[，,、；;]?[ \t\u3000]*(?:[^：:，,、；;\s]{{1,14}}[：:]|$))")
    return [(pat, lambda m: m.group(0) + items[norm(m.group("label"))])]


def find_paragraph(doc, pattern: str, max_len: int = 60, after: Optional[Paragraph] = None) -> Optional[Paragraph]:
    """按正则找第一个（短）段落，常用于定位格式块标题行。"""
    pat = re.compile(pattern)
    started = after is None
    for p in doc.paragraphs:
        if not started:
            started = p._p is after._p
            continue
        t = norm(p.text)
        if t and len(t) <= max_len and pat.search(t):
            return p
    return None


# ---------- 表格 ----------

def distinct_cells(row: _Row) -> list[_Cell]:
    out: list[_Cell] = []
    seen: set[int] = set()
    for c in row.cells:
        if id(c._tc) not in seen:
            seen.add(id(c._tc))
            out.append(c)
    return out


def cells_by_grid(row: _Row) -> dict[int, _Cell]:
    """{起始网格列: 单元格}——表头有横向合并时靠网格列对齐数据列。"""
    out: dict[int, _Cell] = {}
    col = 0
    for c in distinct_cells(row):
        out[col] = c
        col += c._tc.grid_span
    return out


def find_table(doc, keywords: list[str], min_hits: int = 2, head_rows: int = 3) -> Optional[Table]:
    for t in doc.tables:
        if not t.rows:
            continue
        head = norm(" ".join(c.text for r in t.rows[:head_rows] for c in distinct_cells(r)))
        if sum(1 for k in keywords if k in head) >= min_hits:
            return t
    return None


def set_cell_text(cell: _Cell, text: str) -> None:
    """只改文本：保留首段的段落格式与首个 run 的字体。"""
    paras = cell.paragraphs
    p = paras[0]
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    runs = p.runs
    if runs:
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        runs[0].text = text
    else:
        p.add_run(text)


def set_kv_cells(table: Table, mapping: dict[str, str], row_scoped: Optional[dict[str, dict[str, str]]] = None) -> int:
    """"标签格 | 值格" 式表格（基本情况表）：标签命中则填右侧空格。row_scoped={行首关键词: {标签: 值}} 处理同名标签。"""
    filled = 0
    for row in table.rows:
        cells = distinct_cells(row)
        if not cells:
            continue
        scoped: dict[str, str] = {}
        head = norm(cells[0].text)
        for kw, m in (row_scoped or {}).items():
            if kw in head:
                scoped = m
        for i in range(len(cells) - 1):
            key = norm(cells[i].text).rstrip("：:")
            val = scoped.get(key) or mapping.get(key)
            if val and not cells[i + 1].text.strip():
                set_cell_text(cells[i + 1], val)
                filled += 1
    return filled


def _is_template_row(row: _Row) -> bool:
    texts = [c.text.strip() for c in distinct_cells(row)]
    if not texts:
        return False
    k = 0
    while k < len(texts) and texts[k]:
        k += 1
    if any(texts[k:]):
        return False
    return k == 0 or (k <= 2 and bool(_TEMPLATE_LEAD.match(norm(texts[0]))))


def data_rows(table: Table) -> list[int]:
    """表格里第一段连续的"模板空行"（全空，或只有序号/岗位两列预填），即待填数据区。"""
    idx: list[int] = []
    for i, row in enumerate(table.rows):
        if _is_template_row(row):
            idx.append(i)
        elif idx:
            break
    return idx


def fill_rows(table: Table, records: list[list[Optional[str]]]) -> int:
    """把 records 填进数据区：不够则克隆模板行追加，多余模板行删除（至少留一行）。None 表示保留模板格内容。"""
    idx = data_rows(table)
    if not idx:
        return 0
    trs = [table.rows[i]._tr for i in idx]
    template = copy.deepcopy(trs[0])
    n = max(len(records), 1)
    while len(trs) < n:
        new = copy.deepcopy(template)
        trs[-1].addnext(new)
        trs.append(new)
    while len(trs) > n:
        tr = trs.pop()
        tr.getparent().remove(tr)
    for tr, rec in zip(trs, records):
        cells = distinct_cells(_Row(tr, table))
        for c, v in zip(cells, rec):
            if v is not None:
                set_cell_text(c, v)
    return len(records)


def _header_map(table: Table, colmap: list[tuple[str, str]]) -> tuple[list[int], dict[int, str]]:
    """数据区行号 + {网格列: 字段}。表头正则按 colmap 顺序先到先得；同一字段只映射第一列（人员关系表左右两半同名）。"""
    idx = data_rows(table)
    if not idx or idx[0] == 0:
        return idx, {}
    header = table.rows[idx[0] - 1]
    pats = [(re.compile(k), f) for k, f in colmap]
    grid_field: dict[int, str] = {}
    used: set[str] = set()
    for col, cell in cells_by_grid(header).items():
        text = norm(cell.text)
        if not text:
            continue
        if re.fullmatch(r"序号", text):
            grid_field[col] = "__seq__"
            continue
        for pat, field in pats:
            if pat.search(text) and field not in used:
                grid_field[col] = field
                used.add(field)
                break
    return idx, grid_field


def header_fields(table: Table, colmap: list[tuple[str, str]]) -> set[str]:
    """表头实际映射到的字段集合（调用方据此调整记录，如没有"姓名"列就把姓名并入职务列）。"""
    return {f for f in _header_map(table, colmap)[1].values() if f != "__seq__"}


def fill_by_header(table: Table, records: list[dict], colmap: list[tuple[str, str]],
                   min_cols: int = 2) -> int:
    """按表头关键词把 dict 记录映射到列再填行；"序号"列自动编号。

    映射到的实质列不足 min_cols 时视为表头不匹配，不填（返回 0），由调用方另起新表。
    """
    idx, grid_field = _header_map(table, colmap)
    if sum(1 for f in grid_field.values() if f != "__seq__") < min_cols:
        return 0
    sample = cells_by_grid(table.rows[idx[0]])
    rows: list[list[Optional[str]]] = []
    for i, rec in enumerate(records):
        vals: list[Optional[str]] = []
        for col in sample:
            f = grid_field.get(col)
            if f == "__seq__":
                vals.append(str(i + 1))
            elif f is None or rec.get(f) is None:
                vals.append(None)
            else:
                vals.append(str(rec[f]))
        rows.append(vals)
    return fill_rows(table, rows)


def clone_table_after(table: Table, anchor_el) -> Table:
    new = copy.deepcopy(table._tbl)
    anchor_el.addnext(new)
    return Table(new, table._parent)


# ---------- 插入 ----------

def _style_or_none(doc, name: Optional[str]):
    if not name:
        return None
    try:
        return doc.styles[name]
    except KeyError:
        return None


def insert_paragraph_after(doc, anchor_el, text: str = "", style: Optional[str] = None,
                           center: bool = False, bold: bool = False) -> Paragraph:
    new_p = OxmlElement("w:p")
    anchor_el.addnext(new_p)
    para = Paragraph(new_p, doc._body)
    st = _style_or_none(doc, style)
    if st is not None:
        para.style = st
    if text:
        run = para.add_run(text)
        if bold:
            run.bold = True
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return para


def insert_heading_after(doc, anchor_el, text: str, level: int = 2) -> Paragraph:
    """母版有 Heading N 样式就用，没有就加粗段落（母版多数只有 Normal/自定义样式）。"""
    style = f"Heading {level}"
    if _style_or_none(doc, style) is not None:
        return insert_paragraph_after(doc, anchor_el, text, style=style)
    return insert_paragraph_after(doc, anchor_el, text, bold=True)


def _set_borders(tbl) -> None:
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def insert_table_after(doc, anchor_el, header: list[str], rows: list[list[str]],
                       like: Optional[Table] = None) -> Table:
    """新建表格并移到锚点之后；like 给定时沿用其表格属性，否则 Table Grid / 手工边框。"""
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    if like is not None and like._tbl.tblPr is not None:
        t._tbl.remove(t._tbl.tblPr)
        t._tbl.insert(0, copy.deepcopy(like._tbl.tblPr))
    elif _style_or_none(doc, "Table Grid") is not None:
        t.style = "Table Grid"
    else:
        _set_borders(t._tbl)
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
    for r, vals in enumerate(rows, start=1):
        for i, v in enumerate(vals):
            t.rows[r].cells[i].text = "" if v is None else str(v)
    anchor_el.addnext(t._tbl)
    return t


def insert_picture_after(doc, anchor_el, path: str, width_cm: float = 15.0) -> Optional[Paragraph]:
    para = insert_paragraph_after(doc, anchor_el, center=True)
    try:
        para.add_run().add_picture(path, width=Cm(width_cm))
    except Exception:
        para._p.getparent().remove(para._p)
        return None
    return para


def prepend_cover(doc, lines: list[tuple[str, bool]]) -> None:
    """在正文最前插入封面（文本, 是否居中），末尾分页。"""
    first = doc.paragraphs[0] if doc.paragraphs else None
    if first is None:
        for text, center in lines:
            p = doc.add_paragraph(text)
            if center:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        return
    for text, center in lines:
        p = first.insert_paragraph_before(text)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first.insert_paragraph_before().add_run().add_break(WD_BREAK.PAGE)


def has_cover(doc, look: int = 25) -> bool:
    for p in doc.paragraphs[:look]:
        t = norm(p.text)
        if t in ("投标文件", "应答文件", "技术/商务应答文件"):
            return True
        if "文件类别" in t:
            return True
    return False


def last_body_element(doc):
    """正文最后一个非 sectPr 元素（追加内容的锚点）。"""
    body = doc.element.body
    for el in reversed(list(body.iterchildren())):
        if el.tag != qn("w:sectPr"):
            return el
    return None
