"""投标文件生成器：按第六章格式目录装配 商务文件.docx / 技术文件.docx。

母版驱动：格式文本（投标函/授权书/承诺书条款）取自招标文件第六章原文（TRM.format_blocks），
变量由企业档案与 TRM 直填；事实缺失处写【待补充】。LLM 不参与本模块。
输出结构对齐真实投标文件成品（见 docs/国网真实招标文件结构分析.md §6）。
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from typing import Optional

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from jb_kb.models import CompanyProfile, Product
from jb_kb.selection import certificate_validity
from jb_parser.trm import TRM, PackageTRM

from .evidence import certificates_for, insert_credit, insert_scans
from .placeholders import todo
from .techparams import TechParamResult, deviation_rows, fill_spec
from .techparams import star_text as _star


def narrative_reference(ctx, sd) -> str:
    if ctx.spec_paths.get(sd.spec_id) or ctx.spec_paths.get(sd.source):
        return f"规范书 {sd.spec_id} 为叙述式规范，逐项响应见后附原文响应章节。"
    return todo(f"规范书 {sd.spec_id} 逐项响应（缺少规范书原件）")


def append_spec_responses(doc, anchor_el, ctx) -> object:
    """在锚点后追加"技术规范书逐条响应"章节（母版/自建版式共用）。返回新锚点。"""
    from . import master as M
    from .spec_response import append_spec
    paths = [(sd, ctx.spec_paths.get(sd.spec_id) or ctx.spec_paths.get(sd.source)) for sd in ctx.pkg.spec_docs]
    paths = [(sd, p) for sd, p in paths if p]
    if not paths:
        return anchor_el
    last = M.insert_heading_after(doc, anchor_el, "技术规范书逐条响应", level=1)._p
    for sd, path in paths:
        try:
            last, n, skipped = append_spec(
                doc, last, path, ctx.product, f"技术规范书 {sd.spec_id or ''} {sd.title}".strip(),
                plans=ctx.response_plans.get(sd.spec_id),
            )
        except Exception as exc:  # 单本规范书失败不拖垮整份文件
            ctx.notes.append(f"规范书 {sd.spec_id or sd.source} 逐条响应失败：{exc!r}")
            last = M.insert_paragraph_after(doc, last, todo(f"规范书 {sd.spec_id} 逐条响应（原文复制失败）"))._p
            continue
        if skipped:
            ctx.notes.append(f"规范书 {sd.spec_id or sd.source} 仅封面页，未纳入逐条响应")
        elif not n:
            last = M.insert_paragraph_after(doc, last, todo(f"规范书 {sd.spec_id} 未识别可自动响应条目，请人工逐条编制响应"))._p
    return last


@dataclass
class GenContext:
    trm: TRM
    pkg: PackageTRM
    profile: CompanyProfile
    product: Optional[Product] = None
    drafts: Optional[list] = None               # writer Agent 的 DraftSection 列表（可选）
    attachment_paths: dict = field(default_factory=dict)   # 知识库附件 id → 本地绝对路径（扫描件插图用）
    spec_paths: dict = field(default_factory=dict)         # 规范书 spec_id（或 source）→ 原 docx 路径（逐条响应用）
    response_plans: dict = field(default_factory=dict)
    relationship_confirmation: Optional[dict] = None
    todos: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)          # 生成过程说明（母版回退原因等）

    def need(self, value: Optional[str], what: str) -> str:
        if value:
            return str(value)
        if what not in self.todos:
            self.todos.append(what)
        return todo(what)


# ---------- 变量填充（第六章原文 → 填空） ----------

def fill_blanks(text: str, ctx: GenContext) -> str:
    t = text
    p, k, b = ctx.pkg, ctx.profile, ctx.trm
    subs = [
        (r"招标编号：\s*(?=，)", "招标编号：" + ctx.need(b.batch_no, "招标编号")),
        (r"项目名称：\s*(?=，)", "项目名称：" + ctx.need(p.project_name or b.batch_name, "项目名称")),
        (r"分标名称：\s*(?=，)", "分标名称：" + ctx.need(p.sub_name, "分标名称")),
        (r"包名称：\s*(?=招标文件)", "包名称：" + ctx.need(p.pkg_no, "包名称")),
        (r"_{3,}（投标人名称）", ctx.need(k.name, "投标人名称") + "（投标人名称）"),
        (r"_{3,}（应答人名称）", ctx.need(k.name, "应答人名称") + "（应答人名称）"),
        (r"法定地址：(?=。)", "法定地址：" + ctx.need(k.address, "法定地址")),
        (r"_{3,}（营业执照法定代表人（单位负责人））", ctx.need(k.legal_person, "法定代表人姓名") + "（营业执照法定代表人（单位负责人））"),
        (r"特授权_{3,}", "特授权" + ctx.need(k.authorized_rep, "被授权人姓名")),
        (r"XX（招标编号）", ctx.need(b.batch_no, "招标编号") + "（招标编号）"),
        (r"（分标编号及名称）", f"（{p.sub_no} {p.sub_name}）" if p.sub_no or p.sub_name else "（分标编号及名称）"),
        (r"（包号）（包名称）", f"（{p.pkg_no}）" if p.pkg_no else "（包号）（包名称）"),
        (r"职务：_{3,}", "职务：" + ctx.need(k.authorized_rep_title, "被授权人职务")),
    ]
    for pat, rep in subs:
        t = re.sub(pat, lambda _m, r=rep: r, t)
    return t


# ---------- docx 基础 ----------

def _doc():
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "宋体"
    st.font.size = Pt(12)
    return d


def _h(d, text: str, level: int = 1):
    return d.add_heading(text, level=level)


def _p(d, text: str, center: bool = False):
    para = d.add_paragraph(text)
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return para


def _table(d, header: list[str], rows: list[list[str]]):
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = "" if v is None else str(v)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for cell in t.rows[0].cells:
        for para in cell.paragraphs:
            para.paragraph_format.keep_with_next = True
    if header == ["序号", "姓名", "职务/职称", "专业", "证件材料"]:
        t.autofit = False
        for col, width in zip(t.columns, [0.8, 2.0, 2.5, 2.2, 7.5]):
            col.width = Cm(width)
        for row in t.rows:
            for cell, width in zip(row.cells, [0.8, 2.0, 2.5, 2.2, 7.5]):
                cell.width = Cm(width)
    return t


_SRC_TAG = re.compile(r"\s*\[来源:[^\]]+\]")


def iter_markdown(md: str, keep_sources: bool = False, section_title: str = ""):
    """极简 Markdown 流：产出 ("h", level, text) / ("li", 0, text) / ("p", 0, text)；
    **粗体**去标记，table 的 text 为二维单元格列表；来源标注默认剥离。母版/自建共用。"""
    def clean(raw):
        line = raw.rstrip()
        if not keep_sources:
            line = _SRC_TAG.sub("", line)
        return line.replace("**", "")

    def cells(line):
        return [cell.strip().replace(r"\|", "|")
                for cell in re.split(r"(?<!\\)\|", line.strip().strip("|"))]

    lines = [clean(raw) for raw in md.split("\n")]
    position = 0
    while position < len(lines):
        line = lines[position]
        if position + 1 < len(lines) and "|" in line:
            separator = cells(lines[position + 1])
            if separator and all(re.fullmatch(r":?-{3,}:?", cell) for cell in separator):
                rows = [cells(line)]
                position += 2
                while position < len(lines) and "|" in lines[position] and lines[position].strip():
                    rows.append(cells(lines[position]))
                    position += 1
                width = max(map(len, rows))
                rows = [row + [""] * (width - len(row)) for row in rows]
                yield "table", 0, rows
                continue
        position += 1
        if not line.strip() or re.fullmatch(r"[-_*]{3,}", line.strip()):
            continue
        m = re.match(r"^(#{1,6})\s*(.+)$", line)
        if m:
            if section_title and re.sub(r"\s+", "", m.group(2)) == re.sub(r"\s+", "", section_title):
                continue
            yield "h", min(3, 1 + len(m.group(1))), m.group(2).strip()
            continue
        m = re.match(r"^\s*(?:[-*]|\d+[.、])\s+(.+)$", line)
        if m:
            yield "li", 0, m.group(1).strip()
            continue
        yield "p", 0, line.strip()


def _markdown_to_doc(d, md: str, keep_sources: bool = False, section_title: str = ""):
    for kind, level, text in iter_markdown(md, keep_sources, section_title):
        if kind == "h":
            _h(d, text, level)
        elif kind == "table":
            _table(d, text[0], text[1:])
        elif kind == "li":
            d.add_paragraph(text, style="List Bullet")
        else:
            _p(d, text)


def _cover(d, ctx: GenContext, kind: str):
    _p(d, "投标文件" if ctx.trm.terminology == "投标" else "应答文件", center=True)
    _p(d, f"（{kind}）", center=True)
    _p(d, f"项目名称：{ctx.need(ctx.pkg.project_name or ctx.trm.batch_name, '项目名称')}")
    _p(d, f"招标编号：{ctx.need(ctx.trm.batch_no, '招标编号')}")
    _p(d, f"分标/包：{ctx.pkg.sub_no} {ctx.pkg.sub_name} {ctx.pkg.pkg_no}".strip())
    _p(d, f"投标人：{ctx.need(ctx.profile.name, '投标人名称')}")
    _p(d, "法定代表人（单位负责人）或其授权代表人：　　　　（签字）")
    d.add_page_break()


def _format_block(ctx: GenContext, title_kw: str) -> list[str]:
    """优先精确标题（如"投标函"），否则取第一个包含关键词且段落最多的块。"""
    exact = [b for b in ctx.trm.format_blocks if b.title == title_kw]
    if exact:
        return max(exact, key=lambda b: len(b.paragraphs)).paragraphs
    cands = [b for b in ctx.trm.format_blocks if title_kw in b.title]
    return max(cands, key=lambda b: len(b.paragraphs)).paragraphs if cands else []


# ---------- 商务文件 ----------

def build_commercial(ctx: GenContext, out_path: str) -> str:
    d = _doc()
    k = ctx.profile
    _cover(d, ctx, "商务文件")

    _h(d, "投标函", 1)
    paras = _format_block(ctx, "投标函") or ["致：招标人", todo("投标函正文（第六章格式未解析）")]
    for line in paras:
        _p(d, fill_blanks(line, ctx))
    _p(d, f"投标人：{k.name}（盖章）")
    _p(d, f"授权代表姓名：{ctx.need(k.authorized_rep, '授权代表姓名')}　职务：{ctx.need(k.authorized_rep_title, '授权代表职务')}")
    _p(d, "日期：　　年　　月　　日")

    _h(d, "商务偏差表", 1)
    _table(d, ["序号", "招标文件条目号", "招标文件条款", "投标文件条款", "偏差说明"],
           [["1", "/", "/", "/", "无偏差"]])

    _h(d, "投标人与国家电网公司系统人员关系说明", 1)
    from .relationships import append_plain
    append_plain(d, ctx)

    _h(d, "投标人基本情况表", 1)
    _table(d, ["项目", "内容"], [
        ["单位名称", k.name], ["统一社会信用代码", ctx.need(k.credit_code, "统一社会信用代码")],
        ["成立时间", ctx.need(k.founded, "成立时间")], ["注册资金（万元）", ctx.need(str(k.registered_capital_wan or ""), "注册资金")],
        ["单位性质", ctx.need(k.company_type, "单位性质")], ["单位地址", ctx.need(k.address, "单位地址")],
        ["开户银行及账号", ctx.need(k.bank, "开户银行及账号")], ["联系人/电话", f"{k.contact} {k.phone}".strip() or todo("联系人电话")],
        ["员工总数/技术人员", f"{k.staff_total or todo('员工总数')} / {k.staff_technical or todo('技术人员数')}"],
        ["经营范围", k.business_scope or todo("经营范围")],
    ])

    _h(d, "企业营业执照（扫描件）", 1)
    insert_scans(ctx, d, d.paragraphs[-1]._p, [a for c in certificates_for(k, "business") for a in c.attachments], "营业执照")
    _h(d, "有效的税务登记证明（扫描件）", 1)
    insert_scans(ctx, d, d.paragraphs[-1]._p, [a for c in certificates_for(k, "tax") for a in c.attachments], "税务登记证明")

    _h(d, "符合招标文件投标人资格要求的证明文件", 1)
    if k.certificates:
        _table(d, ["序号", "证书名称", "类别", "有效期"],
               [[str(i + 1), c.name, c.cert_type, certificate_validity(c) or todo(f"{c.name}有效期")] for i, c in enumerate(k.certificates)])
    else:
        _p(d, todo("资质证书清单"))
    _h(d, "查询报告及截图", 2)
    insert_credit(ctx, d, d.paragraphs[-1]._p)
    _h(d, "基本账户证明", 2)
    insert_scans(ctx, d, d.paragraphs[-1]._p, [a for c in certificates_for(k, "bank") for a in c.attachments], "银行基本账户证明")

    _h(d, "财务状况", 1)
    if k.financials:
        _table(d, ["年度", "营业收入（万元）", "净利润（万元）", "资产总额（万元）", "资产负债率"],
               [[f.year + ("（比较数据，非独立审计）" if f.report_kind == "comparative" else ""),
                 f.revenue_wan, f.net_profit_wan, f.asset_wan, f.liability_ratio] for f in k.financials])
        last = d.tables[-1]._tbl
        for f in k.financials:
            last = insert_scans(ctx, d, last, f.attachments, f"{f.year}审计报告")
    else:
        _p(d, todo("近三年财务审计报告"))

    _h(d, "法定代表人（单位负责人）授权委托书", 1)
    for line in _format_block(ctx, "授权委托书") or [todo("授权委托书正文")]:
        _p(d, fill_blanks(line, ctx))

    blk = _format_block(ctx, "信用承诺书")
    if blk:
        _h(d, "投标保证信用承诺书", 1)
        for line in blk:
            _p(d, fill_blanks(line, ctx))
        _p(d, f"投标人：{k.name}（盖章）　日期：　　年　　月　　日")

    _h(d, "公章对投标专用章授权书", 1)
    _p(d, f"我代表 {k.name}（投标人），在此作如下说明：在此次投标中，所用“投标专用章”与我公司公章具有同等的效力。特此说明。")

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    d.save(out_path)
    return out_path


# ---------- 技术文件 ----------

def build_technical(ctx: GenContext, out_path: str) -> tuple[str, list[TechParamResult]]:
    d = _doc()
    k, p = ctx.profile, ctx.pkg
    _cover(d, ctx, "技术文件")
    results: list[TechParamResult] = []

    # 1. 技术偏差表（汇总所有规范书的偏差）
    _h(d, "技术偏差表", 1)
    all_dev: list[dict] = []
    for sd in p.spec_docs:
        r = fill_spec(sd, ctx.product)
        results.append(r)
        all_dev += deviation_rows(r)
    if all_dev:
        _table(d, ["序号", "对应条款", "招标文件要求", "投标响应", "偏差说明"],
               [[x["序号"], x["对应条款"], x["招标文件要求"], x["投标响应"], x["偏差说明"]] for x in all_dev])
    else:
        _table(d, ["序号", "对应条款", "招标文件要求", "投标响应", "偏差说明"], [["1", "/", "/", "/", "无偏差"]])

    # 2. 业绩文件
    _h(d, "业绩文件", 1)
    _h(d, "供货及运行业绩表", 2)
    if k.performances:
        _table(d, ["序号", "项目名称", "买方", "签约时间", "金额（万元）", "证明材料"],
               [[str(i + 1), x.project, x.buyer or todo("买方"), x.signed_date or todo("签约日期"),
                 x.amount_wan if x.amount_wan is not None else todo("合同金额"), "/".join(x.evidence)]
                for i, x in enumerate(k.performances)])
        for x in k.performances:
            _h(d, x.project, 3)
            if x.attachments:
                insert_scans(ctx, d, d.paragraphs[-1]._p, x.attachments, x.project)
            else:
                for ev in x.evidence or ["合同", "发票"]:
                    _p(d, f"{ev}：{todo(f'{x.project} 的{ev}扫描件')}")
    else:
        _p(d, todo("业绩清单及合同/发票扫描件"))

    # 3. 技术特性参数表（逐本规范书）
    _h(d, "技术特性参数表", 1)
    for sd, r in zip(p.spec_docs, results):
        _h(d, f"技术规范书 {sd.spec_id or sd.title}", 2)
        if not sd.structured:
            _p(d, "（非结构化技术规范：以下为逐项响应，另以附件形式上传）")
        if sd.param_rows:
            _table(d, ["序号", "参数名称", "单位", "项目需求值或表述", "投标人保证值"],
                   [[row.row, row.name, row.unit, _star(row), resp.response]
                    for row, resp in zip(sd.param_rows, r.responses)])
        else:
            _p(d, narrative_reference(ctx, sd))

    # 3b. 技术规范书逐条响应（规范书原文复制 + 每节响应）
    append_spec_responses(d, d.paragraphs[-1]._p, ctx)

    # 4. 技术/服务方案（写作 Agent 起草；未起草时给结构占位）
    if ctx.drafts:
        _h(d, "技术/服务方案", 1)
        if p.scope:
            _p(d, "招标范围：" + p.scope)
        for sec in ctx.drafts:
            _h(d, sec.title, 2)
            _markdown_to_doc(d, sec.text, section_title=sec.title)
    elif p.scope:
        _h(d, "项目需求响应", 1)
        _p(d, "招标范围：" + p.scope)
        _p(d, todo("总体技术/服务方案（由写作 Agent 依据评分项起草）"))

    # 5. 人员
    _h(d, "主要技术人员、项目经理及其相关证书", 1)
    if k.personnel:
        _table(d, ["序号", "姓名", "职务/职称", "专业", "证件材料"],
               [[str(i + 1), x.name, x.title or todo(f"{x.name}职称"), x.major, "、".join(x.credentials)] for i, x in enumerate(k.personnel)])
        for x in k.personnel:
            _h(d, x.name, 3)
            if x.attachments:
                insert_scans(ctx, d, d.paragraphs[-1]._p, x.attachments, x.name)
            else:
                for c in (x.credentials or ["身份证", "学历证", "职称证书", "社保证明", "劳动合同"]):
                    _p(d, f"{c}：{todo(f'{x.name} 的{c}扫描件')}")
    else:
        _p(d, todo("拟投入人员清单及七件套材料"))

    # 6. 售后 / 质量 / 认证（话术库或待补充）
    for topic, heading in (("售后服务", "售后服务保障及承诺"), ("质量保证", "质量保证措施方案"), ("培训", "技术服务，包括培训、安装指导等")):
        _h(d, heading, 1)
        bps = [b for b in k.boilerplates if b.topic == topic and b.approved]
        if bps:
            for b in bps:
                _p(d, b.text)
        else:
            _p(d, todo(f"{heading}（话术库暂无审核通过的段落）"))
    _h(d, "认证证书", 1)
    certs = [c for c in k.certificates if c.cert_type == "体系认证"]
    _table(d, ["序号", "证书", "有效期"], [[str(i + 1), c.name, certificate_validity(c) or todo(f"{c.name}有效期")] for i, c in enumerate(certs)]) if certs else _p(d, todo("体系认证证书"))
    if certs:
        last = d.tables[-1]._tbl
        for c in certs:
            last = insert_scans(ctx, d, last, c.attachments, c.name)

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    d.save(out_path)
    return out_path, results
