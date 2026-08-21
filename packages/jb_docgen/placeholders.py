"""【待补充】标记：事实缺失时的唯一合法输出；导出前必须清零（export 阶段阻断）。"""
from __future__ import annotations

import re

MARK = "【待补充：{}】"
MARK_PAT = re.compile(r"【待补充：([^】]*)】")


def todo(what: str) -> str:
    return MARK.format(what)


def find_all(text: str) -> list[str]:
    return MARK_PAT.findall(text)


def scan_docx(path: str) -> list[str]:
    """扫描 docx 全文（段落+表格）中的待补充标记。"""
    import docx
    d = docx.Document(path)
    found: list[str] = []
    for p in d.paragraphs:
        found += find_all(p.text)
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                found += find_all(c.text)
    return found
