"""缺项处理闭环：企业长期事实、标包确认与旧产出提示分离。"""
from __future__ import annotations

import json
import re
import tempfile
from datetime import date, datetime, timezone
from functools import lru_cache
from pathlib import Path

from docx import Document
from jb_kb import attachments, repo
from jb_kb.models import CompanyProfile
from jb_kb.selection import approved_profile
from jb_parser.trm import TRM, PackageTRM
from jb_store import Task
from jb_store import projects as pm
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import select

STATS = ('staff_total', 'staff_technical', 'senior_engineers', 'engineers',
         'assistant_engineers', 'technicians', 'staff_as_of', 'staff_source')
REGISTRATION = ('主管部门', '批准部门', '批准文号')
TITLES = {'authorization': '本次授权', 'relationship': '国网人员关系确认',
          'team': '本次项目团队', 'statistics': '企业人员统计',
          'registration': '企业基本情况核对', 'performance': '历史业绩合同日期',
          'system': '系统引用与生成修复',
          'other': '其他原文或生成问题'}


class Authorization(BaseModel):
    model_config = ConfigDict(extra='forbid')
    person_id: str
    title: str = Field(min_length=1, max_length=200)
    identity_attachments: list[str] = Field(default_factory=list)


class TeamMember(BaseModel):
    model_config = ConfigDict(extra='forbid')
    person_id: str
    role: str = Field(min_length=1, max_length=200)
    available: bool  # 必须明确本包可投入状态，不从全局状态推断


class Team(BaseModel):
    model_config = ConfigDict(extra='forbid')
    members: list[TeamMember] = Field(default_factory=list, max_length=300)
    arrival_commitment: bool = False


def _key(index, company):
    return json.dumps([index, company], ensure_ascii=False)


def group_for(message: str) -> str:
    from .gaps import is_team_gap
    if any(w in message for w in ('数值需核实来源', '引用来源需核实', 'AI角色', 'AI 起草失败', '模型未配置')):
        return 'system'
    if any(w in message for w in ('人员关系', '关系说明')):
        return 'relationship'
    if '被授权人' in message:
        return 'authorization'
    if any(w in message for w in ('员工总数', '技术人员数', '工程师人数', '统计日期')):
        return 'statistics'
    if any(w in message for w in ('股东', '出资', '主管部门', '批准部门', '批准文号', '基本情况表')):
        return 'registration'
    if any(w in message for w in ('签约日期', '历史业绩', '合同签订日期补充')):
        return 'performance'
    if is_team_gap(message):
        return 'team'
    return 'other'


def _basis(p, index, cp, kind, data):
    from .workflow import fingerprint
    ids = ([data.get('person_id')] if kind == 'authorization'
           else [x['person_id'] for x in data.get('members', [])])
    records = {x.id: x.model_dump() for x in cp.personnel if x.id in ids}
    return fingerprint({'trm': p.trm_confirmed or p.trm, 'index': index, 'company': cp.name,
                        'deadline': p.deadline, 'people': records, 'ids': ids})


def fact_state(p, index, cp):
    saved = (p.results or {}).get('bid_facts', {}).get(_key(index, cp.name), {})
    return {kind: {'confirmation': saved.get(kind),
                   'stale': bool(saved.get(kind) and (
                       (kind == 'team' and not saved[kind]['data'].get('members')) or
                       saved[kind]['basis'] != _basis(p, index, cp, kind, saved[kind]['data'])))}
            for kind in ('authorization', 'team')}


def apply_facts(p, index, cp):
    """只覆盖生成用副本，绝不将项目岗位和承诺写回企业人员主档。"""
    full = cp
    states = fact_state(p, index, full)
    result = cp.model_copy(deep=True)
    context = {'personnel_candidates': [x.model_dump(exclude={'id_no'}) for x in full.personnel
                                       if x.status == 'approved']}
    for kind, state in states.items():
        saved = state['confirmation']
        if not saved or state['stale']:
            if saved and kind == 'authorization':
                result.authorized_rep = result.authorized_rep_title = ''
            continue
        data = saved['data']
        people = {x.id: x for x in full.personnel if x.status == 'approved'}
        if kind == 'authorization':
            person = people.get(data['person_id'])
            if person:
                result.authorized_rep, result.authorized_rep_title = person.name, data['title']
                context[kind] = {'name': person.name, 'title': data['title'],
                                 'attachments': data.get('identity_attachments', [])}
        else:
            members = []
            for member in data['members']:
                person = people.get(member['person_id'])
                if person:
                    members.append({**member, 'name': person.name})
                    next(x for x in result.personnel if x.id == person.id).available = member['available']
            context[kind] = {**data, 'members': members}
    result.bid_context = context
    return result


def ensure_idle(s, p):
    if s.scalar(select(Task.id).where(Task.project_id == p.id,
                Task.kind.in_(['parse', 'generate']), Task.status.in_(['queued', 'running']))):
        raise ValueError('解析或生成正在进行，请任务完成后再保存')


def save_facts(s, p, index, cp, kind, data, actor):
    if kind not in ('authorization', 'team'):
        raise ValueError('不支持的本次投标任务')
    ensure_idle(s, p)
    validated = (Authorization if kind == 'authorization' else Team).model_validate(data).model_dump()
    people = {x.id: x for x in cp.personnel if x.status == 'approved'}
    ids = ([validated['person_id']] if kind == 'authorization'
           else [x['person_id'] for x in validated['members']])
    if len(ids) != len(set(ids)) or any(x not in people for x in ids):
        raise ValueError('人员必须属于本企业且已审核，不允许重复选择')
    if kind == 'authorization':
        person = people[validated['person_id']]
        owned = {a.id for a in attachments.list_attachments(s, cp.name)}
        if any(key not in person.attachments or key not in owned for key in validated['identity_attachments']):
            raise ValueError('身份证明附件必须属于所选人员及本企业')
    if kind == 'team':
        if not validated['members']:
            raise ValueError('请先选择团队成员；空名单不能保存为已确认团队，也不能确认到岗承诺')
        if validated['arrival_commitment'] and any(not x['available'] for x in validated['members']):
            raise ValueError('仍有不可投入人员，不能确认全体到岗承诺')
        from .workflow import material_selection, save_selection
        selection = material_selection(p, index, cp.name).model_dump()
        selection['personnel'] = ids
        save_selection(p, index, cp, selection)
    all_facts = dict((p.results or {}).get('bid_facts', {}))
    key = _key(index, cp.name)
    all_facts[key] = {**all_facts.get(key, {}), kind: {
        'data': validated, 'basis': _basis(p, index, cp, kind, validated),
        'confirmed_by': actor, 'confirmed_at': datetime.now(timezone.utc).isoformat()}}
    pm.set_result(p, 'bid_facts', all_facts)


def validate_statistics(data):
    for key in STATS[:6]:
        value = data.get(key)
        if value is not None and (isinstance(value, bool) or not isinstance(value, int) or value < 0):
            raise ValueError('人数必须为非负整数，未知请留空')
    if any(data.get(key) is not None for key in STATS[:6]):
        try:
            date.fromisoformat(data.get('staff_as_of', ''))
        except ValueError:
            raise ValueError('填写人数时必须提供有效统计日期 YYYY-MM-DD') from None
        if not data.get('staff_source', '').strip():
            raise ValueError('请说明统计依据或口径')
    if (data.get('staff_total') is not None and data.get('staff_technical') is not None
            and data['staff_technical'] > data['staff_total']):
        raise ValueError('技术人员数不能大于员工总数')


def save_company_task(s, cp, kind, data):
    if kind == 'statistics':
        if set(data) - set(STATS):
            raise ValueError('包含非人员统计字段')
        combined = {k: data.get(k, getattr(cp, k)) for k in STATS}
        validate_statistics(combined)
        return repo.save_main(s, cp.name, combined)
    if kind == 'registration':
        if set(data) - {'registration', 'shareholders'}:
            raise ValueError('包含非登记股权字段')
        merged = CompanyProfile.model_validate({**cp.model_dump(), **data})
        if set(merged.registration) - set(REGISTRATION):
            raise ValueError('不支持的登记字段')
        for value in merged.registration.values():
            if value.not_applicable and (not value.reason.strip() or value.value.strip()):
                raise ValueError('不适用须填写理由，且不能同时填写适用值')
            if value.value.strip() and not value.source.strip():
                raise ValueError('登记信息须注明资料出处')
        if sum(x.ratio or 0 for x in merged.shareholders) > 100.001:
            raise ValueError('出资比例合计不能超过100%')
        for owner in merged.shareholders:
            if not owner.name.strip() or not owner.source.strip():
                raise ValueError('股东须填写名称与资料出处')
        return repo.save_main(s, cp.name, data)
    if kind == 'performance':
        item = next((x for x in cp.performances if x.id == data.get('id')), None)
        if item is None:
            raise ValueError('历史业绩不存在或不属于本企业')
        value = data.get('signed_date', '')
        if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
            raise ValueError('请填写合同签署页证明的日期 YYYY-MM-DD；约略日期不能冒充准确签署日')
        date.fromisoformat(value)
        if not data.get('source', '').strip():
            raise ValueError('请注明合同名称、页码或其他日期证明依据')
        return repo.upsert_item(s, cp.name, 'performances', {
            **item.model_dump(), 'signed_date': value,
            'source': item.source + '\n签约日期核对：' + data['source']})
    raise ValueError('不支持的企业资料任务')


@lru_cache(maxsize=16)
def _requirements(path, mtime, size, package_json):
    from jb_agents.draft_context import requirement_sources
    from jb_parser import package_files
    package = PackageTRM.model_validate_json(package_json)
    with tempfile.TemporaryDirectory(prefix='jb-preparation-') as folder:
        files = package_files(path, package, workdir=folder)
        sources, _ = requirement_sources(TRM(), package, files.spec_docs)
        roles = []
        for spec_path in dict.fromkeys(files.spec_docs.values()):
            for table in Document(spec_path).tables:
                if not table.rows:
                    continue
                heads = [c.text.strip() for c in table.rows[0].cells]
                if '岗位名称' not in heads or '数量' not in heads:
                    continue
                for row in table.rows[1:]:
                    values = [c.text.strip() for c in row.cells]
                    role, number = values[heads.index('岗位名称')], values[heads.index('数量')]
                    if number.isdigit() and role:
                        roles.append({'role': role, 'count': int(number),
                                      'requirement': '；'.join(dict.fromkeys(values)),
                                      'source': Path(spec_path).name})
        return sources, roles


def requirements(p, index):
    trm = TRM.model_validate(p.trm_confirmed or p.trm)
    try:
        st = Path(p.zip_path).stat()
        return _requirements(p.zip_path, st.st_mtime_ns, st.st_size,
                             trm.packages[index].model_dump_json())
    except (OSError, ValueError, KeyError):
        return {}, []


def state(s, p, index, cp):
    from .workflow import inputs, latest, material_selection, relationship_state
    revision = inputs(s, p, cp.name)[2]
    artifact = latest(p, index)
    task = s.get(Task, artifact['id']) if artifact and artifact['profile'] == cp.name else None
    result = task.result or {} if task else {}
    groups = {key: {'id': key, 'title': title, 'items': [], 'occurrences': 0}
              for key, title in TITLES.items()}
    for filename, messages in result.get('todos', {}).items():
        for message in messages:
            group = groups[group_for(message)]
            group['occurrences'] += 1
            if not any(x['message'] == message and x['file'] == filename for x in group['items']):
                group['items'].append({'message': message, 'file': filename})
            if group['id'] == 'registration' and any(w in message for w in ('助理工程师', '技术员')):
                groups['statistics']['items'].append({'message': '企业助理工程师和技术员人数', 'file': filename})
    sources, roles = requirements(p, index)
    for detail in result.get('gap_details', {}).get('team', []):
        groups['team']['items'].append({'message': detail['message'],
            'file': f"{result.get('summary', {}).get('technical', '技术文件')} · 原段落{detail['paragraph']}"})
    facts = fact_state(p, index, cp)
    approved = approved_profile(cp)
    selection = material_selection(p, index, cp.name)
    perfs = [x for x in approved.performances if not x.signed_date
             and (selection.performances is None or x.id in selection.performances)]
    relationship = relationship_state(p, index, cp.name)
    for key in ('authorization', 'team'):
        current = facts[key]
        valid = current['confirmation'] and not current['stale']
        groups[key]['complete'] = bool(valid)
        groups[key]['status'] = '已确认' if valid else '确认已失效' if current['stale'] else '待确认本包安排'
        groups[key]['current_notes'] = []
    team = facts['team']
    if team['confirmation'] and not team['confirmation']['data'].get('members'):
        groups['team']['current_notes'].append('原记录为空团队，已停止作为有效确认；请从已有人员中选择或让AI推荐')
    team_issues = team_blockers(p, index, cp)
    groups['team']['current_notes'] += team_issues
    if team_issues:
        groups['team']['complete'] = False
        groups['team']['status'] = '待完善团队安排'
    groups['team']['current_notes'].append(f'知识库有{len(approved.personnel)}名已审核候选；不代表已确认本包投入')
    groups['relationship']['status'] = ('已确认' if relationship['confirmation'] and not relationship['stale'] else '待本次确认')
    groups['relationship']['complete'] = bool(relationship['confirmation'] and not relationship['stale'])
    missing_stats = [key for key in STATS if getattr(cp, key) in (None, '')]
    groups['statistics']['status'] = '已录入，待核对口径' if not missing_stats else '统计字段待核对'
    missing_registration = [label for label in REGISTRATION if not cp.registration.get(label) or not (
        cp.registration[label].value or cp.registration[label].not_applicable and cp.registration[label].reason)]
    groups['registration']['status'] = '登记与股权待核对' if missing_registration or not cp.shareholders else '已录入，待原件核对'
    groups['registration']['current_notes'] = missing_registration + ([] if cp.shareholders else ['股东出资明细尚未录入'])
    groups['performance']['status'] = f'{len(perfs)}项历史合同日期待核对' if perfs else '所选业绩日期已录入'
    groups['performance']['current_notes'] = [x.project for x in perfs]
    groups['statistics']['complete'] = not missing_stats
    groups['registration']['complete'] = not missing_registration and bool(cp.shareholders)
    groups['performance']['complete'] = not perfs
    for key in ('statistics', 'registration', 'performance'):
        if groups[key]['complete']:
            groups[key]['status'] = '资料已齐备，可复用'
    groups['other']['status'] = '系统生成核查，不等同于缺少企业资料'
    groups['system']['status'] = '系统需重试或核查，不要求补交企业资料'
    from .readiness import readiness
    trm = TRM.model_validate(p.trm_confirmed or p.trm)
    registration_requirements = [
        {'title': block.title, 'text': '\n'.join(block.paragraphs)}
        for block in trm.format_blocks
        if any(word in block.title + '\n'.join(block.paragraphs)
               for word in (*REGISTRATION, '股东', '出资比例'))]
    active_suggestion = s.scalar(select(Task.id).where(Task.project_id == p.id,
                                Task.kind == 'suggest', Task.status.in_(['queued', 'running'])))
    return {'active_suggestion': active_suggestion,
            'registration_requirements': registration_requirements,
            'readiness': readiness(p, index, cp), 'revision': revision, 'profile': cp.model_dump(), 'facts': facts,
            'tasks': [v for k, v in groups.items() if k not in ('other', 'system') or v['items']],
            'generation_stale': bool(artifact and artifact['input_hash'] != revision),
            'generation_id': artifact['id'] if task else None,
            'attachments': [a.model_dump() for a in attachments.list_attachments(s, cp.name)],
            'missing_dates': [x.model_dump() for x in perfs],
            'roles': roles,
            'team_requirements': [{'id': k, 'text': v} for k, v in sources.items()
                                  if k.startswith('req-spec') and any(w in v for w in
                                      ('人员', '岗位', '学历', '社保', '工作经验'))],
            'relationship': relationship}


def save(s, p, index, cp, kind, payload, actor):
    from .workflow import inputs
    ensure_idle(s, p)
    if payload.get('revision') != inputs(s, p, cp.name)[2]:
        raise ValueError('资料已被更新，请刷新任务后重新核对，避免覆盖他人的修改')
    data = payload.get('data', {})
    if kind in ('authorization', 'team'):
        save_facts(s, p, index, cp, kind, data, actor)
    else:
        save_company_task(s, cp, kind, data)
    pm.log_event(s, p.id, 'preparation', f'更新{TITLES.get(kind, kind)}，需重新生成并复查',
                 {'company': cp.name, 'pkg_index': index, 'kind': kind, 'actor': actor})


def team_blockers(p, index, cp):
    saved = fact_state(p, index, cp)['team']
    data = (saved['confirmation']['data'] if saved['confirmation'] and not saved['stale']
            else {'members': [], 'arrival_commitment': False})
    _, roles = requirements(p, index)
    issues = []
    for role in roles:
        count = sum(x['role'] == role['role'] and x['available'] for x in data['members'])
        if count < role['count']:
            issues.append(f"本包{role['role']}要求{role['count']}人，已确认可投入{count}人")
    if any(not x['available'] for x in data['members']):
        issues.append('项目团队中仍有暂不可投入人员，请调整本次名单')
    if (roles or data['members']) and not data['arrival_commitment']:
        issues.append('本包到岗承诺尚未确认')
    return issues


def append_team_document(path, profile):
    """将本包岗位与承诺确定性写入技术文件，避免模型漏写或把全局可用状态当成本包承诺。"""
    team = profile.bid_context.get('team')
    if not team:
        return
    from .placeholders import todo
    doc = Document(path)
    doc.add_heading('本次项目团队与到岗安排', level=2)
    doc.add_paragraph('以下安排经本次投标负责人确认；人员资格及工作经历以所附证明为准。')
    table = doc.add_table(rows=1, cols=4)
    for cell, value in zip(table.rows[0].cells, ('姓名', '本次岗位', '可投入状态', '资格与经历依据')):
        cell.text = value
    people = {x.id: x for x in profile.personnel}
    for member in team['members']:
        person = people.get(member['person_id'])
        if person is None:
            continue
        cells = table.add_row().cells
        cells[0].text, cells[1].text = person.name, member['role']
        cells[2].text = '本包已确认可投入' if member['available'] else todo(person.name + '本包可投入状态')
        cells[3].text = '；'.join(x for x in [person.education, person.major, person.title,
                                             person.experience_source, person.source] if x)
    if team['arrival_commitment'] and all(x['available'] for x in team['members']) and team['members']:
        doc.add_heading('本次人员到岗承诺', level=3)
        doc.add_paragraph(f'{profile.name}承诺按本次确认的岗位安排组织上述人员到岗，'
                          '具体到岗时间、采购人审核和人员变更按采购文件及合同约定执行。')
        doc.add_paragraph('投标人（盖章）：\n授权代表（签字）：\n日期：')
    else:
        doc.add_paragraph(todo('拟投入本包人员名单及到岗承诺'))
    doc.save(path)


def append_authorization_document(path, ctx):
    authority = ctx.profile.bid_context.get('authorization')
    if not authority:
        return
    from .master_builder import _insert_scans
    from .placeholders import todo
    doc = Document(path)
    heading = doc.add_heading('本次授权人身份证明附件', level=2)
    if authority['attachments']:
        _insert_scans(ctx, doc, heading._p, authority['attachments'], authority['name'] + '身份证明')
    else:
        doc.add_paragraph(todo('被授权人' + authority['name'] + '身份证明附件'))
    doc.save(path)
