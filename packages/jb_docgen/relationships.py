"""人员关系声明：企业确认事实，程序选择招标原条款，不从档案空值推断。"""
from __future__ import annotations

import re
from copy import deepcopy
from typing import Literal

from docx.table import Table
from pydantic import BaseModel, ConfigDict, Field, field_validator

TITLE = re.compile(r"(?:投标人|应答人|响应供应商)与国家电网公司系统人员关系说明")
OPTION = re.compile(r"【\s*([12])[.．、]?\s*(不存在|存在)相关情形的适用\s*】")
LABELS = ["本企业人员姓名", "性别", "身份证号", "职务", "任职时间", "与国网系统人员关系",
          "国网人员姓名", "性别", "身份证号", "（曾）任职单位名称", "职务", "任职状态", "离职/退休时间"]
HEADERS = ["人员姓名", "性别", "身份证号", "职务", "任职时间", "与国网公司系统人员关系",
           "人员姓名", "性别", "身份证号", "（曾）任职单位名称", "职务", "任职状态", "离职/退休时间"]


class RelationshipConfirmation(BaseModel):
    model_config = ConfigDict(extra="forbid")
    status: Literal["none", "exists"]
    people: list[list[str]] = Field(default_factory=list, max_length=100)

    @field_validator("people")
    @classmethod
    def validate_people(cls, value):
        if any(len(row) != len(LABELS) or any(len(cell) > 300 for cell in row) for row in value):
            raise ValueError("人员关系每行须包含 13 个字段，每项不超过 300 字")
        return [[cell.strip() for cell in row] for row in value]


def source_options(trm) -> list[str]:
    return list(dict.fromkeys(p for block in trm.format_blocks if TITLE.search(block.title)
                             for p in block.paragraphs if OPTION.search(p)))


def _cell_text(cell, text):
    # 保留母版单元格段落属性及首个 run 的字体，避免替换 text 重置版式。
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(text)
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def _fill_table(table, ctx, state):
    # 只对识别完整的标准双人员表操作，未知格式交给复核，不猜列义。
    header = next((i for i, row in enumerate(table.rows)
                   if len(row.cells) == 13 and all(
                       expected in re.sub(r"\s", "", c.text)
                       for expected, c in zip(HEADERS, row.cells))), None)
    if header is None or len(table.rows) <= header + 1:
        return False
    data_rows = table.rows[header + 1:]
    if state["status"] == "none":
        for row in data_rows:
            for c in row.cells:
                _cell_text(c, "")
        _cell_text(data_rows[0].cells[0], "不适用")
        return True
    people = state.get("people") or [[""] * 13]
    template = deepcopy(data_rows[0]._tr)
    for row in data_rows:
        table._tbl.remove(row._tr)
    for values in people:
        table._tbl.append(deepcopy(template))
        row = table.rows[-1]
        for i, (cell, value) in enumerate(zip(row.cells, values)):
            if i == 12 and values[11] == "在职" and not value:
                value = "不适用"
            _cell_text(cell, value or ctx.need(None, f"关系说明：{LABELS[i]}"))
    return True


def apply_master(doc, ctx):
    """只处理具有两条明确互斥原文的关系小节，保留表格、注释及签章格式。"""
    paragraphs = doc.paragraphs
    candidates = [(i, p) for i, p in enumerate(paragraphs) if TITLE.search(p.text)
                  and len(p.text.strip()) < 65]
    pairs = {}
    for index, heading in candidates:
        options = []
        for p in paragraphs[index + 1:index + 12]:
            match = OPTION.search(p.text)
            if match:
                options.append((p, match.group(1)))
            elif options and p.text.strip() and not p.text.strip().startswith("《"):
                break
        if sorted(n for _, n in options) == ["1", "2"]:
            pairs[tuple(p._p for p, _ in options)] = (heading, options)
    from . import master as M
    if len(pairs) != 1:
        anchor = candidates[-1][1]._p if candidates else doc.paragraphs[-1]._p
        M.insert_paragraph_after(doc, anchor, ctx.need(None, "人员关系说明原文格式需核对（无需另交承诺书模板）"))
        return
    heading, options = next(iter(pairs.values()))
    state = ctx.relationship_confirmation
    if not state:
        M.insert_paragraph_after(doc, heading._p, ctx.need(None, "本包国网系统人员关系情况待企业确认"))
        return
    selected = "1" if state["status"] == "none" else "2"
    # 寻找紧随条款的关系表，不跨到下一节。
    table = None
    node = options[-1][0]._p.getnext()
    for _ in range(8):
        if node is None:
            break
        if node.tag.endswith("}tbl"):
            table = Table(node, doc._body)
            break
        text = "".join(node.itertext())
        if re.match(r"\s*\d+[.．]", text):
            break
        node = node.getnext()
    for p, number in options:
        if number != selected:
            p._p.getparent().remove(p._p)
    if table is None or not _fill_table(table, ctx, state):
        if state["status"] == "exists" or table is not None:
            M.insert_paragraph_after(doc, heading._p, ctx.need(None, "人员关系表原文列格式及人员信息需核对"))


def append_plain(doc, ctx):
    options = source_options(ctx.trm)
    state = ctx.relationship_confirmation
    if not state:
        doc.add_paragraph(ctx.need(None, "本包国网系统人员关系情况待企业确认"))
        return
    selected = "1" if state["status"] == "none" else "2"
    matches = [p for p in options if OPTION.search(p).group(1) == selected]
    if len(matches) != 1:
        doc.add_paragraph(ctx.need(None, "人员关系说明原文格式需核对（无需另交承诺书模板）"))
        return
    doc.add_paragraph(matches[0])
    # 无母版不能声称已保留原表及签章格式，保留明确复核门禁。
    doc.add_paragraph(ctx.need(None, "人员关系说明原文排版需核对（未取得可用母版）"))
    if selected == "2":
        table = doc.add_table(rows=1, cols=13)
        for cell, label in zip(table.rows[0].cells, LABELS):
            cell.text = label
        for values in state.get("people") or [[""] * 13]:
            for cell, value, label in zip(table.add_row().cells, values, LABELS):
                cell.text = value or ctx.need(None, f"关系说明：{label}")
