"""技术规范书逐条响应：规范书原文（段落+表格，保留版式）整体复制进技术文件，每个叶子小节标题后插【投标人响应】。

客户实际做法（纪要 2026-09-10）：规范书全文复制到技术文件，每个标题下插响应，多本规范书逐一纳入；
真实成品也证实"技术文件把技术规范书的投标人响应部分整体嵌入并填写"（结构分析 §6）。
与客户口头做法的差别：这里**不写"我司承诺完全响应"套话**（案例库 11(2)、SG-16 禁词），响应只有两种来源——
产品库事实（参数数值比较复用 techparams，功能条目按产品特性关键词匹配）或【待补充】。

响应形式：
- 规范书里的参数表：有"保证值"列就填该列，没有就追加一列"投标人响应"；
- "(1)…(2)…"功能条目：标题后逐条给出匹配到的产品特性，匹配不到的条目写【待补充】；
- 纯叙述小节（概述/工程概况）不插响应；只有封面的 docx 整本跳过。
"""
from __future__ import annotations

import copy
import re
from dataclasses import dataclass, field
from typing import Optional

import docx
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.table import Table
from docx.text.paragraph import Paragraph
from jb_kb.models import Product
from jb_parser.trm import SpecDoc, SpecParamRow

from . import master as M
from .placeholders import todo
from .techparams import _keywords, fill_spec, match_param_text

HEAD_PAT = re.compile(r"^(\d+(?:\.\d+){0,3})[\s.、．)）]*(\S.{0,38})$")
ITEM_PAT = re.compile(r"^[（(]?\d{1,2}[)）.、．]\s*\S")
RESP_TAG = "【投标人响应】"
_STRIP_XPATH = ".//w:drawing | .//w:pict | .//w:object | .//w:numPr | .//w:sectPr"
_TABLE_HEADER_KW = ("要求", "参数", "指标", "名称", "项目", "内容", "保证值", "响应")


@dataclass
class Section:
    heading_el: object
    title: str
    level: int
    items: list[str] = field(default_factory=list)
    tables: list = field(default_factory=list)      # 源表格元素
    has_children: bool = False
    prose: int = 0


def _heading_level(p) -> Optional[int]:
    t = p.text.strip()
    if not t or len(t) > 45:
        return None
    name = p.style.name if p.style is not None else ""
    m = re.match(r"(?:Heading|标题)\s*(\d)", name)
    if m:
        return int(m.group(1))
    m = HEAD_PAT.match(t)
    if m and not t.endswith(("；", ";", "，", ",", "。")):
        return m.group(1).count(".") + 1
    return None


def analyse(src) -> list[Section]:
    """按正文顺序切小节；返回的 heading_el 用于定位插入点。"""
    sections: list[Section] = []
    cur: Optional[Section] = None
    body = src.element.body
    for el in body.iterchildren():
        if el.tag == qn("w:p"):
            p = Paragraph(el, src)
            lvl = _heading_level(p)
            if lvl is not None:
                for s in reversed(sections):
                    if s.level < lvl:
                        s.has_children = True
                        break
                    if s.level == lvl:
                        break
                cur = Section(heading_el=el, title=p.text.strip(), level=lvl)
                sections.append(cur)
            elif cur is not None and p.text.strip():
                if ITEM_PAT.match(p.text.strip()):
                    cur.items.append(p.text.strip())
                else:
                    cur.prose += 1
        elif el.tag == qn("w:tbl") and cur is not None:
            cur.tables.append(el)
    return sections


def is_cover_only(src, sections: list[Section]) -> bool:
    """只有封面信息（"国家电网公司集中规模招标采购 / 工程概况 / 网省公司："）的 docx：没有可响应内容。"""
    return not src.tables and not any(s.items or s.prose for s in sections)


# ---------- 响应生成 ----------

def _respond_items(items: list[str], product: Optional[Product]) -> list[str]:
    out = []
    for i, text in enumerate(items, start=1):
        body = re.sub(r"^[（(]?\d{1,2}[)）.、．]\s*", "", text)
        hit_text = ""
        if product is not None:
            kw = _keywords(body)
            best, best_score = "", 0
            for feat in product.features:
                score = len(kw & _keywords(feat))
                if score > best_score:
                    best, best_score = feat, score
            # 二元组重叠须既够多又占条目本身的相当比例，否则"支持/接口/配置"这类通用词会把无关特性拉进来
            if best_score >= 3 and best_score >= 0.35 * max(len(kw), 1):
                hit_text = best
            else:
                hit = match_param_text(body, product)
                if hit is not None and hit[0] in body:
                    hit_text = f"{hit[0]} {hit[1]}"
        out.append(f"（{i}）{hit_text}" if hit_text else f"（{i}）{todo(f'对「{body[:30]}」的具体响应')}")
    return out


def _table_rows(table: Table) -> list[list[str]]:
    return [[c.text.strip() for c in M.distinct_cells(r)] for r in table.rows]


def _looks_header(cells: list[str]) -> bool:
    joined = "".join(cells)
    return not re.search(r"\d", joined) and any(k in joined for k in _TABLE_HEADER_KW)


def respond_table(table: Table, product: Optional[Product]) -> int:
    """参数表：有"保证值/响应"列就填该列，否则追加"投标人响应"列。表头可多行、可横向合并（按网格列对齐）。返回填写行数。"""
    rows = _table_rows(table)
    if not rows or len(rows[0]) < 2:
        return 0
    n_header = 0
    while n_header < len(rows) - 1 and _looks_header(rows[n_header]):
        n_header += 1
    resp_col = req_col = None
    for hi in range(n_header):
        for col, cell in M.cells_by_grid(table.rows[hi]).items():
            hn = M.norm(cell.text)
            if resp_col is None and ("保证值" in hn or "响应" in hn):
                resp_col = col
            if req_col is None and any(k in hn for k in ("要求", "需求", "表述", "指标")):
                req_col = col
    data_idx = [i for i in range(n_header, len(rows)) if any(rows[i])]
    if not data_idx:
        return 0
    spec = SpecDoc(spec_id="inline")
    for n, i in enumerate(data_idx, start=1):
        grid = M.cells_by_grid(table.rows[i])
        texts = {c: cell.text.strip() for c, cell in grid.items()}
        if req_col is not None and req_col in texts:
            required = texts[req_col]
            name = " ".join(v for c, v in sorted(texts.items()) if c not in (req_col, resp_col) and v)[:60]
        else:
            vals = [v for _, v in sorted(texts.items())]
            required = vals[-1] if vals else ""
            name = " ".join(vals[:-1])[:60]
        spec.param_rows.append(SpecParamRow(row=n, name=name or required, required=required, star="★" in required[:4]))
    responses = {}
    for r in fill_spec(spec, product).responses:
        # 通用匹配器按关键词重叠挑参数，"控制接口"会被"接口类型"沾上；要求命中的参数名与本行至少共享 2 个二元组
        row = spec.param_rows[r.row - 1]
        weak = r.matched_param not in ("", "features") and "," not in r.matched_param \
            and r.matched_param not in row.name \
            and len(_keywords(r.matched_param) & _keywords(row.name + row.required)) < 2
        if weak:
            responses[r.row] = todo(f"「{row.name[:20]}」的具体参数值")
        else:
            responses[r.row] = r.response
    filled = 0
    if resp_col is None:
        table.add_column(Cm(3.5))
        if n_header:
            M.set_cell_text(M.distinct_cells(table.rows[0])[-1], "投标人响应")
        for n, i in enumerate(data_idx, start=1):
            M.set_cell_text(M.distinct_cells(table.rows[i])[-1], responses.get(n, ""))
            filled += 1
    else:
        for n, i in enumerate(data_idx, start=1):
            cell = M.cells_by_grid(table.rows[i]).get(resp_col)
            if cell is not None and not cell.text.strip():
                M.set_cell_text(cell, responses.get(n, ""))
                filled += 1
    return filled


# ---------- 复制 + 插入 ----------

def _clean_copy(el):
    new = copy.deepcopy(el)
    for bad in new.xpath(_STRIP_XPATH):
        bad.getparent().remove(bad)
    return new


def append_spec(doc, anchor_el, spec_path: str, product: Optional[Product], title: str) -> tuple[object, int, bool]:
    """把一本规范书复制到 anchor 之后并逐节响应。返回 (新锚点, 响应小节数, 是否跳过)。"""
    src = docx.Document(spec_path)
    sections = analyse(src)
    if is_cover_only(src, sections):
        return anchor_el, 0, True
    last = M.insert_heading_after(doc, anchor_el, title, level=2)._p
    by_el = {id(s.heading_el): s for s in sections}
    responded = 0
    for el in src.element.body.iterchildren():
        if el.tag not in (qn("w:p"), qn("w:tbl")):
            continue
        if el.tag == qn("w:p") and not el.xpath("string(.)").strip() and id(el) not in by_el:
            continue
        new = _clean_copy(el)
        last.addnext(new)
        last = new
        if el.tag == qn("w:tbl"):
            t = Table(new, doc._body)
            if respond_table(t, product):
                responded += 1
            continue
        sec = by_el.get(id(el))
        if sec is None or sec.has_children or not sec.items:
            continue
        last = M.insert_paragraph_after(doc, last, RESP_TAG, bold=True)._p
        for line in _respond_items(sec.items, product):
            last = M.insert_paragraph_after(doc, last, line)._p
        responded += 1
    return last, responded, False
