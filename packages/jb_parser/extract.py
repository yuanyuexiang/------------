"""从主文件/规范书/清单中抽取 TRM 各部分。"""
from __future__ import annotations

import os
import re
from typing import Optional

import openpyxl

from . import docx_utils
from .trm import (
    Material,
    PrenoticeClause,
    QualificationItem,
    RejectionRule,
    SpecDoc,
    SpecParamRow,
    SubmissionItem,
)

SPEC_ID_PAT = re.compile(r"\(([A-Z0-9]{4}-\d{9}-\d{5})\)")
BATCH_NO_PAT = re.compile(r"[（(]([A-Z0-9\-]{6,})(?:物资|服务)?[）)]")


# ---------- 主文件：前附表 / 否决表 / 提交方式表 ----------

def extract_prenotice(doc) -> list[PrenoticeClause]:
    rows = docx_utils.find_table(doc, ["条款号", "条款名称", "编列内容"], min_hits=3)
    out: list[PrenoticeClause] = []
    if not rows:
        return out
    import re as _re
    _clause_no = _re.compile(r"^\d+(\.\d+)*")
    for r in rows[1:]:
        if len(r) < 2:
            continue
        cells = list(r)
        # 四列变体（如福建竞谈：大项|条款号|条款名称|编列内容）：首列非条款号则左移
        if len(cells) >= 4 and not _clause_no.match(cells[0]) and _clause_no.match(cells[1]):
            cells = cells[1:]
        clause = PrenoticeClause(
            clause_no=cells[0], name=cells[1] if len(cells) > 1 else "",
            content=cells[2] if len(cells) > 2 else (cells[1] if len(cells) > 1 else ""))
        if clause.clause_no or clause.content:
            out.append(clause)
    return out


def extract_rejection_rules(doc) -> list[RejectionRule]:
    rows = docx_utils.find_table(doc, ["评审内容", "否决情形", "否决事项"], min_hits=2)
    out: list[RejectionRule] = []
    if not rows:
        return out
    last_cat = ""
    for r in rows[1:]:
        cells = [c for c in r if c]
        if not cells:
            continue
        # 形如 [类别, 方面, 内容] 或（合并后）[方面, 内容]
        if len(cells) >= 3:
            cat, aspect, text = cells[0], cells[1], cells[-1]
            last_cat = cat
        elif len(cells) == 2:
            cat, aspect, text = last_cat, cells[0], cells[1]
        else:
            cat, aspect, text = last_cat, "", cells[0]
        if len(text) < 8:  # 过滤表头/噪声行
            continue
        out.append(RejectionRule(category=cat, aspect=aspect, text=text))
    return out


_SECTION_NAMES = ("价格文件", "商务文件", "技术文件", "资质业绩文件", "报价文件")


def extract_submission_table(doc) -> list[SubmissionItem]:
    """提交方式表，兼容两种实测变体：
    - 陕西物资版：… | 提交方式(电子商务平台/省系统) | 格式要求 | 投标工具上传端口，
      分节行为"一/二/三 | 价格文件（按包制作）"
    - 江苏服务版：双行表头，渠道列为 ECP2.0（国网）/网盘/天源e采（江苏），
      分节行为"1 | 价格文件"，明细行以 √/× 标注各渠道
    """
    rows = docx_utils.find_table(doc, ["提交方式", "投标工具上传端口"], min_hits=2)         or docx_utils.find_table(doc, ["提交方式", "是否有格式要求"], min_hits=2)
    out: list[SubmissionItem] = []
    if not rows:
        return out
    # 渠道名从前两行表头收集
    channel_names = []
    for r in rows[:2]:
        for c in r:
            cc = c.replace(chr(10), "")
            if any(k in cc for k in ("平台", "系统", "ECP", "网盘", "e采")) and len(cc) < 20:
                if cc not in channel_names:
                    channel_names.append(cc)
    section = ""
    for r in rows[1:]:
        cells = [c for c in r if c]
        if not cells:
            continue
        seq = cells[0].replace(chr(10), "")
        item = cells[1].replace(chr(10), " ") if len(cells) > 1 else ""
        if seq in ("序号",) or item in ("内容",):
            continue
        is_section = (seq in ("一", "二", "三", "四") or seq.isdigit()) and             any(item.startswith(n) for n in _SECTION_NAMES)
        if is_section:
            section = item
            continue
        marks = [c for c in cells[2:]]
        channels = []
        if channel_names and any(m in ("√", "×") for m in marks):
            for name, m in zip(channel_names, marks):
                if m == "√":
                    channels.append(name)
        port = cells[-1] if len(cells) > 2 and ("端口" in cells[-1] or "页签" in cells[-1] or "生成" in cells[-1]) else ""
        out.append(SubmissionItem(seq=seq, item=item, channels=channels,
                                  port=port, section=section))
    return out


def batch_name_from_doc(doc, limit: int = 60) -> str:
    name = ""
    for p in doc.paragraphs[:limit]:
        t = p.text.strip()
        if len(t) > len(name) and ("招标采购" in t or "采购项目" in t or "谈判采购" in t) and len(t) < 80:
            name = t
    return name


def extract_batch_info(doc, filename: str) -> tuple[str, str, str]:
    """返回 (批次名, 批次号, 术语体系)。批次名主文件取不到时由调用方从招标公告兜底。"""
    name = batch_name_from_doc(doc)
    m = BATCH_NO_PAT.search(os.path.basename(filename))
    no = m.group(1) if m else ""
    text_head = "\n".join(p.text for p in doc.paragraphs[:200])
    term = "应答" if ("应答人" in text_head and "投标人" not in text_head[:2000]) else "投标"
    return name, no, term


# ---------- 技术规范书：标准技术参数表 ----------

def extract_spec(path: str, relpath: str) -> Optional[SpecDoc]:
    try:
        doc = docx_utils.load(path)
    except Exception:
        return None
    m = SPEC_ID_PAT.search(relpath)
    spec_id = m.group(1) if m else ""
    rows = docx_utils.find_table(doc, ["参数名称", "投标人保证值"], min_hits=2) \
        or docx_utils.find_table(doc, ["项目需求值", "保证值"], min_hits=1)
    title = ""
    for p in doc.paragraphs[:20]:
        t = p.text.strip()
        if t and len(t) < 60:
            title = t
            break
    sd = SpecDoc(spec_id=spec_id, title=title, source=relpath,
                 structured=spec_id.startswith("9999"))
    if rows:
        for i, r in enumerate(rows[1:], start=1):
            cells = [c for c in r if c]
            if not cells:
                continue
            # 典型五列：序号|参数名称|单位|项目需求值|投标人保证值（保证值列为空会被去重掉）
            name = cells[1] if len(cells) > 1 else cells[0]
            unit = cells[2] if len(cells) > 3 else ""
            required = cells[3] if len(cells) > 3 else (cells[-1] if len(cells) > 1 else "")
            sd.param_rows.append(SpecParamRow(
                row=i, name=name, unit=unit, required=required,
                star=required.startswith("★") or "★" in required[:4]))
    return sd


# ---------- xlsx：货物清单 / 资质业绩一览表 ----------

def _detect_header(ws, keywords: list[str], scan: int = 6) -> Optional[int]:
    for idx, row in enumerate(ws.iter_rows(max_row=scan, values_only=True)):
        vals = [str(c) for c in row if c is not None]
        if sum(1 for k in keywords if any(k in v for v in vals)) >= 2:
            return idx
    return None


def _load_ws(path: str):
    """ECP 生成的 xlsx 常见 dimension 元数据缺失（只写 A1），read_only 下必须
    reset_dimensions 强制全表扫描，否则 iter_rows 只返回 1 行。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    try:
        ws.reset_dimensions()
    except AttributeError:
        pass
    return wb, ws


def extract_goods_list(path: str) -> list[Material]:
    wb, ws = _load_ws(path)
    head_idx = _detect_header(ws, ["分标编号", "物资描述", "包名称", "数量"])
    out: list[Material] = []
    if head_idx is None:
        return out
    header: list[str] = []
    for idx, row in enumerate(ws.iter_rows(values_only=True)):
        vals = ["" if c is None else str(c).strip() for c in row]
        if idx == head_idx:
            header = vals
            continue
        if idx < head_idx or not any(vals):
            continue
        rec = dict(zip(header, vals))
        out.append(Material(
            sub_no=rec.get("分标编号", ""), pkg=rec.get("包名称", rec.get("分包", "")),
            project=rec.get("项目名称", ""), desc=rec.get("物资描述", rec.get("物料描述", "")),
            unit=rec.get("单位", ""), qty=rec.get("数量", ""),
            deliver_date=rec.get("最后一批交货日期", rec.get("交货时间", "")),
            deliver_place=rec.get("交货地点", ""),
            spec_id=rec.get("技术规范书ID", rec.get("技术规范书Id", ""))))
    wb.close()
    return out


def extract_qualification(path: str) -> list[QualificationItem]:
    wb, ws = _load_ws(path)
    head_idx = _detect_header(ws, ["分标名称", "业绩要求", "资质要求"])
    out: list[QualificationItem] = []
    if head_idx is None:
        return out
    rows = list(ws.iter_rows(values_only=True))
    for row in rows[head_idx + 2:]:  # 双层表头：主表头 + 资质要求子列
        vals = ["" if c is None else str(c).strip() for c in row]
        if len(vals) < 6 or not any(vals) or not vals[1]:
            continue
        out.append(QualificationItem(
            sub_name=vals[1], pkg=vals[3] if len(vals) > 3 else "",
            performance_req=vals[4] if len(vals) > 4 else "",
            test_report_req=vals[5] if len(vals) > 5 else "",
            accept_agent=vals[12] if len(vals) > 12 else ""))
    wb.close()
    return out


# ---------- xlsx：服务类本批需求一览表（每包一行：评审办法 + 资格要求） ----------

def _f(v: str) -> Optional[float]:
    try:
        return float(str(v).replace(",", "").replace("%", "")) if v not in ("", "/") else None
    except ValueError:
        return None


def extract_service_demand(path: str, pkg_no: str) -> Optional[dict]:
    """返回本包的需求行（表头→值 dict），找不到返回 None。"""
    wb, ws = _load_ws(path)
    rows = [["" if c is None else str(c).strip() for c in r] for r in ws.iter_rows(values_only=True)]
    wb.close()
    head_idx = next((i for i, r in enumerate(rows[:6]) if "包号" in r and "分标名称" in r), None)
    if head_idx is None:
        return None
    header = rows[head_idx]
    for r in rows[head_idx + 1:]:
        rec = dict(zip(header, r))
        if rec.get("包号") == pkg_no:
            return rec
    return None


def apply_service_demand(pkg, rec: dict) -> None:
    """把需求行写入 PackageTRM（评审引用 + 资格要求 + 项目信息）。"""
    pkg.project_name = rec.get("招标项目标段名称", "")
    pkg.project_unit = rec.get("项目单位", "")
    pkg.scope = rec.get("招标范围", "")
    pkg.budget_yuan = _f(rec.get("招标金额（元）", ""))
    pkg.max_price = rec.get("最高限价（元或折扣比例）", "")
    d = _f(rec.get("标段工期（日）", ""))
    pkg.duration_days = int(d) if d is not None else None
    pkg.price_mode = rec.get("报价方式", "")
    ac = rec.get("是否允许联合体", "")
    pkg.allow_consortium = True if ac == "是" else (False if ac == "否" else None)
    sr = pkg.scoring_ref
    sr.price_method = rec.get("价格分计算方法", "")
    sr.benchmark_c = _f(rec.get("基准价浮动系数C", ""))
    sr.biz_template = rec.get("商务详评模板", "")
    sr.tech_template = rec.get("技术详评模板", "")
    sr.weight_biz = _f(rec.get("商务分权重（%）", ""))
    sr.weight_tech = _f(rec.get("技术分权重（%）", ""))
    sr.weight_price = _f(rec.get("价格分权重（%）", ""))
    if not pkg.sub_name:
        pkg.sub_name = rec.get("分标名称", "")
    if not pkg.sub_no:
        pkg.sub_no = rec.get("分标编号", "")
    q = QualificationItem(
        sub_name=rec.get("分标名称", ""), pkg=pkg_no_of(rec),
        performance_req=_slash(rec.get("企业业绩要求", "")),
        other_reqs={k: v for k, v in (("企业资质条件", _slash(rec.get("企业资质条件", ""))),
                                     ("主要人员要求", _slash(rec.get("主要人员要求", ""))),
                                     ("技术规范书ID", rec.get("技术规范书ID", ""))) if v})
    pkg.qualification.append(q)


def pkg_no_of(rec: dict) -> str:
    return rec.get("包号", "")


def _slash(v: str) -> str:
    return "" if v.strip() == "/" else v.strip()
