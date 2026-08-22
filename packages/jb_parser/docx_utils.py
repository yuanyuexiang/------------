"""docx 通用读取：章节切分与表格结构化。"""
from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Optional

import docx

CHAPTER_PAT = re.compile(r"^第([一二三四五六七八九十]+)章\s*(.*)$")
_CN_NUM = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}


@dataclass
class Chapter:
    no: int
    title: str
    start: int  # 段落下标（含）
    end: int    # 段落下标（不含）


def load(path: str) -> docx.document.Document:
    return docx.Document(path)


def full_text(doc) -> str:
    """正文段落 + 表格单元格文本（按行拼接），供正则扫描整篇（如招标公告）。"""
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def split_chapters(doc) -> list[Chapter]:
    """按"第N章"标记切分正文。目录区的章名行（连续出现在文首）会被跳过：
    取每个章号最后一次出现的位置作为正文起点。"""
    hits: dict[int, list[int]] = {}
    titles: dict[int, str] = {}
    for i, p in enumerate(doc.paragraphs):
        m = CHAPTER_PAT.match(p.text.strip())
        if not m:
            continue
        no = _CN_NUM.get(m.group(1))
        if not no:
            continue
        hits.setdefault(no, []).append(i)
        titles[no] = m.group(2).strip() or titles.get(no, "")
    start_by_no = {no: idxs[-1] for no, idxs in hits.items()}
    # 正文标题可能省略"第N章"前缀（实测第六章正文只写"投标文件格式"）：
    # 若某章起点早于编号更小的章（说明命中的是目录行），则在后文按裸标题重新定位
    paras = [p.text.strip() for p in doc.paragraphs]
    for no in sorted(start_by_no):
        prev = [start_by_no[k] for k in start_by_no if k < no]
        if prev and start_by_no[no] < max(prev):
            title = titles.get(no, "")
            if title:
                for i in range(max(prev) + 1, len(paras)):
                    if paras[i] == title:
                        start_by_no[no] = i
                        break
    starts = sorted((st, no) for no, st in start_by_no.items())
    chapters: list[Chapter] = []
    for k, (start, no) in enumerate(starts):
        end = starts[k + 1][0] if k + 1 < len(starts) else len(doc.paragraphs)
        chapters.append(Chapter(no=no, title=titles.get(no, ""), start=start, end=end))
    return chapters


def table_rows(table) -> list[list[str]]:
    """表格→二维文本；同一行内被合并单元格产生的重复值去重（保序）。"""
    rows: list[list[str]] = []
    for r in table.rows:
        cells = [c.text.strip() for c in r.cells]
        dedup: list[str] = []
        for c in cells:
            if not dedup or dedup[-1] != c:
                dedup.append(c)
        rows.append(dedup)
    return rows


def find_table(doc, header_keywords: list[str], min_hits: int = 2) -> Optional[list[list[str]]]:
    """按表头关键词定位表格（检查前两行），返回结构化行。"""
    for t in doc.tables:
        if not t.rows:
            continue
        head = " ".join(c.text.strip() for r in t.rows[:2] for c in r.cells)
        hits = sum(1 for k in header_keywords if k in head)
        if hits >= min_hits:
            return table_rows(t)
    return None


def find_tables(doc, header_keywords: list[str], min_hits: int = 2) -> list[list[list[str]]]:
    out = []
    for t in doc.tables:
        if not t.rows:
            continue
        head = " ".join(c.text.strip() for r in t.rows[:2] for c in r.cells)
        if sum(1 for k in header_keywords if k in head) >= min_hits:
            out.append(table_rows(t))
    return out
