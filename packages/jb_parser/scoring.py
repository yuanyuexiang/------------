"""评分细则模板结构化（评标办法前附表之三~六引用的模板库）。

样本形态（福建服务竞谈包）：
  评分细则（NNNNNN）/价格评分模板/区间平均价浮动法.docx
  评分细则（NNNNNN）/商务评分模板/FWSW01：服务类通用商务详评细则.xlsx
  评分细则（NNNNNN）/技术评分模板/1.JS-FWTY：服务技术详评模板（通用）.xlsx
xlsx 结构：首行模板名，次行表头（项目|评审要素|评审内容 或 评审要素|评审内容），
评审要素文本内嵌分值区间，如"项目团队（12-20 分）"。
"""
from __future__ import annotations

import os
import re
from typing import Optional

import openpyxl

from .trm import ScoringItem, ScoringTemplate

SCORE_RANGE_PAT = re.compile(r"[（(]\s*(-?\d+(?:\.\d+)?)\s*[-~—]\s*(-?\d+(?:\.\d+)?)\s*分\s*[）)]")
SINGLE_SCORE_PAT = re.compile(r"[（(]\s*(-?\d+(?:\.\d+)?)\s*分\s*[）)]")


def _score_range(text: str) -> tuple[Optional[float], Optional[float]]:
    m = SCORE_RANGE_PAT.search(text)
    if m:
        return float(m.group(1)), float(m.group(2))
    m = SINGLE_SCORE_PAT.search(text)
    if m:
        v = float(m.group(1))
        return (v, v)
    return None, None


def kind_from_path(relpath: str) -> str:
    if "价格" in relpath:
        return "price"
    if "商务" in relpath:
        return "biz"
    if "技术" in relpath:
        return "tech"
    return "other"


def parse_xlsx(path: str, relpath: str) -> Optional[ScoringTemplate]:
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return None
    ws = wb[wb.sheetnames[0]]
    try:
        ws.reset_dimensions()
    except AttributeError:
        pass
    rows = [["" if c is None else str(c).strip() for c in r]
            for r in ws.iter_rows(values_only=True)]
    wb.close()
    rows = [r for r in rows if any(r)]
    if not rows:
        return None
    name = rows[0][0] if rows[0] else os.path.basename(relpath)
    tpl = ScoringTemplate(name=name, kind=kind_from_path(relpath), source=relpath)
    element = ""
    group = ""
    for r in rows[1:]:
        cells = [c for c in r if c]
        if not cells or cells[0] in ("评审要素", "项目"):
            continue
        # 三列（项目|评审要素|评审内容）：cells[0] 为分组；两列：cells[0] 即要素；单列：续行
        if len(cells) >= 3:
            group, element, content = cells[0], cells[1], cells[-1]
        elif len(cells) == 2:
            element = cells[0] if "分" in cells[0] or len(cells[0]) < 30 else element
            content = cells[-1]
        else:
            content = cells[0]
            if not element and _score_range(content)[0] is not None:
                element = content
        lo, hi = _score_range(element)
        if lo is None:
            lo, hi = _score_range(content)
        if element or content:
            tpl.items.append(ScoringItem(
                element=element, content=content, score_min=lo, score_max=hi, group=group))
    # 合并同要素续行
    merged: list[ScoringItem] = []
    for it in tpl.items:
        if merged and it.element == merged[-1].element:
            merged[-1].content = (merged[-1].content + "\n" + it.content).strip()
            if merged[-1].score_min is None:
                merged[-1].score_min, merged[-1].score_max = it.score_min, it.score_max
        else:
            merged.append(it)
    tpl.items = merged
    return tpl


def parse_docx_name_only(relpath: str) -> ScoringTemplate:
    """价格评分模板为 docx 公式文档，S1 只登记名称，公式参数化在价格模块做。"""
    name = os.path.splitext(os.path.basename(relpath))[0]
    return ScoringTemplate(name=name, kind=kind_from_path(relpath), source=relpath)
