"""产业单位编码包、无规范ID正文及混入封面的回归。"""
from pathlib import Path

import pytest
from docx import Document
from jb_agents.draft_report import confirmation_items
from jb_agents.preflight import generation_issues
from jb_parser import parse
from jb_parser.extract import extract_batch_info
from jb_parser.trm import TRM, PackageTRM


def test_cover_identity_wins_over_instruction_paragraph():
    doc = Document()
    doc.add_paragraph('采购编号：CY8526JFF17C')
    doc.add_paragraph('国网信通产业集团2026年产业单位第十七次服务类公开谈判采购（三）')
    doc.add_paragraph('（2）响应供应商收到短信后需登陆ECP2.0非招标采购模块查看谈判采购通知路径。')
    assert extract_batch_info(doc, '采购文件.docx') == (
        '国网信通产业集团2026年产业单位第十七次服务类公开谈判采购（三）', 'CY8526JFF17C', '应答')


def test_missing_identity_is_system_issue_and_duplicates_are_grouped():
    assert any(x['blocking'] for x in generation_issues(TRM(), PackageTRM()))
    items = confirmation_items(None, {'商务.docx': ['分标名称', '分标名称'],
                                      '技术.docx': ['分标名称', '被授权人姓名', '数值需核实来源']})
    assert items[0].category == '招标原文待核对'
    assert items[0].occurrences == 3 and len(items[0].locations) == 2
    assert items[1].category == '本次投标事实待确认'
    assert items[2].category == '系统生成待修复'


SAMPLE = Path(__file__).resolve().parents[1] / '国网信通产业集团2026年产业单位第十七次服务类公开谈判采购（三）_采购文件包(1).zip'


@pytest.mark.skipif(not SAMPLE.exists(), reason='本地产业单位采购样本不可用')
def test_industry_package_uses_matching_body_not_foreign_cover(tmp_path):
    trm = parse(str(SAMPLE), str(tmp_path))
    assert trm.batch_no == 'CY8526JFF17C'
    assert len(trm.packages) == 1
    pkg = trm.packages[0]
    assert pkg.pkg_no == 'FF17-229'
    assert pkg.project_name == '能源互联网营销服务系统服务支撑'
    assert '4个月' in pkg.duration_text
    assert pkg.max_price == '132.000000万元'
    assert (pkg.scoring_ref.weight_tech, pkg.scoring_ref.weight_price, pkg.scoring_ref.weight_biz) == (40, 50, 10)
    assert len(pkg.spec_docs) == 1 and not pkg.spec_docs[0].spec_id
    assert '能源互联网' in pkg.spec_docs[0].source
    assert not any(x['blocking'] for x in generation_issues(trm, pkg))
    assert any('湖北思极' in x for x in trm.warnings)


def test_transient_draft_failure_retries_and_successful_sections_are_reused(monkeypatch):
    import jb_llm
    from jb_agents.writer import draft_package
    from jb_kb.models import CompanyProfile
    calls = []

    def chat(*args, **kwargs):
        calls.append(args)
        if len(calls) == 1:
            raise RuntimeError('temporary')
        return '拟建立交付记录并进行复核。'

    monkeypatch.setattr(jb_llm, 'available', lambda: True)
    monkeypatch.setattr(jb_llm, 'chat', chat)
    draft = draft_package(TRM(), PackageTRM(), CompanyProfile(name='测试'))
    assert all(s.by_llm for s in draft.sections)
    assert len(calls) == len(draft.sections) + 1
    previous = [s.model_dump() for s in draft.sections]
    previous[0]['by_llm'] = False
    calls.clear()
    result = draft_package(TRM(), PackageTRM(), CompanyProfile(name='测试'), previous_sections=previous)
    assert len(calls) == 2
    assert all(s.by_llm for s in result.sections)


def test_service_deliverable_table_uses_plan_without_claiming_product_performance():
    from jb_docgen.spec_response import respond_service_table, service_table_rows
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    for cell, text in zip(table.rows[0].cells, ['序号', '成果名称', '成果说明', '应答人应答']):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ['1', '验收报告', '验收报告', '']):
        cell.text = text
    key = service_table_rows(table)[0][1]
    respond_service_table(table, {key: '拟整理验收记录并提交采购方审核。[来源:req-service-0]'})
    assert len(table.columns) == 4
    assert table.rows[1].cells[3].text == '拟整理验收记录并提交采购方审核。'


def test_service_table_without_response_column_keeps_original_width():
    from docx import Document
    from jb_docgen.spec_response import respond_service_table, service_table_rows

    table = Document().add_table(rows=2, cols=3)
    for cell, text in zip(table.rows[0].cells, ['序号', '岗位名称', '数量']):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ['1', '项目负责人', '1']):
        cell.text = text
    widths = [column.width for column in table.columns]
    key = service_table_rows(table)[0][1]
    following = respond_service_table(table, {key: '拟核对负责人资格。'})
    assert [column.width for column in table.columns] == widths
    assert following == ['表中第1项响应：拟核对负责人资格。']


def test_compound_requirement_sentence_and_deposit_do_not_become_ability_claims():
    from jb_agents.writer import _check_numbers
    text = '招标要求服务期4个月、保证金21120元；不据此声称本企业已具备履约能力。[来源:req-package]'
    sources = {'req-package': '服务期4个月；保证金21120元'}
    assert _check_numbers(text, sources) == (text, [])
    false_claim = '招标要求4个月；我司已达到4个月交付能力。[来源:req-package]'
    assert _check_numbers(false_claim, sources)[1]
    ordinal = '| 序号 | 内容 |\n| （1） | 项目调研 |'
    assert _check_numbers(ordinal, {}) == (ordinal, [])


def test_contract_duration_attribution_requires_exact_source_and_keeps_claim_guard():
    from jb_agents.draft_context import attribute_contract_duration
    from jb_agents.writer import _check_numbers
    sources = {'req-package': '服务期原文：自合同签订之日起，4个月内完成服务'}
    text = '服务期按合同签订之日起4个月编排。'
    attributed = attribute_contract_duration(text, sources)
    assert _check_numbers(attributed, sources)[1] == []
    unmatched = '服务期按合同签订之日起5个月编排。'
    assert attribute_contract_duration(unmatched, sources) == unmatched
    claim = attribute_contract_duration('我司已实现合同签订之日起4个月交付。', sources)
    assert _check_numbers(claim, sources)[1]


def test_industry_basic_info_layout_fills_known_fields_and_marks_missing_counts():
    from jb_docgen.builder import GenContext
    from jb_docgen.master_builder import _fill_basic_info
    from jb_kb.models import CompanyProfile
    doc = Document()
    labels = ['单位名称', '单位地址及邮编', '主管部门', '单位性质', '开户银行及账号',
              '注册资金（万元）', 'E-mail', '员工总数', '其中：技术人员数']
    table = doc.add_table(rows=len(labels), cols=2)
    for row, label in zip(table.rows, labels):
        row.cells[0].text = label
    profile = CompanyProfile(name='测试公司', address='测试地址', company_type='有限责任公司',
                             bank='测试银行', registered_capital_wan=100, email='test@example.com',
                             staff_total=0)
    _fill_basic_info(doc, GenContext(TRM(), PackageTRM(), profile))
    assert table.rows[0].cells[1].text == '测试公司'
    assert table.rows[1].cells[1].text == '测试地址'
    assert table.rows[4].cells[1].text == '测试银行'
    assert table.rows[5].cells[1].text == '100万元'
    assert table.rows[6].cells[1].text == 'test@example.com'
    assert table.rows[7].cells[1].text == '0'
    assert '待补充' in table.rows[8].cells[1].text
    assert any('主管部门' in p.text and '待补充' in p.text for p in doc.paragraphs)
