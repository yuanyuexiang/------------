"""正式导出的 API 回归：真实规则引擎 + 合成文档，不依赖客户样本或 PDF 服务。"""
from pathlib import Path

import docx
import pytest
from jb_docgen import GenResult
from jb_docgen.placeholders import scan_docx
from jb_parser.trm import TRM, PackageTRM, SubmissionItem
from jb_store import Project, Task, session


@pytest.fixture()
def bid(client, tmp_path, monkeypatch):
    from jb_api import main

    body = {"text": "已核对的投标正文"}

    def generate(trm, pkg, profile, out_dir, *args, **kwargs):
        paths = []
        for kind in ("商务", "技术"):
            path = Path(out_dir) / f"{kind}文件_{pkg.sub_no}{pkg.pkg_no}.docx"
            d = docx.Document()
            d.core_properties.author = "测试作者"
            d.add_paragraph(profile.name + "\n" + body["text"])
            d.save(path)
            paths.append(str(path))
        return GenResult(*paths, docx_todos={Path(p).name: scan_docx(p) for p in paths})

    monkeypatch.setattr("jb_docgen.generate", generate)
    client.put("/api/profiles/测试企业", json={"name": "测试企业"})
    trm = TRM(packages=[PackageTRM(sub_no="A", pkg_no="包1"),
                        PackageTRM(sub_no="B", pkg_no="包1")],
              submission_table=[SubmissionItem(section="商务文件", item="基本情况")]).model_dump()
    with session() as s:
        p = Project(filename="合成.zip", zip_path="", trm=trm, trm_confirmed=trm,
                    status="confirmed")
        s.add(p)
        s.flush()
        pid = p.id
        folder = Path(main.UPLOAD_DIR) / pid
        folder.mkdir(parents=True)
        p.zip_path = str(folder / "合成.zip")
    return {"client": client, "id": pid, "url": f"/api/projects/{pid}", "folder": folder,
            "body": body, "trm": trm}


def generate(bid, index=0):
    c, url = bid["client"], bid["url"]
    r = c.post(url + "/generate", params={"profile": "测试企业", "pkg_index": index})
    assert r.status_code == 202, r.text
    task = c.get("/api/tasks/" + r.json()["task_id"]).json()
    assert task["status"] == "done", task
    return task["result"]["files"][0]


def review(bid, index=0):
    return bid["client"].post(bid["url"] + "/review",
                              params={"profile": "测试企业", "pkg_index": index})


def export(bid, filename, **params):
    r = bid["client"].post(bid["url"] + "/export",
                           params={"filename": filename, "pdf": False, **params})
    assert r.status_code == 200, r.text
    return r.json()


def state(bid, index=0):
    return bid["client"].get(bid["url"] + "/workflow", params={"pkg_index": index}).json()


def test_confirmation_review_and_force_gate(bid):
    with session() as s:
        s.get(Project, bid["id"]).trm_confirmed = None
    name = generate(bid)
    assert "尚未人工确认" in "；".join(export(bid, name)["blocked_by"])
    assert review(bid).status_code == 409
    bid["client"].put(bid["url"] + "/trm", json=bid["trm"])
    assert "尚未审查" in "；".join(export(bid, name)["blocked_by"])
    assert review(bid).json()["blocked"] is False
    assert export(bid, name, force=True)["ok"] is False
    assert export(bid, name)["ok"] is True


def test_formal_copy_repeat_export_and_restored_review(bid, monkeypatch):
    name = generate(bid)
    path = bid["folder"] / "out" / name
    before = path.read_bytes()
    rep = review(bid).json()
    assert state(bid)["review"]["findings"] == rep["findings"]
    assert state(bid)["generation"]["files"][0] == name

    def pdf(path, out_dir=None):
        target = Path(path).with_suffix(".pdf")
        target.write_bytes(b"%PDF-1.4\n%%EOF")
        return str(target), "test converter"

    monkeypatch.setattr("jb_docgen.export.to_pdf", pdf)
    result = export(bid, name, pdf=True)
    assert result["ok"] and result["pdf"]
    assert path.read_bytes() == before
    formal = bid["folder"] / "out" / result["docx"]
    assert docx.Document(formal).core_properties.author == ""
    response = bid["client"].get(bid["url"] + "/files/" + result["pdf"])
    assert response.status_code == 200 and response.headers["content-type"] == "application/pdf"
    assert export(bid, name)["ok"]
    assert len(state(bid)["exports"]) == 2


@pytest.mark.parametrize("change", ["trm", "profile", "attachment", "date"])
def test_changed_inputs_require_regeneration_and_revoke_download(bid, change):
    name = generate(bid)
    review(bid)
    formal = export(bid, name)["docx"]
    c, url = bid["client"], bid["url"]
    if change == "trm":
        bid["trm"]["key_terms"]["validity_days"] = 120
        c.put(url + "/trm", json=bid["trm"])
    elif change == "profile":
        c.put("/api/profiles/测试企业", json={"name": "测试企业", "credit_code": "changed"})
    elif change == "date":
        c.patch(url, json={"open_time": "2030-01-01 09:00"})
    else:
        from jb_store.models import KbAttachment
        with session() as s:
            s.add(KbAttachment(company="测试企业", filename="证明.pdf", content_type="application/pdf",
                               size=10, sha256="new", storage_path="synthetic"))
    assert export(bid, name)["ok"] is False
    assert review(bid).status_code == 409
    assert c.get(url + "/files/" + formal).status_code == 409
    assert state(bid)["exports"] == []
    new_name = generate(bid)
    assert new_name != name
    assert export(bid, new_name)["ok"] is False
    assert review(bid).status_code == 200
    assert export(bid, new_name)["ok"] is True
    assert export(bid, name)["ok"] is False


def test_rule_change_requires_review_and_todos_cannot_be_disabled(bid):
    name = generate(bid)
    review(bid)
    c = bid["client"]
    assert c.put("/api/config/rules/DOC-01", json={"enabled": False}).status_code == 200
    assert "规则配置已变化" in "；".join(export(bid, name)["blocked_by"])
    assert review(bid).status_code == 200
    assert export(bid, name)["ok"]
    bid["body"]["text"] = "【待补充：营业执照】"
    name = generate(bid)
    assert review(bid).json()["blocked"] is False
    assert export(bid, name)["ok"] is False


def test_veto_blocks_even_without_placeholders(bid):
    bid["body"]["text"] = "W-2025-OTHER-Z01"
    name = generate(bid)
    rep = review(bid).json()
    assert rep["blocked"] and any(f["rule_id"] == "DOC-04" for f in rep["findings"])
    assert export(bid, name)["ok"] is False
    assert export(bid, name, force=True)["ok"] is False


def test_identical_package_numbers_do_not_share_review(bid):
    first = generate(bid, 0)
    second = generate(bid, 1)
    assert first != second
    review(bid, 0)
    assert export(bid, first)["ok"] is True
    assert export(bid, second)["ok"] is False
    assert state(bid, 0)["generation_id"] != state(bid, 1)["generation_id"]
    assert state(bid, 1)["review"] is None
    matrix = bid["client"].get(bid["url"] + "/submission-matrix", params={"pkg_index": 1}).json()
    assert matrix["rows"][0]["generated_file"] == second
    r = bid["client"].post(bid["url"] + "/review", params={"profile": "另一企业"})
    assert r.status_code == 409


@pytest.mark.parametrize("damage", ["delete", "edit"])
def test_any_changed_document_blocks_entire_version(bid, damage):
    name = generate(bid)
    review(bid)
    other = bid["folder"] / "out" / state(bid)["generation"]["files"][1]
    if damage == "delete":
        other.unlink()
    else:
        other.write_bytes(b"broken docx")
    assert export(bid, name)["ok"] is False
    assert review(bid).status_code == 409


def test_legacy_draft_and_running_generation(bid):
    folder = bid["folder"] / "out"
    folder.mkdir()
    d = docx.Document()
    d.save(folder / "legacy.docx")
    assert bid["client"].get(bid["url"] + "/files/legacy.docx").status_code == 200
    assert export(bid, "legacy.docx")["ok"] is False
    name = generate(bid)
    review(bid)
    with session() as s:
        s.add(Task(project_id=bid["id"], kind="generate", status="queued"))
    assert export(bid, name)["ok"] is False
    assert review(bid).status_code == 409
    assert state(bid)["active_task_id"]
    assert bid["client"].post(bid["url"] + "/generate", params={"profile": "测试企业"}).status_code == 409


def test_open_date_used_by_review(bid):
    c = bid["client"]
    c.put("/api/profiles/测试企业", json={"name": "测试企业", "certificates": [
        {"name": "过期证书", "valid_until": "2025-01-01"}]})
    c.patch(bid["url"], json={"open_time": "2030-01-01 09:00"})
    name = generate(bid)
    rep = review(bid).json()
    assert any(f["rule_id"] == "SG-13" for f in rep["findings"])
    assert export(bid, name)["ok"] is False


def test_change_during_conversion_prevents_publication(bid, monkeypatch):
    from jb_docgen import workflow
    from jb_store import Profile

    name = generate(bid)
    review(bid)
    original = workflow.export

    def convert(path, want_pdf=True):
        result = original(path, want_pdf=False)
        with session() as s:
            row = s.get(Profile, "测试企业")
            row.data = {**row.data, "credit_code": "updated-during-conversion"}
        return result

    monkeypatch.setattr(workflow, "export", convert)
    assert export(bid, name)["ok"] is False
    assert not list((bid["folder"] / "out").glob("正式_*"))


def test_duplicate_package_summaries_are_isolated(bid):
    generate(bid, 0)
    generate(bid, 1)
    review(bid, 0)
    review(bid, 1)
    c, url = bid['client'], bid['url']
    assert c.post(url + '/qualify', params={'profile': '测试企业'}).status_code == 200
    for index in (0, 1):
        assert c.post(url + '/score', params={'profile': '测试企业', 'pkg_index': index}).status_code == 200
    detail = c.get(url + '/detail').json()
    keys = {row['result_key'] for row in detail['package_list']}
    assert len(keys) == 2
    for kind in ('generate', 'review', 'score'):
        assert set(detail['results'][kind]) == keys
    assert set(detail['results']['qualify']['verdicts']) == keys


def test_score_rejects_wrong_package_project_and_stale_inputs(bid):
    generate(bid, 0)
    task = state(bid, 0)['generation_id']
    c, url = bid['client'], bid['url']
    params = {'profile': '测试企业', 'task_id': task}
    assert c.post(url + '/score', params=params).status_code == 200
    assert c.post(url + '/score', params={**params, 'pkg_index': 1}).status_code == 409
    assert c.post(url + '/score', params={**params, 'pkg_index': -1}).status_code == 400
    assert c.post(url + '/score', params={**params, 'pkg_index': 9}).status_code == 400
    with session() as s:
        original = s.get(Project, bid['id'])
        other = Project(filename='other.zip', zip_path='', trm=original.trm,
                        trm_confirmed=original.trm_confirmed)
        s.add(other)
        s.flush()
        other_id = other.id
    assert c.post(f'/api/projects/{other_id}/score', params=params).status_code == 409
    c.put('/api/profiles/测试企业', json={'name': '测试企业', 'credit_code': 'changed'})
    assert c.post(url + '/score', params=params).status_code == 409


def test_package_selection_is_scoped_and_revokes_previous_exports(bid):
    c, url = bid['client'], bid['url']
    a = c.post('/api/profiles/测试企业/personnel', json={'name': 'A'}).json()
    b = c.post('/api/profiles/测试企业/personnel', json={'name': 'B', 'status': 'draft'}).json()
    name = generate(bid)
    review(bid)
    formal = export(bid, name)['docx']
    params = {'profile': '测试企业', 'pkg_index': 0}
    assert c.put(url + '/materials', params=params, json={'personnel': [b['id']]}).status_code == 400
    assert c.put(url + '/materials', params=params, json={'personnel': [a['id']]}).status_code == 200
    assert c.get(url + '/materials', params=params).json()['selection']['personnel'] == [a['id']]
    assert c.get(url + '/materials', params={**params, 'pkg_index': 1}).json()['selection']['personnel'] is None
    assert c.get(url + '/files/' + formal).status_code == 409


def test_generation_defaults_to_ai_and_restores_confirmation_list(bid, monkeypatch):
    from jb_agents.writer import Draft, DraftSection

    calls = []

    def draft(trm, pkg, profile, **kwargs):
        calls.append(kwargs)
        kwargs['on_progress'](0, 1, '实施方案')
        return Draft(pkg_no=pkg.pkg_no, sections=[DraftSection(
            title='实施方案', text='拟开展需求调研并交付报告。', by_llm=True)])

    monkeypatch.setattr('jb_agents.writer.draft_package', draft)
    generate(bid)
    restored = state(bid)['generation']
    assert len(calls) == 1 and 'spec_paths' in calls[0]
    assert restored['draft_sections'][0]['by_llm']
    assert restored['confirmations'][0]['category'] == '拟实施方案待确认'
    result = bid['client'].post(bid['url'] + '/generate', params={
        'profile': '测试企业', 'pkg_index': 0, 'with_draft': False,
    })
    assert result.status_code == 202
    assert len(calls) == 1


def test_relationship_confirmation_revokes_export_and_reaches_generation(bid, monkeypatch):
    c, url = bid['client'], bid['url']
    name = generate(bid)
    review(bid)
    formal = export(bid, name)['docx']
    params = {'profile': '测试企业', 'pkg_index': 0}
    source = c.get(url + '/relationship', params=params).json()['source']
    assert c.put(url + '/relationship', params=params,
                 json={'source': source, 'status': 'none'}).status_code == 200
    assert c.get(url + '/files/' + formal).status_code == 409
    import jb_docgen
    original = jb_docgen.generate
    captured = []

    def capture(*args, **kwargs):
        captured.append(kwargs['relationship_confirmation'])
        return original(*args, **kwargs)

    monkeypatch.setattr(jb_docgen, 'generate', capture)
    generate(bid)
    assert captured[0]['status'] == 'none' and captured[0]['confirmed_by'] == 'admin'
    generate(bid, 1)
    assert captured[1] is None
