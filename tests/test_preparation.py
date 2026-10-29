"""投标准备闭环：作用域、失效、证据保留、模板填充及候选不自动采信。"""

import pytest
from docx import Document
from jb_agents.preparation import validate_candidates
from jb_docgen import generate, preparation, workflow
from jb_kb import repo
from jb_kb.models import CompanyProfile, Performance, Person
from jb_parser.trm import TRM, PackageTRM
from jb_store import Project, Task, session


@pytest.fixture()
def preparation_project(client):
    company = CompanyProfile(name='准备测试企业', authorized_rep='常用联系人', authorized_rep_title='职员',
        personnel=[Person(id='person-a', name='甲', title='工程师', source='学历证明', attachments=['attachment-a']),
                   Person(id='person-b', name='乙', title='技术员', status='draft')],
        performances=[Performance(id='perf-a', project='历史项目A', source='签字页日期空白',
                                  attachments=['contract-a'], evidence=['合同', '发票'])])
    assert client.put('/api/profiles/准备测试企业', json=company.model_dump()).status_code == 200
    trm = TRM(batch_no='TEST', batch_name='测试采购', packages=[PackageTRM(pkg_no='包1'), PackageTRM(pkg_no='包2')])
    with session() as s:
        p = Project(filename='test.zip', zip_path='', trm=trm.model_dump(), trm_confirmed=trm.model_dump())
        s.add(p)
        s.flush()
        pid = p.id
    return client, pid


def state(case, index=0):
    c, pid = case
    r = c.get(f'/api/projects/{pid}/preparation', params={'profile': '准备测试企业', 'pkg_index': index})
    assert r.status_code == 200, r.text
    return r.json()


def save(case, kind, data, revision=None, index=0):
    c, pid = case
    return c.put(f'/api/projects/{pid}/preparation/{kind}', params={'profile': '准备测试企业', 'pkg_index': index},
                 json={'revision': revision or state(case, index)['revision'], 'data': data})


def test_authorization_scoped_and_reconfirmation_on_person_change(preparation_project):
    case = preparation_project
    assert save(case, 'authorization', {'person_id': 'person-a', 'title': '项目授权代表'}).status_code == 200
    c, pid = case
    assert state(case, 1)['facts']['authorization']['confirmation'] is None
    with session() as s:
        cp = repo.load_profile(s, '准备测试企业')
        assert cp.authorized_rep == '常用联系人'
        effective = workflow.selected_profile(s.get(Project, pid), 0, cp)
        assert effective.authorized_rep == '甲' and effective.authorized_rep_title == '项目授权代表'
        person = cp.personnel[0]
        repo.upsert_item(s, cp.name, 'personnel', {**person.model_dump(), 'title': '高级工程师'})
    assert state(case)['facts']['authorization']['stale']
    with session() as s:
        cp = repo.load_profile(s, '准备测试企业')
        assert workflow.selected_profile(s.get(Project, pid), 0, cp).authorized_rep == ''


@pytest.mark.parametrize('payload', [
    {'person_id': 'person-b', 'title': '代表'}, {'person_id': 'foreign', 'title': '代表'},
    {'person_id': 'person-a', 'title': '代表', 'identity_attachments': ['foreign-attachment']},
])
def test_authorization_rejects_unapproved_or_foreign_records(preparation_project, payload):
    assert save(preparation_project, 'authorization', payload).status_code == 409


def test_team_selection_is_confirmed_once_and_does_not_modify_global_availability(preparation_project, monkeypatch):
    case = preparation_project
    monkeypatch.setattr(preparation, 'requirements', lambda *a: ({}, [{'role': '项目负责人', 'count': 1}, {'role': '技术服务人员', 'count': 9}]))
    payload = {'members': [{'person_id': 'person-a', 'role': '项目负责人', 'available': True}], 'arrival_commitment': True}
    assert save(case, 'team', payload).status_code == 200
    with session() as s:
        cp = repo.load_profile(s, '准备测试企业')
        p = s.get(Project, case[1])
        assert cp.personnel[0].available is None
        selected = workflow.selected_profile(p, 0, cp)
        assert selected.personnel[0].available is True
        assert selected.bid_context['team']['members'][0]['role'] == '项目负责人'
        assert any('要求9人' in x for x in preparation.team_blockers(p, 0, cp))
        workflow.save_selection(p, 0, cp, {'personnel': []})
    assert state(case)['facts']['team']['confirmation'] is None


def test_stale_revision_and_running_generation_cannot_overwrite(preparation_project):
    case = preparation_project
    revision = state(case)['revision']
    assert save(case, 'authorization', {'person_id': 'person-a', 'title': '代表'}).status_code == 200
    assert save(case, 'authorization', {'person_id': 'person-a', 'title': '其他'}, revision).status_code == 409
    with session() as s:
        s.add(Task(project_id=case[1], kind='generate', status='running'))
    assert save(case, 'authorization', {'person_id': 'person-a', 'title': '其他'}).status_code == 409


def test_statistics_zero_is_not_unknown_and_date_and_basis_are_required(preparation_project):
    case = preparation_project
    assert save(case, 'statistics', {'staff_total': 0}).status_code == 409
    valid = {'staff_total': 0, 'staff_technical': 0, 'staff_as_of': '2026-10-29', 'staff_source': '人事确认台账'}
    assert save(case, 'statistics', valid).status_code == 200
    assert state(case)['profile']['staff_total'] == 0
    assert save(case, 'statistics', {**valid, 'staff_technical': 1}).status_code == 409
    assert save(case, 'statistics', {**valid, 'staff_total': -1}).status_code == 409
    assert save(case, 'statistics', {**valid, 'staff_as_of': '2026-02-31'}).status_code == 409


def test_registration_na_needs_reason_and_shareholding_cannot_exceed_100(preparation_project):
    case = preparation_project
    assert save(case, 'registration', {'registration': {'主管部门': {'not_applicable': True}}}).status_code == 409
    data = {'registration': {'主管部门': {'not_applicable': True, 'reason': '经企业核对采购表格要求不适用'}}}
    assert save(case, 'registration', data).status_code == 200
    owners = [{'name': '股东甲', 'ratio': 70, 'source': '章程第1页'}, {'name': '股东乙', 'ratio': 40, 'source': '章程第1页'}]
    assert save(case, 'registration', {'shareholders': owners}).status_code == 409


def test_contract_date_is_item_specific_and_preserves_evidence(preparation_project):
    case = preparation_project
    st = state(case)
    assert st['missing_dates'][0]['project'] == '历史项目A'
    assert save(case, 'performance', {'id': 'perf-a', 'signed_date': '2025-05', 'source': '约为2025年5月'}).status_code == 409
    assert save(case, 'performance', {'id': 'perf-a', 'signed_date': '2025-05-06', 'source': '合同签署页第9页'}).status_code == 200
    p = state(case)['profile']['performances'][0]
    assert p['attachments'] == ['contract-a'] and p['evidence'] == ['合同', '发票']
    assert p['source'].startswith('签字页日期空白') and '第9页' in p['source']
    assert not state(case)['missing_dates']


def test_generation_consumes_project_facts_and_names_missing_historical_contract(tmp_path):
    cp = CompanyProfile(name='公司', authorized_rep='甲', authorized_rep_title='经理',
                        personnel=[Person(id='a', name='甲', source='身份证明')],
                        performances=[Performance(project='独立历史合同')])
    cp.bid_context = {'team': {'members': [{'person_id': 'a', 'name': '甲', 'role': '项目负责人', 'available': True}], 'arrival_commitment': True}}
    result = generate(TRM(), PackageTRM(), cp, str(tmp_path))
    d = Document(result.technical_path)
    text = '\n'.join(p.text for p in d.paragraphs)
    assert '本次人员到岗承诺' in text
    assert any('独立历史合同' in x and '签约日期' in x for x in result.docx_todos[PathName(result.technical_path)])
    assert 'bid_context' not in cp.model_dump()


def PathName(path):
    from pathlib import Path
    return Path(path).name


def test_ai_suggestions_require_verbatim_evidence_and_never_infer_approximate_dates():
    evidence = [{'id': 'a', 'title': '合同第9页', 'text': '客户说明：签订时间约为2025年5月6日', 'attachment_id': 'att'}]
    result = {'candidates': [{'field': 'signed_date', 'value': '2025-05-06', 'source_id': 'a', 'quote': evidence[0]['text']}]}
    assert validate_candidates(result, evidence, 'performance') == []
    evidence[0]['text'] = '合同签署日期：2025年5月6日'
    result['candidates'][0]['quote'] = evidence[0]['text']
    assert validate_candidates(result, evidence, 'performance')[0]['value'] == '2025-05-06'
    result['candidates'][0]['quote'] = '伪造的引用2025年5月6日'
    assert validate_candidates(result, evidence, 'performance') == []


def test_repeated_personnel_gaps_group_without_losing_original_locations(preparation_project):
    case = preparation_project
    with session() as s:
        p = s.get(Project, case[1])
        digest = workflow.inputs(s, p, '准备测试企业')[2]
        t = Task(project_id=p.id, kind='generate', status='done', result={'todos': {'技术.docx': ['拟投入人员名单', '拟投入人员名单', '项目负责人毕业证明']}})
        s.add(t)
        s.flush()
        workflow.save_artifact(p, {'id': t.id, 'pkg_index': 0, 'profile': '准备测试企业', 'created_at': '2026-10-29', 'input_hash': digest, 'files': {}})
    group = next(x for x in state(case)['tasks'] if x['id'] == 'team')
    assert group['occurrences'] == 3 and len(group['items']) == 2
    assert group['items'][0]['file'] == '技术.docx'


def test_prepared_company_fields_fill_original_template():
    from jb_docgen.builder import GenContext
    from jb_docgen.master_builder import _fill_basic_info
    cp = CompanyProfile(name='模板测试企业', staff_total=12, staff_technical=10,
                        assistant_engineers=0, technicians=2,
                        staff_as_of='2026-10-29', staff_source='人员台账',
                        registration={'主管部门': {'not_applicable': True, 'reason': '企业已核实'}},
                        shareholders=[{'name': '股东甲', 'amount_wan': 100, 'ratio': 100, 'source': '章程'}])
    doc = Document()
    table = doc.add_table(rows=5, cols=2)
    for row, label in zip(table.rows, ['投标人名称', '注册地址', '主管部门', '助理工程师', '技术员']):
        row.cells[0].text = label
    owners = doc.add_table(rows=2, cols=3)
    for cell, text in zip(owners.rows[0].cells, ['股东或上级管理单位名称', '出资金额', '出资比例']):
        cell.text = text
    _fill_basic_info(doc, GenContext(TRM(), PackageTRM(), cp))
    assert table.rows[2].cells[1].text == '不适用（企业已核实）'
    assert table.rows[3].cells[1].text == '0'
    assert '股东甲' in owners.rows[1].cells[0].text
    assert owners.rows[1].cells[2].text == '100%'
    assert any('2026-10-29' in p.text and '人员台账' in p.text for p in doc.paragraphs)


def test_suggestion_job_does_not_write_candidate_to_company(preparation_project, monkeypatch):
    import jb_agents.preparation as agent
    from jb_api import tasks
    from jb_store import llm_connections
    case = preparation_project
    monkeypatch.setattr(tasks, 'celery_app', None)
    monkeypatch.setattr(llm_connections, 'resolve', lambda: {})
    monkeypatch.setattr(agent, 'suggest', lambda *a: {
        'candidates': [{'field': 'staff_total', 'value': 20, 'quote': '员工总数20人', 'source': '台账', 'attachment_id': ''}],
        'checked': [], 'unreadable': [], 'message': '待确认'})
    before = state(case)
    response = case[0].post(f'/api/projects/{case[1]}/preparation/suggest/statistics', params={'profile': '准备测试企业'})
    assert response.status_code == 202
    with session() as s:
        task = s.get(Task, response.json()['task_id'])
        assert task.status == 'done', task.message
        assert task.result['revision'] == before['revision']
        assert task.result['candidates'][0]['value'] == 20
    assert state(case)['profile']['staff_total'] is None


def test_empty_team_is_rejected_and_legacy_record_does_not_become_confirmation(preparation_project):
    case = preparation_project
    assert save(case, 'team', {'members': [], 'arrival_commitment': True}).status_code == 409
    assert save(case, 'team', {'members': [], 'arrival_commitment': False}).status_code == 409
    assert save(case, 'team', {'members': [{'person_id': 'person-a', 'role': '项目负责人', 'available': False}], 'arrival_commitment': True}).status_code == 409
    with session() as s:
        p = s.get(Project, case[1])
        cp = repo.load_profile(s, '准备测试企业')
        data = {'members': [], 'arrival_commitment': True}
        p.results = {'bid_facts': {preparation._key(0, cp.name): {'team': {
            'data': data, 'basis': preparation._basis(p, 0, cp, 'team', data)}}}}
        assert preparation.fact_state(p, 0, cp)['team']['stale']
        effective = preparation.apply_facts(p, 0, cp)
        assert 'team' not in effective.bid_context
        assert effective.bid_context['personnel_candidates'][0]['name'] == '甲'


def test_registration_context_and_unknown_are_preserved(preparation_project):
    from jb_parser.trm import FormatBlock
    _, pid = preparation_project
    with session() as s:
        p = s.get(Project, pid)
        trm = TRM.model_validate(p.trm_confirmed)
        trm.format_blocks = [FormatBlock(title='企业基本情况表', paragraphs=['主管部门、批准文号']),
                             FormatBlock(title='其他章节', paragraphs=['无关内容'])]
        p.trm_confirmed = trm.model_dump()
    before = state(preparation_project)
    assert before['registration_requirements'] == [
        {'title': '企业基本情况表', 'text': '主管部门、批准文号'}]
    assert save(preparation_project, 'registration', {
        'registration': {'主管部门': {'value': '', 'not_applicable': False, 'reason': '', 'source': ''}},
        'shareholders': []}).status_code == 200
    after = state(preparation_project)
    assert not after['profile']['registration']['主管部门']['not_applicable']
    assert '主管部门' in next(t for t in after['tasks'] if t['id'] == 'registration')['current_notes']


def test_registration_retrieval_prioritizes_direct_source(monkeypatch):
    from jb_agents import preparation as agent
    evidence = [
        {'id': 'general', 'title': '通用信息', 'attachment_id': '', 'text': '一般资料' * 12000},
        {'id': 'owner', 'title': '公示报告 第2页', 'attachment_id': 'report',
         'text': '股东：测试股东，出资比例尚未记载'},
    ]
    monkeypatch.setattr(agent, 'collect_evidence', lambda *args: (evidence, []))

    class Model:
        def chat_json(self, prompt, **kwargs):
            assert '公示报告 第2页' in prompt
            return {'candidates': []}

    result = agent.suggest(None, CompanyProfile(name='测试'), 'registration', '', Model())
    assert result['checked'] == [{'title': '公示报告 第2页', 'attachment_id': 'report'}]
    assert any('通用信息' in value for value in result['unreadable'])


def test_suggestion_reconnects_active_job_without_duplicate(preparation_project, monkeypatch):
    from jb_api import tasks
    calls = []
    monkeypatch.setattr(tasks, 'submit_suggestion', lambda *args: calls.append(args) or 'test')
    c, pid = preparation_project
    url = f'/api/projects/{pid}/preparation/suggest/'
    params = {'profile': '准备测试企业'}
    first = c.post(url + 'registration', params=params).json()
    repeated = c.post(url + 'registration', params=params)
    assert repeated.status_code == 202
    assert repeated.json()['task_id'] == first['task_id']
    assert not repeated.json()['busy_other']
    other = c.post(url + 'team', params=params)
    assert other.status_code == 202 and other.json()['busy_other']
    assert other.json()['task_id'] == first['task_id']
    assert state(preparation_project)['active_suggestion'] == first['task_id']
    assert len(calls) == 1


def test_completed_authorization_and_incomplete_team_are_distinct(preparation_project):
    case = preparation_project
    assert save(case, 'authorization', {'person_id': 'person-a', 'title': '授权代表'}).status_code == 200
    current = {t['id']: t for t in state(case)['tasks']}
    assert current['authorization']['complete']
    assert current['authorization']['status'] == '已确认'
    assert not current['team']['complete']
    assert save(case, 'team', {'members': [{'person_id': 'person-a', 'role': '技术服务', 'available': False}],
                               'arrival_commitment': False}).status_code == 200
    team = next(t for t in state(case)['tasks'] if t['id'] == 'team')
    assert not team['complete'] and team['status'] == '待完善团队安排'
    assert any('不可投入' in note for note in team['current_notes'])
