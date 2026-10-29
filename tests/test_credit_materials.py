from types import SimpleNamespace

from docx import Document
from jb_agents.qualify import check_general
from jb_docgen.evidence import insert_credit
from jb_kb.credit import credit_check, credit_materials
from jb_kb.models import Certificate, CompanyProfile
from jb_parser.trm import TRM, PackageTRM
from jb_rules import Context, r_general_credit


def test_credit_checks_reuse_evidence_without_claiming_qualification():
    cp = CompanyProfile(name='企业', certificates=[
        Certificate(name='信用中国报告', valid_from='2026-10-21', attachments=['a']),
        Certificate(name='全国建筑市场监管公共服务平台不良行为查询截图', attachments=['b'],
                    source='国网名单/中国裁判文书网截图/混合目录.png'),
        Certificate(name='国网供应商不良行为', status='draft', attachments=['c']),
        Certificate(name='国家企业信用信息公示报告'),
    ])
    reason, evidence = credit_check(cp)
    assert '已关联查询材料：信用中国' in reason
    assert '2026-10-21' in evidence[0]
    assert not credit_materials(cp)['国网供应商不良行为']
    assert not credit_materials(cp)['中国裁判文书网']
    assert '尚未关联：国家企业信用信息公示系统' in reason
    pkg = PackageTRM()
    result = check_general(pkg, cp)[0]
    assert result.status == '需人工'
    assert result.reason == reason
    finding = r_general_credit(Context(trm=TRM(), pkg=pkg, profile=cp))[0]
    assert reason in finding.message and evidence[0] in finding.message


def test_credit_assembly_includes_split_sites_and_deduplicates_attachments(monkeypatch):
    cp = CompanyProfile(name='企业', certificates=[
        Certificate(name='信用中国报告', attachments=['report', 'screen']),
        Certificate(name='企业信用信息公示报告', attachments=['public']),
        Certificate(name='中国裁判文书网截图', attachments=['court']),
        Certificate(name='中国执行信息公开网查询', attachments=['execution', 'screen']),
        Certificate(name='全国建筑市场监管公共服务平台不良行为', attachments=['building']),
        Certificate(name='国网供应商不良行为名单', attachments=['grid']),
    ])
    calls = []
    def insert(ctx, doc, anchor, ids, label):
        calls.append((label, ids))
        return anchor
    monkeypatch.setattr('jb_docgen.evidence.insert_scans', insert)
    d = Document()
    insert_credit(SimpleNamespace(profile=cp), d, d.add_paragraph('查询材料')._p)
    ids = [aid for _, items in calls for aid in items]
    assert ids == ['report', 'screen', 'public', 'court', 'execution', 'building', 'grid']
    assert len(calls) == 6
    assert '中国执行信息公开网查询材料' in '\n'.join(p.text for p in d.paragraphs)


def test_grid_requirement_distinguishes_condition_attachment_and_declaration():
    from jb_kb.credit import grid_credit_requirement
    from jb_parser.trm import FormatBlock, PrenoticeClause, RejectionRule
    cp = CompanyProfile(name='企业')
    pkg = PackageTRM()
    trm = TRM(rejection_rules=[RejectionRule(aspect='不良行为', text=
        '根据国家电网有限公司供应商关系管理办法，投标人存在不良行为且在处理有效期内的将被否决。')])
    assert grid_credit_requirement(trm, pkg)['mode'] == 'condition'
    message, evidence = credit_check(cp, trm, pkg)
    assert '不列为截图缺项' in message and '处理有效期' in evidence[-1]
    assert '国网供应商不良行为' not in message.split('尚未关联：')[1].split('；')[0]
    trm.prenotice = [PrenoticeClause(content='投标人须提交国家电网供应商不良行为查询截图。')]
    assert grid_credit_requirement(trm, pkg)['mode'] == 'evidence'
    assert '无不良行为也应提供对应查询结果' in check_general(pkg, cp, trm)[0].reason
    assert '无不良行为也应提供对应查询结果' in r_general_credit(Context(trm=trm, pkg=pkg, profile=cp))[0].message
    trm.prenotice = [PrenoticeClause(content='采购人自行核查国网供应商不良行为，投标人无需提交查询截图。')]
    assert grid_credit_requirement(trm, pkg)['mode'] == 'exempt'
    trm.prenotice = []
    trm.format_blocks = [FormatBlock(title='国网供应商不良行为承诺书', paragraphs=['本企业声明……'])]
    assert grid_credit_requirement(trm, pkg)['mode'] == 'declaration'
    assert '审阅生成稿中的原文承诺' in credit_check(cp, trm, pkg)[0]


def test_grid_requirement_does_not_borrow_other_site_or_package_requirements():
    from jb_kb.credit import grid_credit_requirement
    from jb_parser.trm import PrenoticeClause, QualificationItem
    a = PackageTRM(qualification=[QualificationItem(other_reqs={
        '信用': '投标人须提供国网供应商不良行为查询截图'})])
    b = PackageTRM()
    trm = TRM(packages=[a, b], prenotice=[PrenoticeClause(content=
        '国网供应商不良行为处理有效期内不得参与。投标人须提供信用中国查询截图。')])
    assert grid_credit_requirement(trm, b)['mode'] == 'condition'
    assert grid_credit_requirement(trm, a)['mode'] == 'evidence'
    assert grid_credit_requirement(TRM(), b)['mode'] == 'unspecified'
    trm.prenotice.append(PrenoticeClause(content='国网供应商不良行为查询截图无需提交。'))
    assert grid_credit_requirement(trm, a)['mode'] == 'conflict'


def test_grid_missing_attachment_marker_only_for_explicit_submission_requirement(monkeypatch):
    from jb_parser.trm import PrenoticeClause
    calls = []
    def insert(ctx, doc, anchor, ids, label):
        calls.append(label)
        return anchor
    monkeypatch.setattr('jb_docgen.evidence.insert_scans', insert)
    ctx = SimpleNamespace(profile=CompanyProfile(name='企业'), trm=TRM(), pkg=PackageTRM())
    d = Document()
    insert_credit(ctx, d, d.add_paragraph('查询材料')._p)
    assert '国网供应商不良行为' not in calls
    calls.clear()
    ctx.trm.prenotice = [PrenoticeClause(content='投标人必须提交国网供应商不良行为查询截图。')]
    insert_credit(ctx, d, d.add_paragraph('查询材料')._p)
    assert '国网供应商不良行为' in calls
