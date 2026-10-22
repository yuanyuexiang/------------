"""需求表有多个分标的同号包时，不得串行；福建双层表头和万元限价保留口径。"""
from pathlib import Path

import docx
import openpyxl
import pytest
from jb_parser import parse
from jb_parser.extract import extract_service_demand
from jb_parser.pipeline import parse_package_dir
from jb_parser.trm import TRM


def demand(path):
    w = openpyxl.Workbook()
    s = w.active
    s.append(['分标编号', '分标名称', '包号', '包名称（采购项目名称）', '项目概况',
              '行报价最高限价（含税/万元）', '应答人专用资格要求', '', '',
              '技术权重（%）', '商务权重（%）', '价格权重（%）', '技术规范ID号'])
    s.append(['', '', '', '', '', '', '资质要求', '业绩要求', '人员要求'])
    s.append(['123456-9013006-2030', '其他运维', '包16', '错误项目', '错误范围', 20,
              '错误资质', '错误业绩', '错误人员', 50, 10, 40, 'OTHER'])
    s.append(['123456-9012008-2037', '科技项目-科技项目', '包16', '研究项目', '研究范围', 170,
              '/', '研究业绩', '研究人员', 60, 10, 30, 'B00D-300000000-00001'])
    w.save(path)


def test_ambiguous_same_package_number_requires_identity(tmp_path):
    path = tmp_path / '需求一览表.xlsx'
    demand(path)
    assert extract_service_demand(str(path), '包16') is None
    row = extract_service_demand(str(path), '包16', sub_no='123456-9012008-2037')
    assert row['企业业绩要求'] == '研究业绩'
    assert row['主要人员要求'] == '研究人员'
    assert extract_service_demand(str(path), '包16', sub_no='123456-9999999-9999') is None


def test_ecp_path_and_fujian_headers_populate_the_correct_package(tmp_path):
    folder = tmp_path / '123456-9012008-2037、科技项目-科技项目、包16'
    folder.mkdir()
    main = folder / '演示采购文件.docx'
    d = docx.Document()
    d.add_paragraph('模拟采购文件')
    d.save(main)
    sheet = tmp_path / '采购公告附件-需求一览表（123456）.xlsx'
    demand(sheet)
    trm = TRM()
    parse_package_dir(str(tmp_path), [str(main.relative_to(tmp_path)), sheet.name], trm, pkg_no='包16')
    pkg = trm.packages[0]
    assert (pkg.sub_no, pkg.sub_name, pkg.project_name) == ('123456-9012008-2037', '科技项目-科技项目', '研究项目')
    assert trm.batch_no == '123456'
    assert pkg.scope == '研究范围' and pkg.max_price == '170万元（含税）'
    assert (pkg.scoring_ref.weight_tech, pkg.scoring_ref.weight_biz, pkg.scoring_ref.weight_price) == (60, 10, 30)
    assert pkg.qualification[0].performance_req == '研究业绩'


SAMPLE = Path(__file__).resolve().parents[1] / '服务/2/国网福建电力2026年第三次服务类竞争性谈判采购_采购文件包.zip'


@pytest.mark.skipif(not SAMPLE.exists(), reason='本地福建采购样本不可用')
def test_fujian_real_packages_are_not_assigned_another_subdivision():
    trm = parse(str(SAMPLE))
    assert trm.batch_no == '132673'
    assert {p.pkg_no for p in trm.packages} == {'包14', '包15', '包16'}
    for pkg in trm.packages:
        assert pkg.sub_no == '132673-9012008-2037'
        assert pkg.sub_name == '科技项目-科技项目'
        assert pkg.project_name and pkg.scope
        assert (pkg.scoring_ref.weight_tech, pkg.scoring_ref.weight_biz, pkg.scoring_ref.weight_price) == (60, 10, 30)
    pkg16 = next(p for p in trm.packages if p.pkg_no == '包16')
    assert pkg16.max_price == '170万元（含税）'
    assert pkg16.qualification[0].other_reqs['技术规范书ID'] == 'B00D-300013753-00364'
