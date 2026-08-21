"""模拟评分 + 导出加固（合成数据）。"""
import docx
from jb_agents.scorer import score_package
from jb_docgen.export import clean_metadata, export
from jb_kb.models import Certificate, CompanyProfile, Performance, Person
from jb_parser.trm import TRM, PackageTRM, ScoringItem, ScoringRef, ScoringTemplate


def _trm():
    tech = ScoringTemplate(name="JS-FWTY 服务技术详评模板（通用）", kind="tech", items=[
        ScoringItem(element="1.总体评价（6-10 分）", content="优：9-10 分 良：7-8 分 一般：6 分"),
        ScoringItem(element="3．业绩情况（12-20 分）", content="具有 1 项合格业绩的得 12 分，每多提供一项合格业绩得 2 分，最高 20 分"),
        ScoringItem(element="6.专利数量（0-2 分）", content="每提供 1 项加 1 分，上限不超过 2 分"),
        ScoringItem(element="7．绩效评价（1-5 分）", content="A：5分 …… 未参加绩效评价：4分"),
    ])
    biz = ScoringTemplate(name="FWSW01：服务类通用商务详评细则", kind="biz", items=[
        ScoringItem(group="2.绿色低碳评价（3 分）", element="绿电绿证（0-3 分）", content="取得证书得 3 分"),
        ScoringItem(group="3.科研创新评价", element="管理体系认证（3-5 分）", content="取得4项=5分；1-3项=4分"),
    ])
    pkg = PackageTRM(pkg_no="包1", scoring_ref=ScoringRef(tech_template="JS-FWTY", biz_template="FWSW01",
                                                          weight_tech=60, weight_biz=10, weight_price=30))
    return TRM(scoring_templates=[tech, biz], packages=[pkg]), pkg


def test_hard_indicators_counted():
    trm, pkg = _trm()
    prof = CompanyProfile(name="T", performances=[Performance(project=f"P{i}", evidence=["合同", "发票"], buyer_is_end_user=True, signed_date="2025-01-01") for i in range(3)],
                          certificates=[Certificate(name="质量管理体系认证证书", cert_type="体系认证"), Certificate(name="环境管理体系认证证书", cert_type="体系认证")],
                          personnel=[Person(name="A", title="高级工程师")])
    rep = score_package(trm, pkg, prof, use_llm=False)
    by = {i.element: i for i in rep.items}
    assert by["3．业绩情况（12-20 分）"].predicted == 16          # 12 + 2*2
    assert by["6.专利数量（0-2 分）"].predicted == 0
    assert by["7．绩效评价（1-5 分）"].predicted == 4
    assert by["绿电绿证（0-3 分）"].predicted == 0
    assert by["管理体系认证（3-5 分）"].predicted == 4.0           # 2/4 → 3 + 2*0.5
    assert by["1.总体评价（6-10 分）"].predicted is None           # 软指标无 LLM → 区间
    assert rep.tech_max == 37 and rep.biz_max == 8
    assert rep.heatmap()[0]["loss"] >= rep.heatmap()[-1]["loss"]


def test_export_blocks_on_todo_then_cleans(tmp_path):
    p = tmp_path / "a.docx"
    d = docx.Document()
    d.core_properties.author = "张三"
    d.add_paragraph("正文 【待补充：营业执照扫描件】")
    d.save(str(p))
    r = export(str(p), want_pdf=False)
    assert not r.ok and r.blocked_by == ["营业执照扫描件"]
    d = docx.Document(str(p))
    d.paragraphs[0].text = "正文 已补齐"
    d.save(str(p))
    r = export(str(p), want_pdf=False)
    assert r.ok and docx.Document(str(p)).core_properties.author == ""


def test_clean_metadata(tmp_path):
    p = tmp_path / "b.docx"
    d = docx.Document()
    d.core_properties.author = "某公司"
    d.core_properties.comments = "x"
    d.save(str(p))
    clean_metadata(str(p))
    cp = docx.Document(str(p)).core_properties
    assert cp.author == "" and cp.comments == ""
