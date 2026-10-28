"""复现空团队、合并引用、重复人员缺项及附件检索问题。"""
from types import SimpleNamespace

from docx import Document
from jb_agents.draft_review import review_batch, semantic_issues
from jb_agents.preparation import validate_recommendations
from jb_agents.writer import _check_numbers, _facts_block
from jb_docgen.gaps import consolidate_team_gaps
from jb_docgen.placeholders import scan_docx
from jb_kb.models import CompanyProfile, Person
from jb_kb.text_evidence import read_evidence


def test_merged_citations_are_normalized_but_unknown_sources_are_blocked():
    sources = {'req-package': '总服务期4个月', 'req-spec-0': '运营3个月'}
    text, issues = _check_numbers('招标要求服务期4个月，运营3个月[来源:req-package；来源:req-spec-0]。', sources)
    assert not issues and '待补充' not in text
    text, issues = _check_numbers('招标要求9人[来源:req-package；来源:fake]。', sources)
    assert '无效来源:fake' in issues and '9' in issues
    assert '待补充' in text


def test_requirement_numbers_still_cannot_prove_enterprise_capability():
    _, issues = _check_numbers('我方已具备9人团队[来源:req-a；来源:req-b]。', {'req-a': '要求9人', 'req-b': '要求团队'})
    assert '9' in issues


def test_team_consolidation_preserves_locations_other_issues_and_format(tmp_path):
    path = tmp_path / 'draft.docx'
    d = Document()
    p = d.add_paragraph()
    p.add_run('正文：').bold = True
    p.add_run('【待补充：拟投入')
    p.add_run('人员名单】')
    p.add_run('，继续说明【待补充：项目负责人工作年限】。')
    d.add_table(rows=1, cols=1).cell(0, 0).text = '【待补充：拟投入人员名单】'
    d.add_paragraph('【待补充：合同签约日期】')
    d.save(path)
    details = consolidate_team_gaps(str(path))
    assert len(details) == 3 and len({x['paragraph'] for x in details}) == 2
    assert len(scan_docx(str(path))) == 2
    result = Document(path)
    assert result.paragraphs[0].runs[0].bold
    assert result.tables[0].cell(0, 0).text == '（见“本次项目团队与资格核对”）'
    assert any('项目负责人工作年限' in x.text for x in result.paragraphs)
    assert any('本包团队名单' in x for x in scan_docx(str(path)))


def test_authority_is_not_team_and_candidate_facts_are_available():
    person = Person(id='p1', name='张某', education='本科', major='计算机科学与技术')
    context = {'authorization': {'name': '张某'}, 'personnel_candidates': [person.model_dump()]}
    cp = CompanyProfile(name='公司', bid_context=context)
    facts, _ = _facts_block(cp)
    assert 'fact-candidate-p1' in facts and '本科' in facts and '未承诺本包投入' in facts
    assert semantic_issues('拟将张某作为项目负责人候选。', context)
    assert semantic_issues('团队到岗承诺已确认。', context)
    assert semantic_issues('【待补充：张某学历、工作年限】', context)


def test_failed_role_repair_remains_blocking():
    class Offline:
        def chat_json(self, *a, **kw):
            raise RuntimeError('offline')
    context = {'authorization': {'name': '张某'}}
    output, checks = review_batch({'a': '拟将张某作为项目负责人候选。'}, {}, '', context, Offline())
    assert 'AI角色或团队状态一致性需修复' in output['a']
    assert checks['a']['semantic']


def test_repair_cannot_remove_real_missing_facts_to_pass():
    class Removing:
        def chat_json(self, *a, **kw):
            return {'a': '招标要求服务期四个月，本公司将依据项目安排组织实施并提交所有成果资料。'}
    output, checks = review_batch({'a': '招标要求4个月。【待补充：企业开户证明】'}, {}, '', {}, Removing())
    assert '企业开户证明' in output['a'] and checks['a']['violations']


def test_recommendation_requires_same_person_quote_and_known_role():
    cp = CompanyProfile(name='公司', personnel=[Person(id='p', name='甲')])
    evidence = [{'id': 'e', 'title': '毕业证', 'person_ids': ['p'], 'text': '甲本科计算机专业', 'attachment_id': 'a'}]
    entry = {'person_id': 'p', 'role': '项目负责人', 'source_id': 'e', 'quote': '本科计算机专业'}
    result = validate_recommendations({'recommendations': [entry]}, evidence, cp, [{'role': '项目负责人'}])
    assert len(result) == 1
    assert not validate_recommendations({'recommendations': [{**entry, 'person_id': 'foreign'}]}, evidence, cp, [{'role': '项目负责人'}])
    assert not validate_recommendations({'recommendations': [entry]}, evidence, cp, [{'role': '技术人员'}])
    assert not validate_recommendations({'recommendations': [{**entry, 'quote': '甲博士电力专业'}]}, evidence, cp, [{'role': '项目负责人'}])


def test_ocr_runs_for_scanned_page_and_cache_avoids_duplicate_work(tmp_path, monkeypatch):
    source = tmp_path / 'degree.pdf'
    source.write_bytes(b'test pdf fixture')
    calls = []
    def run(cmd, **kwargs):
        calls.append(cmd[0])
        return SimpleNamespace(returncode=0, stdout='\f' if cmd[0] == 'pdftotext' else '毕业证书：计算机专业，本科学历')
    monkeypatch.setattr('jb_kb.text_evidence.subprocess.run', run)
    result = read_evidence(str(source))
    assert result['pages'][0]['method'] == 'ocr'
    assert result['pages'][0]['page'] == 1 and '本科' in result['pages'][0]['text']
    assert calls == ['pdftotext', 'pdftoppm', 'tesseract']
    assert read_evidence(str(source)) == result and len(calls) == 3
    source.write_bytes(b'changed pdf fixture')
    read_evidence(str(source))
    assert len(calls) == 6
