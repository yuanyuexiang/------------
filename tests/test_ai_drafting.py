"""AI 写作编排回归：合成样本，不依赖外部模型与客户资料。"""
from docx import Document
from jb_agents.draft_context import requirement_sources, section_requirements
from jb_agents.draft_report import confirmation_items
from jb_agents.writer import _check_numbers, draft_package
from jb_kb.models import Boilerplate, CompanyProfile, Performance
from jb_parser.trm import TRM, PackageTRM, SpecDoc


def test_original_spec_is_included_and_other_package_is_excluded(tmp_path):
    path = tmp_path / 'spec.docx'
    doc = Document()
    doc.add_paragraph('量子密钥研究技术路线与验证方法')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '交付研究报告和验证记录'
    doc.save(path)
    pkg = PackageTRM(pkg_no='包1', scope='开展量子密钥研究',
                     spec_docs=[SpecDoc(spec_id='S', title='研究任务书')])
    trm = TRM(packages=[pkg, PackageTRM(scope='不相关项目保密原文')])
    sources, warnings = requirement_sources(trm, pkg, {'S': str(path)})
    selected = section_requirements(sources, '量子密钥研究交付报告')
    text = str(selected)
    assert not warnings
    assert '验证方法' in text and '验证记录' in text
    assert '不相关项目' not in text


def test_requirement_numbers_do_not_prove_company_capability():
    sources = {'req-param-0': '招标要求吞吐率1000次'}
    checked, violations = _check_numbers('我司现有设备达到1000次。[来源:req-param-0]', sources)
    assert violations and '达到1000' not in checked
    text = '招标要求吞吐率1000次。[来源:req-param-0]'
    assert _check_numbers(text, sources) == (text, [])


def test_drafting_uses_original_and_approved_facts(monkeypatch, tmp_path):
    import jb_llm

    prompts = []
    monkeypatch.setattr(jb_llm, 'available', lambda: True)

    def chat(prompt, **kwargs):
        prompts.append(prompt)
        return '拟实施方案，供投标负责人确认。拟组织需求调研并交付验证报告。'

    monkeypatch.setattr(jb_llm, 'chat', chat)
    doc = Document()
    doc.add_paragraph('专属研究任务：量子网络验证与报告')
    path = tmp_path / 'spec.docx'
    doc.save(path)
    profile = CompanyProfile(name='测试企业', performances=[
        Performance(project='已审核案例', amount_wan=25, status='approved'),
        Performance(project='未审核案例', amount_wan=99, status='draft'),
    ], boilerplates=[Boilerplate(topic='质量', text='未审核话术绝不进入模型', approved=False)])
    pkg = PackageTRM(scope='量子网络研究', spec_docs=[SpecDoc(spec_id='S')])
    progress = []
    draft = draft_package(TRM(packages=[pkg]), pkg, profile, spec_paths={'S': str(path)},
                          on_progress=lambda *args: progress.append(args))
    assert len(progress) == len(draft.sections) == len(prompts)
    assert all(section.by_llm and not section.todos for section in draft.sections)
    assert any('专属研究任务' in prompt for prompt in prompts)
    assert all('已审核案例' in prompt and '25' in prompt for prompt in prompts)
    assert all('未审核案例' not in prompt and '未审核话术' not in prompt for prompt in prompts)
    assert all('不得因缺少现成方案而要求用户撰写正文' in prompt for prompt in prompts)
    items = confirmation_items(draft, {})
    assert all(item.category == '拟实施方案待确认' for item in items)


def test_failed_or_empty_model_output_is_not_success(monkeypatch):
    import jb_llm

    monkeypatch.setattr(jb_llm, 'available', lambda: True)
    monkeypatch.setattr(jb_llm, 'chat', lambda *args, **kwargs: '  ')
    draft = draft_package(TRM(), PackageTRM(scope='研究服务'), CompanyProfile(name='测试'))
    assert draft.warnings
    assert all(not section.by_llm and section.todos for section in draft.sections)
    items = confirmation_items(draft, {'技术.docx': draft.sections[0].todos})
    assert items and all(item.category == 'AI 起草需重试' for item in items)


def test_unconfigured_model_is_explicit(monkeypatch):
    import jb_llm

    monkeypatch.setattr(jb_llm, 'available', lambda: False)
    draft = draft_package(TRM(), PackageTRM(scope='服务'), CompanyProfile(name='测试'))
    assert any('模型未配置' in warning for warning in draft.warnings)
    assert all(not section.by_llm and section.todos for section in draft.sections)


def test_plan_response_is_not_mislabeled_as_missing_company_fact():
    items = confirmation_items(None, {'技术.docx': ['对「研究路线」的具体响应', '合同金额']})
    assert items[0].category == '方案响应待完善'
    assert items[1].category == '企业事实待补充'


def test_service_response_fills_exact_requirement_only(monkeypatch, tmp_path):
    import jb_llm
    from jb_docgen.spec_response import _respond_items

    monkeypatch.setattr(jb_llm, 'available', lambda: True)
    monkeypatch.setattr(jb_llm, 'chat', lambda *args, **kwargs: '拟组织需求调研。')
    monkeypatch.setattr(jb_llm, 'chat_json', lambda *args, **kwargs: {
        'req-service-0-0': '拟实施响应（供投标负责人确认）：拟开展对照实验并提交验证记录。',
    })
    doc = Document()
    doc.add_heading('1 研究任务', level=1)
    doc.add_paragraph('（1）提交算法验证报告')
    path = tmp_path / 'service.docx'
    doc.save(path)
    pkg = PackageTRM(sub_name='科技项目-科技项目', scope='算法研究',
                     spec_docs=[SpecDoc(spec_id='S')])
    draft = draft_package(TRM(), pkg, CompanyProfile(name='测试'), spec_paths={'S': str(path)})
    responses = _respond_items(['（1）提交算法验证报告', '（2）其他要求'], None,
                               draft.response_plans['S'])
    assert '对照实验' in responses[0] and '待补充' not in responses[0]
    assert '待补充' in responses[1]


def test_material_requirements_do_not_use_service_plans(monkeypatch, tmp_path):
    import jb_llm
    from jb_parser.trm import Material

    monkeypatch.setattr(jb_llm, 'available', lambda: True)
    monkeypatch.setattr(jb_llm, 'chat', lambda *args, **kwargs: '拟核对产品证据。')

    def unexpected(*args, **kwargs):
        raise AssertionError('物资参数不得生成服务方案替代证据')

    monkeypatch.setattr(jb_llm, 'chat_json', unexpected)
    pkg = PackageTRM(sub_name='设备设计', materials=[Material(desc='设备')])
    draft = draft_package(TRM(), pkg, CompanyProfile(name='测试'))
    assert not draft.response_plans


def test_json_client_forwards_generation_limits(monkeypatch):
    import jb_llm

    monkeypatch.setattr(jb_llm, 'available', lambda: True)
    monkeypatch.setenv('LLM_API_KEY', 'test-only')
    payloads = []

    class Response:
        def raise_for_status(self):
            pass

        def json(self):
            return {'choices': [{'message': {'content': '{"ok": true}'}}]}

    def post(*args, **kwargs):
        payloads.append(kwargs)
        return Response()

    monkeypatch.setattr(jb_llm.httpx, 'post', post)
    assert jb_llm.chat_json('test', timeout=120, temperature=0.3, max_tokens=6500) == {'ok': True}
    assert payloads[0]['json']['max_tokens'] == 6500
    assert payloads[0]['timeout'] == 120


def test_requirement_quote_does_not_cover_company_claim_in_same_sentence():
    text = '招标要求1000次，我司设备已达到1000次。[来源:req-param-0]'
    checked, violations = _check_numbers(text, {'req-param-0': '指标1000次'})
    assert violations and '达到1000' not in checked


def test_heading_numbers_and_date_format_do_not_create_false_gaps():
    text = '### 1.1 工作安排\n投标截止为2026年5月25日9:30。[来源:req-terms]'
    sources = {'req-terms': '投标截止2026-05-25 09:30'}
    assert _check_numbers(text, sources) == (text, [])
    checked, violations = _check_numbers('### 1.1 保证1.1小时', {})
    assert violations and '保证1.1小时' not in checked


def test_requirement_cannot_support_unqualified_guarantee():
    checked, violations = _check_numbers('保证1000次。[来源:req-param-0]', {'req-param-0': '指标1000次'})
    assert violations and '保证1000' not in checked


def test_numeric_guard_keeps_existing_placeholder_intact():
    checked, violations = _check_numbers('1. 保证1小时；【待补充：表1的完整内容】', {})
    assert violations and checked.startswith('1. 保证【待补充')
    assert '【待补充：表1的完整内容】' in checked


def test_scoring_outline_reuses_fact_tables_and_avoids_duplicate_overview():
    from jb_agents.writer import build_outline
    from jb_parser.trm import ScoringItem, ScoringTemplate

    trm = TRM(scoring_templates=[ScoringTemplate(kind='tech', items=[
        ScoringItem(element='投标文件技术部分总体评价'), ScoringItem(element='业绩情况'),
        ScoringItem(element='服务方案'), ScoringItem(element='工作进度及保证措施'),
    ])])
    titles = [section.title for section in build_outline(trm, PackageTRM(scope='研究服务'))]
    assert '业绩情况' not in titles
    assert '项目理解与需求分析' not in titles and '招标范围逐条响应' not in titles
    assert '工作进度及保证措施' in titles and '质量保证体系' in titles


def test_bold_parenthesized_outline_is_not_a_business_number():
    text = '**（1）研究方法**\n（2）验证方法'
    assert _check_numbers(text, {}) == (text, [])
    for text in ('（20）人参与', '3.14万元', '**（1）拟投入1人**'):
        checked, violations = _check_numbers(text, {})
        assert violations and '数值需核实来源' in checked


def test_fullwidth_outline_and_project_items_are_formatting():
    text = '#### 1．项目理解\n（1）项目管理\n**（2）人员配置**'
    assert _check_numbers(text, {}) == (text, [])
    checked, violations = _check_numbers('（20）人参与；保证1小时', {})
    assert violations and '数值需核实来源' in checked


def test_citation_guard_is_separate_from_missing_company_material():
    items = confirmation_items(None, {'技术.docx': ['数值需核实来源', '引用来源需核实', '合同金额']})
    assert [item.category for item in items] == ['引用依据待核对', '引用依据待核对', '企业事实待补充']


def test_markdown_table_ordinal_is_formatting_but_values_need_evidence():
    text = '| 序号 | 安排 |\n| --- | --- |\n| 1 | 拟开展调研 |'
    assert _check_numbers(text, {}) == (text, [])
    text += '\n| 2 | 保证2小时 |'
    checked, violations = _check_numbers(text, {})
    assert violations and '| 2 |' in checked and '保证2小时' not in checked


def test_docx_run_boundaries_do_not_split_a_year(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run('计划在202')
    paragraph.add_run('7年9月完成')
    path = tmp_path / 'runs.docx'
    doc.save(path)
    pkg = PackageTRM(spec_docs=[SpecDoc(spec_id='S')])
    sources, _ = requirement_sources(TRM(), pkg, {'S': str(path)})
    assert '2027年9月' in str(sources)


def test_schedule_heading_dates_get_only_unique_existing_schedule_citation():
    from jb_agents.draft_context import attribute_schedule_dates

    sources = {'req-spec-0-1': '预计时间段：2026年10月～2026年12月；2027年1月～2027年3月'}
    text = '##### 第二阶段（2026年10月～12月）'
    attributed = attribute_schedule_dates(text, sources)
    assert attributed.endswith('[来源:req-spec-0-1]')
    assert _check_numbers(attributed, sources) == (attributed, [])
    for unsupported in ('##### 第三阶段（2028年1月～3月）', '本公司成立于2026年10月。', '##### 第三阶段（2026年10月～2027年3月）',
                        '| 项目 | 签约日期 |\n| 历史业绩 | 2026年10月 |'):
        assert attribute_schedule_dates(unsupported, sources) == unsupported
    duplicate = {**sources, 'req-spec-1-1': sources['req-spec-0-1']}
    assert attribute_schedule_dates(text, duplicate) == text


def test_table_row_citation_covers_its_cells_but_not_next_row():
    text = '| 第一阶段 | 2026年10月 | 拟编制交底书2项；组织评审 | [来源:req-spec-0-1] |'
    sources = {'req-spec-0-1': '预计时间段2026年10月，要求编制交底书2项'}
    assert _check_numbers(text, sources) == (text, [])
    checked, violations = _check_numbers(text + '\n| 第二阶段 | 保证999小时 |', sources)
    assert violations and '999小时' not in checked
