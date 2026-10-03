"""母版切片装配·真实样本回归（陕西物资 / 江西物资，缺样本自动 skip）。

断言母版模式生效、国网原表格被原位填充、母版其他章节的媒体没有残留。
"""
import os
import zipfile

import docx
import pytest
from jb_docgen import generate
from jb_kb.models import CompanyProfile, FinancialYear, Performance
from jb_parser import main_doc_for_package, parse

WS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHAANXI_044 = os.path.join(
    WS, "物资/1/国网陕西省电力有限公司2026年第三次物资集中招标采购项目_招标文件包/"
        "分标044接入设备/包1_完整招标文件_1813604751495232.zip")
JIANGXI_PKG = os.path.join(WS, "物资/通信监控系统/包1_完整招标文件_102421217810232908.zip")


def _profile():
    return CompanyProfile(name="样本测试企业", credit_code="91320000MA1234567X", address="南京市江北新区",
                          legal_person="张三", authorized_rep="李四", authorized_rep_title="销售经理",
                          performances=[Performance(project="某变电站通信改造", buyer="国网某公司", amount_wan=88,
                                                    signed_date="2025-03-01", evidence=["合同", "发票"])],
                          financials=[FinancialYear(year="2024", revenue_wan=5000, asset_wan=8000, liability_ratio="40%")])


def _full_text(path):
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            parts.append(" | ".join(c.text for c in r.cells))
    return "\n".join(parts), d


def _run(zip_path, tmp_path):
    trm = parse(zip_path)
    pkg = trm.packages[0]
    master = main_doc_for_package(zip_path, pkg.sub_no, pkg.pkg_no, str(tmp_path / "src"))
    assert master and master.endswith(".docx")
    res = generate(trm, pkg, _profile(), str(tmp_path / "out"), master_docx=master)
    assert res.mode == {"commercial": "master", "technical": "master"}, res.notes
    return trm, pkg, res


@pytest.mark.skipif(not os.path.exists(SHAANXI_044), reason="样本缺失")
def test_shaanxi_master(tmp_path):
    trm, pkg, res = _run(SHAANXI_044, tmp_path)
    text, d = _full_text(res.commercial_path)
    # 授权委托书原文填空（跨 run）+ 投标函标签行
    assert "样本测试企业（投标人名称）为中华人民共和国合法企业，法定地址：南京市江北新区。" in text
    assert f"特授权李四代表我方全权办理{trm.batch_no}（招标编号）（{pkg.sub_no} {pkg.sub_name}）（{pkg.pkg_no}）" in text
    assert "投标人名称：样本测试企业" in text and "授权代表职务：销售经理" in text
    # 国网原表格原位填充：基本情况表 / 人员关系表 / 财务表 / 商务偏差表
    assert "投标人名称 | 样本测试企业" in text and "统一社会信用代码 | 91320000MA1234567X" in text
    assert "张三 | 【待补充：性别】" in text and "2024年 | XX年 | XX年 | 平均" in text
    assert "资产总额 | 8000万元" in text and "1 | / | / | / | 无偏差" in text
    # 第六章以外的章节与投标工具指引行不在成品里
    heads = {p.text.strip() for p in d.paragraphs}
    assert not any(h.startswith(("第一章", "第二章", "第三章", "第四章", "第五章", "第六章")) for h in heads)
    assert "评标办法（综合评估法）" not in heads and "上传投标工具路径" not in text

    text, d = _full_text(res.technical_path)
    assert "某变电站通信改造 | 88万元" in text and "技术偏差表" in text
    assert "商务偏差表" not in text
    for path in (res.commercial_path, res.technical_path):
        with zipfile.ZipFile(path) as z:
            assert not [n for n in z.namelist() if n.startswith("word/media/")]


@pytest.mark.skipif(not os.path.exists(JIANGXI_PKG), reason="样本缺失")
def test_jiangxi_master(tmp_path):
    """主文件名无括号编号（…招标采购项目招标文件.docx）：classify 兼容分支 + 母版切片。"""
    _, _, res = _run(JIANGXI_PKG, tmp_path)
    text, _ = _full_text(res.commercial_path)
    assert "样本测试企业（投标人名称）为中华人民共和国合法企业" in text
    assert "投标人名称 | 样本测试企业" in text
