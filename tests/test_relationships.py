"""本包事实确认、条款选择及原版式保护。全部为合成事实。"""
import docx
import pytest
from docx.shared import Pt
from jb_docgen.builder import GenContext
from jb_docgen.relationships import HEADERS, apply_master, source_options
from jb_docgen.workflow import relationship_snapshot, relationship_state, save_relationship
from jb_kb.models import CompanyProfile
from jb_parser.docx_utils import Chapter
from jb_parser.extract import extract_format_blocks
from jb_parser.trm import TRM, PackageTRM
from jb_store import Project, Task, session


@pytest.fixture()
def template():
    d = docx.Document()
    d.add_paragraph('2.10 应答人与国家电网公司系统人员关系说明')
    p = d.add_paragraph('【1.不存在相关情形的适用】截至首次应答截止日，不存在相关情形。')
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(12)
    d.add_paragraph('【2.存在相关情形的适用】截至首次应答截止日，存在相关情形。具体情况如下。')
    d.add_paragraph('《本企业员工与国家电网公司系统人员关系说明表》')
    t = d.add_table(rows=3, cols=13)
    for c, h in zip(t.rows[1].cells, HEADERS):
        c.text = h
    d.add_paragraph('1.本表的人员只统计企业法人、出资人、高管涉及的具体情形。')
    d.add_paragraph('3.应答人基本情况表')
    d.add_paragraph('法定代表人（签字）：')
    return d


def context(state=None):
    return GenContext(TRM(), PackageTRM(), CompanyProfile(name='测试企业', legal_person='不能自动列为关系人',
                      sgcc_relationship='无'), relationship_confirmation=state)


def test_none_preserves_exact_clause_style_table_and_notes(template):
    first = template.paragraphs[1]._p.xml
    ctx = context({'status': 'none'})
    apply_master(template, ctx)
    assert template.paragraphs[1]._p.xml == first
    assert not any('【2.' in p.text for p in template.paragraphs)
    assert template.tables[0].rows[2].cells[0].text == '不适用'
    assert template.paragraphs[-1].text == '法定代表人（签字）：'
    assert any('只统计' in p.text for p in template.paragraphs)
    assert not ctx.todos


def test_unknown_never_uses_global_none(template):
    ctx = context()
    apply_master(template, ctx)
    assert ctx.todos == ['本包国网系统人员关系情况待企业确认']
    assert template.tables[0].rows[2].cells[0].text == ''
    assert sum('【1.' in p.text or '【2.' in p.text for p in template.paragraphs) == 2


def test_exists_fills_both_sides_and_missing_fields(template):
    row = ['甲', '男', '测试身份A', '股东', '2020', '配偶', '乙', '女', '测试身份B', '测试单位', '职员', '在职', '']
    ctx = context({'status': 'exists', 'people': [row, [''] * 13]})
    apply_master(template, ctx)
    assert not any('【1.' in p.text for p in template.paragraphs)
    assert any('【2.' in p.text for p in template.paragraphs)
    t = template.tables[0]
    assert len(t.rows) == 4
    assert t.rows[2].cells[0].text == '甲' and t.rows[2].cells[6].text == '乙'
    assert t.rows[2].cells[12].text == '不适用'
    assert '待补充' in t.rows[3].cells[0].text
    assert ctx.todos


def test_ambiguous_or_unsupported_source_blocks_completion(template):
    template.add_paragraph('应答人与国家电网公司系统人员关系说明')
    template.add_paragraph('【1.不存在相关情形的适用】不同原文。')
    template.add_paragraph('【2.存在相关情形的适用】不同原文。')
    ctx = context({'status': 'none'})
    apply_master(template, ctx)
    assert ctx.todos and '原文' in ctx.todos[0]
    assert sum('【2.' in p.text for p in template.paragraphs) == 2


def test_service_chapter_five_numbered_heading(template):
    blocks = extract_format_blocks(template, [Chapter(5, '应答文件格式', 0, len(template.paragraphs))])
    options = source_options(TRM(format_blocks=blocks))
    assert len(options) == 2 and options[0].startswith('【1.')


def test_confirmation_scope_and_stale():
    p = Project(trm=TRM(packages=[PackageTRM(), PackageTRM()]).model_dump(), zip_path='sample.zip')
    state = relationship_state(p, 0, '测试企业')
    save_relationship(p, 0, '测试企业', {'source': state['source'], 'status': 'none'}, 'operator')
    assert relationship_snapshot(p, 0, '测试企业')['confirmed_by'] == 'operator'
    assert relationship_snapshot(p, 1, '测试企业') is None
    assert relationship_snapshot(p, 0, '另一企业') is None
    p.deadline = '2026-12-31'
    assert relationship_state(p, 0, '测试企业')['stale']
    assert relationship_snapshot(p, 0, '测试企业') is None
    with pytest.raises(ValueError, match='刷新'):
        save_relationship(p, 0, '测试企业', {'source': state['source'], 'status': 'none'}, 'operator')


def test_api_confirmation_validation_and_running_lock(client):
    client.put('/api/profiles/测试企业', json={'name': '测试企业'})
    with session() as s:
        p = Project(filename='sample.zip', zip_path='sample.zip', trm=TRM(packages=[PackageTRM()]).model_dump())
        s.add(p)
        s.flush()
        pid = p.id
    url = f'/api/projects/{pid}/relationship?profile=测试企业'
    state = client.get(url).json()
    assert state['confirmation'] is None
    assert client.put(url, json={'source': state['source'], 'status': 'none', 'confirmed_by': 'fake'}).status_code == 400
    assert client.put(url, json={'source': state['source'], 'status': 'exists', 'people': [['bad']]}).status_code == 400
    assert client.put(url, json={'source': state['source'], 'status': 'none'}).status_code == 200
    assert client.get(url).json()['confirmation']['confirmed_by'] == 'admin'
    assert client.get(url + '&pkg_index=3').status_code == 400
    with session() as s:
        s.add(Task(project_id=pid, kind='generate', status='running'))
    assert client.put(url, json={'source': state['source'], 'status': 'exists'}).status_code == 409


def test_old_project_previews_original_without_reparse(template, tmp_path, monkeypatch):
    from types import SimpleNamespace

    from jb_docgen.workflow import _relationship_original
    source = tmp_path / 'old.zip'
    source.write_bytes(b'synthetic fixture')
    monkeypatch.setattr('jb_parser.package_files', lambda *a, **k: SimpleNamespace(main_doc='original.docx'))
    monkeypatch.setattr('jb_docgen.master.carve', lambda *a: template)
    p = Project(trm=TRM(packages=[PackageTRM()]).model_dump(), zip_path=str(source))
    assert not relationship_state(p, 0, '测试企业')['options']
    preview = relationship_state(p, 0, '测试企业', preview=True)
    assert len(preview['options']) == 2
    assert '截至首次应答截止日' in preview['options'][0]
    assert preview['confirmation'] is None
    _relationship_original.cache_clear()
