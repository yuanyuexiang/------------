"""技术规范书逐条响应：规范书原文复制 + 参数表加响应列 + 功能条目匹配/待补充（合成规范书，不依赖样本）。"""
import docx
from docx.shared import Pt
from jb_docgen import generate
from jb_docgen import spec_response as SR
from jb_kb.models import CompanyProfile, Product
from jb_parser.trm import TRM, PackageTRM, SpecDoc


def _spec_doc(path: str) -> str:
    d = docx.Document()
    d.add_paragraph("技术规范")
    d.add_paragraph("1.调度专用接入设备（IAD）--8口")
    d.add_paragraph("调度专用接入设备可将用户端设备接入到分组交换网络中。")      # 纯叙述，不插响应
    d.add_paragraph("1.1技术参数")
    t = d.add_table(rows=4, cols=3)
    for r, vals in enumerate([["FXS及FXO端口", "端口数量", "支持8路"], ["FXS及FXO端口", "接口类型", "RJ45"],
                              ["录音端口", "录音端口", "最大支持单台设备32路"], ["控制接口", "控制接口", "1个异步串口"]]):
        for c, v in enumerate(vals):
            t.rows[r].cells[c].text = v
    d.add_paragraph("1.2功能特点")
    p = d.add_paragraph("(1)产品采用模块化设计，接口类型及录音功能可灵活配置，便于维护；")
    p.runs[0].font.size = Pt(10.5)
    d.add_paragraph("(2)标准1U高度；")
    d.add_paragraph("(3)支持两种传真模式：透传传真和T.38传真（速率可配置）；")
    d.add_paragraph("2.保证值表")
    t2 = d.add_table(rows=3, cols=4)
    for r, vals in enumerate([["序号", "名称", "项目需求值", "投标人保证值"], ["1", "端口数量", "不少于8路", ""], ["2", "录音端口", "≥16路", ""]]):
        for c, v in enumerate(vals):
            t2.rows[r].cells[c].text = v
    d.add_paragraph().add_run().add_break()
    d.save(path)
    return path


def _cover_doc(path: str) -> str:
    d = docx.Document()
    for line in ("国家电网公司集中规模招标采购", "国网某省电力有限公司", "招标文件", "（技术规范专用部分）", "工程概况", "网省公司：国网某省电力有限公司"):
        d.add_paragraph(line)
    d.save(path)
    return path


def _product():
    return Product(model="IAD-8", name="接入设备", category="接入设备",
                   params={"端口数量": "8路", "接口类型": "RJ45", "录音端口": "32路"},
                   features=["产品采用模块化设计，接口类型及录音功能可灵活配置", "支持透传传真和T.38传真"])


def _all_text(path):
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            parts.append(" | ".join(c.text for c in r.cells))
    return "\n".join(parts), d


def test_analyse_sections(tmp_path):
    src = docx.Document(_spec_doc(str(tmp_path / "s.docx")))
    secs = SR.analyse(src)
    titles = [(s.title, s.level, len(s.items), len(s.tables), s.has_children) for s in secs]
    assert titles[0] == ("1.调度专用接入设备（IAD）--8口", 1, 0, 0, True)
    assert titles[1] == ("1.1技术参数", 2, 0, 1, False)
    assert titles[2] == ("1.2功能特点", 2, 3, 0, False)
    assert not SR.is_cover_only(src, secs)
    cover = docx.Document(_cover_doc(str(tmp_path / "c.docx")))
    assert SR.is_cover_only(cover, SR.analyse(cover))


def test_generate_with_spec_responses(tmp_path):
    spec_path = _spec_doc(str(tmp_path / "spec.docx"))
    cover_path = _cover_doc(str(tmp_path / "cover.docx"))
    trm = TRM(batch_name="测试", batch_no="T-1")
    pkg = PackageTRM(sub_no="分标1", sub_name="接入设备", pkg_no="包1",
                     spec_docs=[SpecDoc(spec_id="G006-500130619-00001", title="技术规范", source="a/spec.docx"),
                                SpecDoc(spec_id="G006-500130619-00002", title="封面", source="b/cover.docx")])
    prof = CompanyProfile(name="测试企业", products=[_product()])
    res = generate(trm, pkg, prof, str(tmp_path / "out"),
                   spec_paths={"G006-500130619-00001": spec_path, "G006-500130619-00002": cover_path})
    text, d = _all_text(res.technical_path)
    assert "技术规范书逐条响应" in text and "技术规范书 G006-500130619-00001 技术规范" in text
    # 原文（含版式）被复制：功能条目保留 10.5 磅字号
    item = next(p for p in d.paragraphs if p.text.startswith("(1)产品采用模块化设计"))
    assert item.runs[0].font.size == Pt(10.5)
    # 三列参数表追加"投标人响应"列：数值匹配到产品库，弱匹配（控制接口≠接口类型）写待补充
    assert "FXS及FXO端口 | 端口数量 | 支持8路 | 8路" in text
    assert "录音端口 | 录音端口 | 最大支持单台设备32路 | 32路" in text
    assert "控制接口 | 控制接口 | 1个异步串口 | 【待补充：「控制接口 控制接口」的具体参数值】" in text
    # 已有"投标人保证值"列的表直接填该列
    assert "1 | 端口数量 | 不少于8路 | 8路" in text and "2 | 录音端口 | ≥16路 | 32路" in text
    # 功能条目：匹配到的给产品特性，匹配不到的待补充；绝不写"完全响应"
    assert "【投标人响应】" in text
    assert "（1）产品采用模块化设计，接口类型及录音功能可灵活配置" in text
    assert "（2）【待补充：对「标准1U高度；」的具体响应】" in text
    assert "（3）支持透传传真和T.38传真" in text
    assert "完全响应" not in text and "满足要求" not in text
    # 纯叙述小节不插响应（响应块只出现 1 次）；封面 docx 跳过并记入 notes
    assert text.count("【投标人响应】") == 1
    assert any("仅封面页" in n for n in res.notes)
    assert "国家电网公司集中规模招标采购" not in text


def test_no_product_all_todo(tmp_path):
    spec_path = _spec_doc(str(tmp_path / "spec.docx"))
    trm = TRM(batch_no="T-1")
    pkg = PackageTRM(pkg_no="包1", spec_docs=[SpecDoc(spec_id="X", title="规范", source="spec.docx")])
    res = generate(trm, pkg, CompanyProfile(name="测试企业"), str(tmp_path / "out"), spec_paths={"X": spec_path})
    text, _ = _all_text(res.technical_path)
    assert "FXS及FXO端口 | 端口数量 | 支持8路 | 【待补充：" in text
    assert "（1）【待补充：对「产品采用模块化设计" in text


def test_unnumbered_parameter_table_is_not_cover(tmp_path):
    src = docx.Document()
    src.add_paragraph("技术参数")
    table = src.add_table(rows=2, cols=3)
    for row, vals in zip(table.rows, [["名称", "要求值", "保证值"], ["端口数量", "8路", ""]]):
        for cell, val in zip(row.cells, vals):
            cell.text = val
    path = str(tmp_path / "spec.docx")
    src.save(path)
    assert not SR.is_cover_only(src, SR.analyse(src))
    dst = docx.Document()
    anchor = dst.add_paragraph("规范书")._p
    _, count, skipped = SR.append_spec(dst, anchor, path, None, "技术规范")
    assert not skipped and count == 1
    assert "【待补充：" in dst.tables[0].rows[1].cells[2].text


def test_schedule_table_is_preserved_without_fake_product_response():
    d = docx.Document()
    t = d.add_table(rows=2, cols=3)
    for row, vals in zip(t.rows, [["序号", "预计时间段", "研究内容"], ["1", "2026年10月至12月", "调研与方案设计"]]):
        for cell, val in zip(row.cells, vals):
            cell.text = val
    assert SR.respond_table(t, None) == 0
    assert len(t.columns) == 3
    assert t.cell(1, 2).text == "调研与方案设计"


def test_service_plan_is_inserted_but_does_not_fill_product_guarantee(tmp_path):
    d = docx.Document()
    d.add_heading('1 研究任务', level=1)
    item = '（1）提交研究报告'
    d.add_paragraph(item)
    d.add_heading('2 性能指标', level=1)
    table = d.add_table(rows=2, cols=3)
    for row, values in zip(table.rows, [['名称', '要求值', '保证值'], ['吞吐率', '1000次', '']]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    path = tmp_path / 'service.docx'
    d.save(path)
    pkg = PackageTRM(pkg_no='包1', spec_docs=[SpecDoc(spec_id='S')])
    result = generate(TRM(), pkg, CompanyProfile(name='测试'), str(tmp_path / 'out'),
                      spec_paths={'S': str(path)}, response_plans={'S': {
                          item: '拟实施响应（供投标负责人确认）：拟提交研究报告并组织评审。',
                      }})
    text, _ = _all_text(result.technical_path)
    assert '拟提交研究报告并组织评审' in text
    assert '对「提交研究报告」的具体响应' not in text
    assert '吞吐率 | 1000次 | 【待补充' in text


def test_ai_markdown_tables_are_word_tables_in_both_layouts():
    from jb_agents.writer import DraftSection
    from jb_docgen.builder import GenContext, _markdown_to_doc
    from jb_docgen.master_builder import _insert_drafts

    text = '| 阶段 | 工作内容 |\n| --- | --- |\n| 启动 | 需求核对 [来源:req-package] |\n\n---\n后续说明'
    plain = docx.Document()
    _markdown_to_doc(plain, text)
    master = docx.Document()
    master.add_paragraph('技术方案')
    ctx = GenContext(TRM(), PackageTRM(), CompanyProfile(), drafts=[DraftSection(title='计划', text=text)])
    _insert_drafts(master, ctx)
    for document in (plain, master):
        assert len(document.tables) == 1
        assert document.tables[0]._tbl.xpath('w:tr[1]/w:trPr/w:tblHeader')
        assert document.tables[0].cell(1, 1).text == '需求核对'
        assert any(p.text == '后续说明' for p in document.paragraphs)
        assert not any(p.text.startswith('|') or p.text == '---' for p in document.paragraphs)


def test_writer_chapter_title_is_not_duplicated():
    from jb_docgen.builder import iter_markdown

    blocks = list(iter_markdown('# 工作进度及保证措施\n## 总体进度\n拟分阶段执行。',
                                section_title='工作进度及保证措施'))
    assert not any(text == '工作进度及保证措施' for _, _, text in blocks)
    assert any(kind == 'h' and text == '总体进度' for kind, _, text in blocks)
