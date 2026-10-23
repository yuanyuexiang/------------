"""选材、审核、证据分类和长期材料的行为回归。"""
from datetime import date
from pathlib import Path

import docx
import pytest
from jb_agents.qualify import qualify
from jb_docgen import generate
from jb_docgen.evidence import certificates_for, scan_ids
from jb_kb.expiry import expiry_report
from jb_kb.models import Certificate, CompanyProfile, FinancialYear, Performance, Person
from jb_kb.models import TestReport as Report
from jb_kb.selection import MaterialSelection, certificate_validity, select_profile
from jb_parser.trm import TRM, PackageTRM, QualificationItem
from jb_rules import Context, review


def test_only_approved_facts_reach_generation_and_qualification(tmp_path):
    cp = CompanyProfile(name='T', personnel=[Person(name='未审核人员', status='draft')],
                        performances=[Performance(project='未审核业绩', status='draft')])
    pkg = PackageTRM(pkg_no='包1', qualification=[QualificationItem(performance_req='不少于1项')])
    trm = TRM(packages=[pkg])
    res = generate(trm, pkg, cp, str(tmp_path))
    text = '\n'.join(p.text for p in docx.Document(res.technical_path).paragraphs)
    assert '未审核人员' not in text and '未审核业绩' not in text
    assert qualify(trm, cp, use_llm=False).packages[0].verdict == '不可投'
    assert cp.personnel[0].status == 'draft'  # 不修改知识库原对象


def test_explicit_selection_rejects_foreign_and_unapproved_ids():
    cp = CompanyProfile(name='T', personnel=[Person(id='a', name='A'), Person(id='b', name='B', status='draft')])
    assert [p.name for p in select_profile(cp, MaterialSelection(personnel=['a'])).personnel] == ['A']
    for ids in [['b'], ['foreign'], ['a', 'a']]:
        with pytest.raises(ValueError):
            select_profile(cp, MaterialSelection(personnel=ids))
    assert select_profile(cp, MaterialSelection(personnel=[])).personnel == []


def test_financial_default_does_not_treat_comparative_data_as_audit():
    cp = CompanyProfile(name='T', financials=[FinancialYear(year=str(y)) for y in (2021, 2022, 2023, 2024)] +
                        [FinancialYear(year='2025', report_kind='comparative')])
    assert [f.year for f in select_profile(cp).financials] == ['2024', '2023', '2022']


def test_long_term_and_snapshot_materials_do_not_need_fake_expiry_dates():
    cp = CompanyProfile(name='T', certificates=[
        Certificate(name='长期许可', validity_kind='permanent'),
        Certificate(name='信用查询', validity_kind='point_in_time', valid_from='2026-10-20'),
        Certificate(name='过期认证', valid_until='2026-06-01')])
    assert certificate_validity(cp.certificates[0]) == '长期'
    assert '2026-10-20' in certificate_validity(cp.certificates[1])
    items = expiry_report(cp, on_date=date(2026, 10, 22))
    assert [(i.name, i.level) for i in items] == [('过期认证', 'expired')]
    findings = review(Context(TRM(), PackageTRM(), cp, open_date='2026-10-22')).findings
    assert any(f.rule_id == 'SG-13' and '过期认证' in f.message for f in findings)
    assert not any(f.rule_id == 'SG-13b' for f in findings)


def test_same_stem_scan_pdf_preferred_but_unrelated_word_still_reports_missing():
    paths = {'word': '/a/123456789abc_项目A.docx', 'pdf': '/a/abcdef123456_项目A.pdf', 'other': '/a/其他.docx'}
    assert scan_ids(['word', 'pdf', 'other', 'pdf'], paths) == ['pdf', 'other']


def test_multi_registration_and_query_materials_are_routed():
    cp = CompanyProfile(name='T', certificates=[
        Certificate(id='a', name='营业执照', includes_tax_registration=True),
        Certificate(id='b', name='基本存款账户信息'),
        Certificate(id='c', name='中国裁判文书网查询截图', cert_type='信用查询材料')])
    assert certificates_for(cp, 'tax')[0].id == 'a'
    assert certificates_for(cp, 'bank')[0].id == 'b'
    assert certificates_for(cp, 'credit')[0].id == 'c'
    assert certificates_for(cp, 'other') == []


def test_demo_watermark_survives_generation(tmp_path):
    cp = CompanyProfile(name='演示企业', is_demo=True)
    res = generate(TRM(), PackageTRM(), cp, str(tmp_path))
    for path in (res.commercial_path, res.technical_path):
        d = docx.Document(path)
        assert '模拟数据' in d.paragraphs[0].text
        assert '不可用于正式投标' in d.sections[0].header.paragraphs[0].text
        assert Path(path).is_file()


def test_simulated_project_plan_only_used_by_demo_profiles(tmp_path):
    from jb_kb.models import Boilerplate
    cp = CompanyProfile(name='模拟企业', is_demo=True, boilerplates=[Boilerplate(
        topic='演示项目方案', title='模拟研究方案', text='模拟实验与联调流程', approved=True)])
    pkg = PackageTRM(scope='研究服务')
    trm = TRM(packages=[pkg])
    demo = generate(trm, pkg, cp, str(tmp_path / 'demo'))
    assert '模拟实验与联调流程' in '\n'.join(p.text for p in docx.Document(demo.technical_path).paragraphs)
    cp.is_demo = False
    real = generate(trm, pkg, cp, str(tmp_path / 'real'))
    assert '模拟实验与联调流程' not in '\n'.join(p.text for p in docx.Document(real.technical_path).paragraphs)
    assert any('总体技术/服务方案' in todo for todo in real.docx_todos[Path(real.technical_path).name])


def test_excluding_required_evidence_does_not_remove_qualification_gate():
    pkg = PackageTRM(qualification=[QualificationItem(report_required=True)])
    cp = CompanyProfile(name='T', test_reports=[Report(id='r', name='检测报告')])
    selected = select_profile(cp, MaterialSelection(test_reports=[]))
    result = review(Context(TRM(), pkg, selected))
    assert any(f.rule_id == 'Q-02' and f.level == '否决' for f in result.findings)
    result = review(Context(TRM(), pkg, select_profile(cp)))
    assert not any(f.rule_id == 'Q-02' and f.level == '否决' for f in result.findings)
    assert any(f.rule_id == 'Q-02' and f.level == '需人工' for f in result.findings)


def test_excluding_required_iso_certificate_is_reported_as_missing():
    pkg = PackageTRM(qualification=[QualificationItem(other_reqs={'企业资质条件': '须具备ISO27001认证'})])
    cp = CompanyProfile(name='T', certificates=[Certificate(name='质量管理体系认证证书', validity_kind='permanent')])
    result = review(Context(TRM(), pkg, cp))
    assert any(f.rule_id == 'Q-02' and f.level == '否决' and '信息安全管理' in f.message
               for f in result.findings)
