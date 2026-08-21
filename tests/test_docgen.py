"""文档引擎：技术参数自动填写 + 商务/技术文件生成（合成数据，不依赖样本）。"""
import os

from jb_docgen import generate
from jb_docgen.placeholders import scan_docx
from jb_docgen.techparams import fill_spec, split_compound
from jb_kb.models import CompanyProfile, Performance, Product
from jb_parser.trm import TRM, FormatBlock, PackageTRM, SpecDoc, SpecParamRow


def _spec():
    return SpecDoc(spec_id="9999-500000001-00001", structured=True, param_rows=[
        SpecParamRow(row=1, name="主机", required="★包含主机1套，不小于128GB内存,4*1920GB SSD,8个GE电口,6个光口", star=True),
        SpecParamRow(row=2, name="授权", required="★配置网络设备授权不小于256个", star=True),
        SpecParamRow(row=3, name="功能", required="支持统一拓扑管理，拓扑自动生成和刷新"),
        SpecParamRow(row=4, name="衰减", required="≤0.21"),
    ])


def _product(**over):
    params = {"内存": "256GB", "SSD": "4*1920GB", "GE电口": "8个", "网络设备授权": "512个", "衰减系数": "13.2"}
    params.update(over)
    return Product(model="T-1", name="综合网管", category="综合网管", params=params,
                   features=["支持统一拓扑管理，拓扑自动生成和刷新，展示设备链路告警"])


def test_split_compound():
    assert len(split_compound("★包含主机1套，不小于128GB内存,4*1920GB SSD,8个GE电口")) == 4
    assert split_compound("不小于128GB内存") == ["不小于128GB内存"]


def test_fill_spec_verdicts():
    r = fill_spec(_spec(), _product())
    by_row = {x.row: x for x in r.responses}
    assert by_row[1].verdict == "missing" and "光口" in by_row[1].response   # 光口无参数→待补充，不冒充
    assert "内存 256GB" in by_row[1].response and "GE电口 8个" in by_row[1].response
    assert by_row[2].verdict == "satisfied" and by_row[2].response == "512个"
    assert by_row[3].verdict == "unknown"                                     # 功能项→人工确认
    assert by_row[4].verdict == "deviation"                                   # 13.2 > 0.21 → 偏差表
    assert len(r.deviations) == 1


def test_forbidden_phrase_never_written():
    r = fill_spec(_spec(), _product(**{"网络设备授权": "完全响应"}))
    resp = {x.row: x for x in r.responses}[2]
    assert resp.verdict == "missing" and "完全响应" not in resp.response.replace("禁止填「完全响应」", "")


def test_no_product_all_missing():
    r = fill_spec(_spec(), None)
    assert all(x.verdict == "missing" for x in r.responses)


def test_generate_docs(tmp_path):
    trm = TRM(batch_name="测试批次", batch_no="T-2026-01", format_blocks=[
        FormatBlock(title="投标函", paragraphs=["致：招标人", "1.我方已仔细研究了招标编号：    ，项目名称：   ，分标名称：   ，包名称：   招标文件的全部内容。"]),
        FormatBlock(title="法定代表人（单位负责人）授权委托书", paragraphs=["_____________（投标人名称）为中华人民共和国合法企业，法定地址：。", "________（营业执照法定代表人（单位负责人））特授权__________代表我方全权办理XX（招标编号）（分标编号及名称）（包号）（包名称）项目。"]),
    ])
    pkg = PackageTRM(sub_no="分标1", sub_name="综合网管", pkg_no="包1", project_name="测试项目", spec_docs=[_spec()])
    prof = CompanyProfile(name="测试企业", credit_code="9100X", address="南京市", legal_person="张三",
                          authorized_rep="李四", authorized_rep_title="经理",
                          performances=[Performance(project="项目A", evidence=["合同", "发票"], signed_date="2025-01-01")],
                          products=[_product()])
    res = generate(trm, pkg, prof, str(tmp_path))
    assert os.path.exists(res.commercial_path) and os.path.exists(res.technical_path)
    import docx
    text = "\n".join(p.text for p in docx.Document(res.commercial_path).paragraphs)
    assert "招标编号：T-2026-01，项目名称：测试项目，分标名称：综合网管，包名称：包1" in text   # 第六章原文填空
    assert "测试企业（投标人名称）" in text and "法定地址：南京市" in text and "张三" in text and "特授权李四" in text
    assert res.export_blocked                       # 扫描件等待补充 → 阻断导出
    assert all("【待补充：" in f"【待补充：{t}】" for t in scan_docx(res.technical_path))
    s = res.summary()
    assert s["tech_params"][0]["deviation"] == 1
